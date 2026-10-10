"""Build the BioExplainR iMeta/Q1 submission improvement roadmap.

This document is deliberately based on the installed package surface, source
modules, tests, research scripts and current manuscript. It is a project
planning document rather than a manuscript and should be revised after each
new frozen evaluation round.
"""

from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/BioExplainR-iMeta-Q1-submission-roadmap.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, font_size=8.6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    if widths:
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=font_size)
        set_cell_shading(cell, "285E75")
    header_trpr = table.rows[0]._tr.get_or_add_trPr()
    header_trpr.append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value, size=font_size)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F2F6F7")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Roadmap")
    p.add_run(text)
    return p


def add_callout(doc, title, text, fill="EAF3F5"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title + "\n")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("174B5D")
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    styles["Normal"].font.size = Pt(10.2)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.16
    for name, size, color in [
        ("Title", 22, "174B5D"),
        ("Subtitle", 12, "52656D"),
        ("Heading 1", 15, "174B5D"),
        ("Heading 2", 12, "285E75"),
        ("Heading 3", 10.5, "3E6875"),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)
    if "Code Roadmap" not in [s.name for s in styles]:
        code = styles.add_style("Code Roadmap", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "DejaVu Sans Mono"
        code._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        code.font.size = Pt(8.3)
        code.font.color.rgb = RGBColor.from_string("263238")
        code.paragraph_format.left_indent = Cm(.45)
        code.paragraph_format.right_indent = Cm(.25)
        code.paragraph_format.space_before = Pt(3)
        code.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("BioExplainR iMeta/Q1 submission roadmap | internal planning document | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def build():
    doc = Document()
    configure(doc)

    # Cover page
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("BioExplainR 面向 iMeta / 一区期刊的\n继续改造与投稿路线图")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("基于完整 R 包源码、测试、研究脚本、用户指南与当前论文交付的系统审查")
    doc.add_paragraph()
    cover = doc.add_table(rows=5, cols=2)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.style = "Light Shading Accent 1"
    cover_data = [
        ("软件版本", "BioExplainR 0.3.1"),
        ("审查日期", str(date.today())),
        ("主要目标", "iMeta；备选为当前投稿年度核验后的一区计算生物学/生物信息学期刊"),
        ("当前结论", "软件完成度高，论文证据仍不足以支撑 iMeta 的强主张"),
        ("路线原则", "停止以增加函数数量为主，改为独立多组学、claim-level 评价和真实任务证据"),
    ]
    for row, (key, value) in zip(cover.rows, cover_data):
        set_cell_text(row.cells[0], key, bold=True, size=9.5)
        set_cell_text(row.cells[1], value, size=9.5)
    doc.add_paragraph()
    add_callout(
        doc,
        "关于“80% 几率”的边界",
        "任何人都不能根据软件审查保证某个期刊的录用概率。本路线把 80% 定义为内部“可投稿准备度”：在没有 P0 阻断项、完成独立评价、公开复现和伦理/作者信息核验后，团队有足够证据把稿件投向 iMeta 或一个匹配的一区备选。它不是录用承诺，也不是编辑决策的统计预测。",
        "FFF2CC",
    )
    doc.add_page_break()

    heading(doc, "执行摘要", 1)
    para(doc, "BioExplainR 已经不是一个简单的 LLM wrapper。它有统一的 bio_context 中间表示、差异表达/富集/单细胞/PCA/图形/ExpressionSet 适配器、离线解释、可选模型调用、证据对象、声明验证、图形交叉核对、错误诊断、显式修复不变量、项目依赖图、报告和 Shiny 工作台。测试和包检查也已达到可以作为软件方法学论文基础的水平。")
    para(doc, "当前论文的核心证据主要证明了三件事：R 适配器可以从完整输入计算确定性 facts；受限模型协议可以把已经提供的 facts 复制到结构化 claims；受控错误注入能够被声明/图形/不变量检查器识别。这些是重要的工程和机制结果，但不是模型自由文本质量、证据支持质量、多组学整合收益或真实用户收益的证明。")
    para(doc, "要提高 iMeta 或其他一区期刊的竞争力，论文必须改成一个更窄、更强的科学主张：BioExplainR 是否能降低组学解释中的可检测统计错误、无支持生物学表述和证据错配，并在真实的整合组学工作流中留下可复核的证据链。")
    add_callout(doc, "最重要的决策", "不要继续优先增加教育题目、更多连接预设或更多简单表格适配器。下一轮开发资源应优先投入：1）typed claim ledger；2）空间转录组与结构层证据链；3）MultiAssayExperiment/SpatialExperiment 真实案例；4）冻结后的独立 claim-level benchmark；5）正规专家/用户任务研究；6）公开仓库、DOI、容器和 CI。", "E8F1E8")

    heading(doc, "一、审查范围与当前包画像", 1)
    para(doc, "本路线基于项目中可见的 R 包源码、NAMESPACE、DESCRIPTION、man/*.Rd、testthat 测试、inst/examples、架构/需求/评测/安装/工作台文档、研究脚本、当前 PAPER_DRAFT 与 SUPPLEMENTARY_MATERIAL，以及已经生成的用户指南和投稿审计文件。未把历史交接文件中的过期数字当作当前状态；以 2026-10-05 交接说明和当前文件内容为准。")
    add_table(doc, ["审查对象", "当前状态", "对投稿的意义"], [
        ("公开导出 API", "74 个函数；另有 S3 print/as.list/format 方法", "功能面很完整，但论文不能按函数数量证明科学创新"),
        ("核心中间表示", "bio_context：facts、top_features、plot_mapping、provenance、uncertainty、metadata", "是最值得形成论文中心贡献的数据契约"),
        ("对象覆盖", "DE、富集、Seurat、SCE、SE、ExpressionSet、PCA、ggplot、细胞组成、generic", "覆盖转录组和单细胞；尚未形成 iMeta 所需的跨模态主案例"),
        ("模型接口", "OpenAI-compatible、Anthropic、DeepSeek、Ollama、自定义预设", "协议层成熟；不等于模型质量已经被证明"),
        ("验证层", "显式 claims、证据 ID、图形映射、修复不变量", "机制贡献明确；自由文本和证据支持仍未完整评价"),
        ("运行记录", "警告、错误栈、对象摘要、包调用、版本和报告", "有利于可复现和审计；尚未形成公开 artifact 标准"),
        ("浏览器工作台", "Shiny 本地运行，连接间工作区隔离", "适合工作流展示；不是安全多租户服务"),
        ("软件验证", "525 项检查，524 通过、1 跳过；R CMD check OK", "软件质量基础强，但不是外部方法效果证据"),
        ("模型规模实验", "120 个程序生成案例、480 facts、三种输入、三轮重复", "主要是事实传递/序列化压力测试，不能作为自然数据泛化"),
        ("人类评价", "3 位评分者和 30 位功能评分的聚合材料", "当前缺少招募、伦理、同意、任务正确率、对照和原始评分审计"),
    ], widths=[3.3, 6.2, 7.1])

    heading(doc, "1.1 当前功能模块审计", 2)
    add_table(doc, ["模块", "已实现功能", "当前证据", "投稿定位"], [
        ("配置与 LLM", "bio_config、预设、端点、重试、OpenAI/Anthropic、脱敏", "mock 协议测试；少量真实调用", "基础设施，不是主要创新"),
        ("对象适配", "DESeq2/limma/edgeR、ORA/GSEA、Seurat/SCE/SE/ExpressionSet、PCA、ggplot", "真实 airway、PBMC3k、pasilla、ALL、bladderbatch", "核心方法；需要跨模态扩展"),
        ("解释", "R facts + 本地模板 + 可选模型分节解释", "离线测试和结构化模型实验", "中心工作流，但要用语义质量评价支撑"),
        ("证据", "用户证据、Europe PMC、GO、Reactome、KEGG、UniProt、缓存", "mock 端点和少量实端点", "需要从可追溯 ID 升级为支持关系评价"),
        ("验证", "claim values/directions/path/citation、图形 cross-check、fix invariants", "90 项声明挑战、9 项图形挑战、修复挑战", "最有辨识度的技术贡献"),
        ("诊断", "错误分类、栈、离线建议、可选模型建议", "规则单元测试", "需要真实错误 benchmark 才能成为结果"),
        ("项目", "依赖图、派生证据提示、跨结果一致性", "循环/未知依赖和示例", "适合承载多组学关系"),
        ("绘图与 Shiny", "PCA、UMAP、富集、热图和本地工作台", "浏览器 smoke test", "展示层；不能作为主要科学贡献"),
        ("教学评估", "题库、随机、专家 rubric、学习分析", "软件测试；没有正式学习实验", "应从主论文主线降为扩展或后续研究"),
    ], widths=[2.2, 6.0, 5.0, 3.4])

    heading(doc, "二、现阶段能发什么，为什么还不能直接冲 iMeta", 1)
    para(doc, "以当前软件和论文，最稳妥的描述是“高完成度、可复现的软件方法学工作，具有机制验证，但缺少独立的科学效果证据”。它已经明显超过只有 demo 和一个模型调用的软件短文，但还没有证明 BioExplainR 在真实组学研究中带来可量化的错误降低、时间收益、证据质量提高或跨模态整合价值。")
    add_table(doc, ["维度", "当前判断", "主要原因"], [
        ("软件工程", "强", "模块化、可选依赖、脱敏、mock、R CMD check、浏览器检查和用户指南完整"),
        ("可复现性", "强但未闭环", "有结果和哈希，但公开仓库、DOI、固定环境和 CI 仍需正式发布"),
        ("统计边界意识", "强", "raw p/padj、pseudoreplication、ORA/NES、batch 和因果边界写得清楚"),
        ("方法创新", "中上", "事实契约、claim validation、plot cross-check、fix invariants 的组合有特色"),
        ("模型能力证据", "弱到中", "facts_only 的 100% 是事实传递，不是模型独立推理；自由文本错误已被负例证明"),
        ("生物学价值", "中等偏弱", "目前以对象接入和解释工作流为主，缺少真实多组学科学问题"),
        ("人类收益", "不足", "当前功能评分不是正式用户实验；伦理和任务终点不完整"),
        ("iMeta 适配", "当前中等；空间/结构案例完成后可强", "iMeta 需要整合组学/生物学应用主线，当前 package 主线仍偏 R 解释层"),
    ], widths=[3.0, 3.0, 10.6])
    add_callout(doc, "投稿判断", "当前不建议把稿件包装成“LLM 生物信息学准确率达到 100%”。如果现在投稿，审稿人最可能指出：目标事实已由 R 计算并写入输入、案例是程序生成且与开发流程相关、证据只检查引用闭集、自由文本存在 shrinkage 语义错误、外部评分不是正式用户研究。", "FDECEC")

    heading(doc, "2.1 iMeta 的正确切入角度", 2)
    para(doc, "iMeta 目标应当不是“我有一个功能很多的 R 包”，而是“整合组学结果的解释容易发生事实、统计和证据错配；BioExplainR 用一个可审计的数据契约和 claim ledger 把这些风险分开，并在跨模态案例中验证风险是否下降”。")
    add_bullets(doc, [
        "主科学问题：在空间转录组与其他组学/结构层信息的联合解释中，确定性 facts、空间关系、证据关系和显式验证能否降低 LLM 的可检测错误？",
        "主技术贡献：将不同 R 对象和模态统一为带来源状态、样本对齐、设计信息、事实路径和不确定性的 context；将模型文本拆成可核对的 typed claims。",
        "主应用贡献：在一个真实空间组学案例上展示从 spot/区域接入、空间统计、候选靶点筛选、结构计算结果、证据追踪到报告输出的完整链条。",
        "必须明确的边界：不声称发现新的生物学机制，不声称模型具有普遍推理能力，不声称本地 callr 进程是安全沙箱，不声称引用 ID 证明了文章支持关系。",
    ])

    heading(doc, "三、阻断项清单：P0 必须解决，P1 强烈建议，P2 可延后", 1)
    add_table(doc, ["等级", "阻断项", "现有证据/代码位置", "通过标准"], [
        ("P0", "缺少真正独立的自由文本/生物学 claim benchmark", "bio_validate_explanation 只核对显式 value/direction/path/citation；当前规模实验把 facts 直接传给模型", "冻结外部 test set；至少 100 个专家标注 claims；报告 unsupported、statistical、biological、citation 指标和 CI"),
        ("P0", "缺少 iMeta 风格空间/跨模态真实案例", "当前 adapter 以 bulk RNA、单细胞、表达对象和结果表为主；没有 SpatialExperiment 或结构层主流程", "至少一个真实 SpatialExperiment/Seurat spatial 案例；空间统计、区域/样本关系、结构层候选和证据链全部可复现"),
        ("P0", "当前 100% 结果容易被解释为答案泄漏/事实复制", "research/run_main_scale_experiment.R:41-46, 200-231；目标 facts 来自 BioExplainR context 并写入 prompt", "将 transport fidelity 与 scientific interpretation 分为两个 endpoint；后者不把答案直接作为待复制 claim 输入"),
        ("P0", "证据只验证 ID 闭集，不验证支持关系", "R/explanation-validation.R:116-121；evidence-live 主要产生检索记录和定义", "建立支持/矛盾/无关/定义-only 标注；专家盲评 evidence-to-claim support precision/recall"),
        ("P0", "人类评价不能作为当前正式用户研究结果", "docs/PAPER_DRAFT.md:102-106；SUPPLEMENTARY S5 明确缺招募、伦理、同意、补偿和任务终点", "取得真实伦理审批/豁免和同意记录；预注册；随机对照或至少专家盲评的任务正确率和时间"),
        ("P0", "当前文档和审计仍有历史状态混杂", "README.md 仍保留 6 案例/54 请求；research/audit_results.R 仍核对旧 fact_scores.csv 225 行/54 调用", "所有 README、audit、manifest、论文、补充材料只引用一个冻结版本；旧结果移动到明确 historical 目录"),
        ("P1", "默认 JSON 序列化精度没有形成正式契约", "R/utils.R:.json 使用 digits=10；规模脚本另行使用 digits=17", "内部 facts JSON 默认保留 17 位或显式 precision；精度回归测试覆盖 PCA、极小 p 和零值"),
        ("P1", "typed claim schema 不够严格", "当前 claims 主要检查 fact/value/direction/evidence；自由 prose 仍未结构化", "加入 claim_type、fact_path、value、unit、direction、evidence_role、scope、status，拒绝未知字段或保留 unverified"),
        ("P1", "大对象隐私保护是约定，不是全面审计", "R/utils.R:.object_summary、bio_context_payload；top feature/metadata 仍可能含用户自定义标识", "敏感字段注入测试、payload allowlist、manifest 记录发送字段和 hash，明确脱敏边界"),
        ("P1", "修复执行不是 OS sandbox", "R/validate.R:163-167；Shiny 同样共享文件/网络权限", "论文删除安全沙箱暗示；研究代理使用 Docker/Apptainer/no-network/no-write baseline，并单独审计"),
        ("P1", "错误诊断和图形核对缺少自然案例评价", "当前主要是规则测试和人工注入 challenge", "真实或盲法构造的错误/图形集，报告 top-1/top-3、false-pass 和 not-checkable coverage"),
        ("P2", "教学和 Shiny 功能分散主线", "teaching.R、evaluation.R、shiny-app.R 功能齐全但不是同一科学终点", "主论文降为扩展模块；教学效果另文或补充研究，不用功能评分支撑主结论"),
    ], widths=[1.0, 4.0, 6.0, 5.6], font_size=8.0)

    heading(doc, "四、必须先做的软件改造", 1)
    para(doc, "下列改造是为了让 BioExplainR 能够被严格评价，而不是为了堆功能。每项都需要代码、测试、文档和研究记录同时更新。")

    heading(doc, "4.1 建立 typed claim ledger", 2)
    para(doc, "当前验证器能判断“某个声明的数字是否等于某个 fact”，但不能系统区分统计事实、统计解释、生物学关联、因果断言和证据支持。建议把 claims 从任意模型返回的 list 升级为固定 schema，并让报告显示每个 claim 的状态和检查范围。")
    add_table(doc, ["字段", "必须含义", "示例", "验证方式"], [
        ("id", "稳定的 claim 标识", "claim_001", "唯一且可追踪"),
        ("claim_type", "事实、统计解释、关联、因果、证据支持、设计、图形映射", "statistical_fact", "枚举值；未知类型为 unverified"),
        ("text", "模型或模板生成的文字", "999 features pass threshold", "不自动视为正确"),
        ("fact_path", "指向 context facts 的路径", "facts.n_significant", "路径存在且值类型可比较"),
        ("value/unit", "标量值和单位", "999 / features", "数字容差、字符串和单位检查"),
        ("direction", "仅在有符号效应时使用", "up", "和事实符号或预设方向比较"),
        ("scope", "样本、模态、比较和阈值范围", "airway; trt vs untrt", "与 metadata/provenance 对齐"),
        ("evidence", "证据 ID 和证据角色", "supports / definition_only", "闭集 + 专家支持关系"),
        ("status", "pass/fail/partial/unverified/not_assessable", "partial", "验证器填充，模型不能自行声明 pass"),
    ], widths=[2.0, 5.0, 4.0, 5.6], font_size=8.3)
    add_code(doc, "# 建议新增的公开接口方向（名称可调整）\nbio_claim_ledger(explanation)\nbio_validate_claims(explanation, policy = bio_claim_policy())\nbio_evidence_support(claim, evidence, annotation = c(\"supports\", \"contradicts\", \"definition_only\", \"irrelevant\"))\nbio_context_manifest(ctx, include_hashes = TRUE)")
    para(doc, "最低实现要求：模型输出中没有 fact_path 的自由文本必须显式标为 unverified；causal claim 默认不允许自动 pass；definition_only 证据不能满足 experiment_support；claim scope 必须和物种、组织、比较方向、模态一致。")

    heading(doc, "4.2 增加跨模态对象契约", 2)
    para(doc, "iMeta 的最小软件增量不是再支持一种结果表，而是能够处理“同一组样本的多个模态”。建议优先支持 SpatialExperiment、Seurat 的空间对象/图像和 MultiAssayExperiment，并把各模态 context 接入现有 bio_project。空间转录组应成为主入口，结构计算作为经过 provenance 和验证的外部证据层。")
    add_table(doc, ["能力", "最小实现", "不能省略的检查"], [
        ("模态注册", "transcriptome、methylation/CNV、proteome、metabolome、microbiome 至少两个", "每个模态的 class、source package、analysis type、feature unit"),
        ("样本对齐", "统一 sample map 和交集/缺失表", "重复 sample ID、模态缺失、样本顺序、配对关系"),
        ("设计信息", "group、batch、donor、timepoint、paired unit", "生物学重复不是 feature/cell；组别和 batch 可识别性"),
        ("结果依赖", "模态内和模态间依赖边", "派生结果不当作独立验证；同一输入的证据相关性"),
        ("跨模态结论", "agreement、discordance、modality-specific", "方向一致不等于因果；缺失模态不能被模型补全"),
        ("报告", "每模态 facts + joint checks + claim ledger", "所有结论能回到对象、路径和证据"),
    ], widths=[3.2, 6.0, 7.4])
    para(doc, "推荐的第一套真实案例：一个有匹配样本的转录组 + DNA 甲基化/CNV 或转录组 + 蛋白组数据集；如果能找到样本级匹配的微生物组 + 代谢组 + 转录组数据，iMeta 相关性更强。数据集必须在开发前冻结，记录下载 URL、版本、SHA-256、样本交集和所有排除规则。不要为了“多组学”把不匹配的公开数据拼在一起。")

    heading(doc, "4.3 空间转录组与分子对接：可以结合，但必须分层", 2)
    para(doc, "空间转录组和分子对接确实是当前很有吸引力的研究热点，但二者回答的是不同问题：空间转录组回答“哪些分子信号在什么组织位置、哪些细胞/区域共同出现”；分子对接回答“给定受体结构、配体结构和参数时，某种结合构象在计算上是否可行”。空间共定位加上较好的 docking score 不能直接证明真实结合、通路激活、药效或因果机制。")
    add_callout(doc, "推荐论文定位", "Spatially resolved and structure-aware evidence-constrained interpretation of omics results。空间层负责定位和组织上下文，结构层负责候选相互作用的计算可行性，文献/数据库负责知识背景；BioExplainR 负责把每层事实、限制和验证状态连接起来。", "EAF3F5")
    add_table(doc, ["层级", "建议支持", "必须计算的 facts", "不能自动推出"], [
        ("空间对象", "SpatialExperiment、Seurat spatial assay/image、空间 metadata", "spot/cell 数、切片/样本、坐标范围、组织覆盖、QC、区域/cluster、邻域图", "坐标接近不等于细胞互作或谱系关系"),
        ("空间统计", "已有空间变异基因、空间 DE、邻域/共定位、Moran's I/Geary's C 等结果", "统计方法、空间邻接定义、显著特征、效应方向、样本/切片分母", "单切片显著不等于跨患者可重复"),
        ("候选分子", "ligand/receptor、靶点、药物或代谢物候选表", "ID 类型、物种、区域表达、共表达、来源、筛选规则", "表达共现不等于物理结合"),
        ("结构计算", "AutoDock Vina/smina/GNINA 或 HADDOCK 等外部结果解析", "引擎版本、结构 ID、预处理、搜索空间、score 分布、pose cluster、RMSD、接触残基、对照/decoy", "docking score 不等于实验结合自由能或药效"),
        ("证据图", "空间 claim → 结构 claim → 文献/数据库 claim 的依赖图", "每条 claim 的类型、scope、evidence role、验证状态", "多层证据一致不自动证明机制"),
    ], widths=[2.6, 6.0, 5.3, 4.0], font_size=8.0)
    para(doc, "优先级建议：第一阶段先做 SpatialExperiment/Seurat spatial 的对象与结果适配；第二阶段接收外部 docking 输出和 provenance，不在 R 包内部手写 docking 引擎；第三阶段做空间-结构 claim graph 和独立专家 benchmark。这样可以利用热点，但不会把一个不可验证的结构模型变成软件的未经审查结论。")

    heading(doc, "4.3.1 空间转录组软件接口", 3)
    add_bullets(doc, [
        "新增 SpatialExperiment adapter：读取 spatialCoords、imgData、colData、assayNames、区域/cluster、sample/section/donor 信息；默认不向模型发送完整坐标、图像或 barcode。",
        "支持已有结果表：空间变异基因、空间 DE、区域 marker、邻域富集、ligand-receptor 候选和 deconvolution 输出；记录方法、空间邻接定义和统计分母。",
        "新增空间 facts：每个 section 的 spot/cell 数、坐标范围、组织覆盖率、缺失坐标、每个样本/区域的观察数、空间统计阈值和多重检验方法。",
        "新增 spatial cross-check：检查 spot/region ID 对齐、坐标和 metadata 对齐、图形映射、区域子集和结果表来源；复杂空间图形无法核查时返回 not_checkable。",
        "所有空间生物学 claim 都要带 scope：sample/section、region、cell type、species、assay 和 comparison；没有 scope 的 claim 不自动通过。",
    ])
    add_code(doc, "# 建议接口方向；具体名称可以在 API 设计评审后确定\nbio_spatial_context(x, sample_col = NULL, section_col = NULL, region_col = NULL)\nbio_spatial_neighbors(x, method = c(\"stored\", \"knn\", \"radius\"))\nbio_spatial_cross_check(plot, source, allow_subset = FALSE)\nbio_docking_context(result, receptor = NULL, ligand = NULL, provenance = list())\nbio_integrate_spatial_structure(spatial_ctx, docking_ctx, evidence = NULL)")

    heading(doc, "4.3.2 分子对接不能只写成一个 score", 3)
    add_bullets(doc, [
        "区分小分子-蛋白对接和蛋白-蛋白对接；二者的输入、搜索空间、评分和验证方式不同。",
        "R 包负责读取和审计外部引擎输出，不负责在模型生成代码的情况下自动执行 docking。外部程序、容器、命令、结构文件 hash 和参数必须进入 provenance。",
        "保存 top pose、score 分布、聚类大小、RMSD、界面残基、氢键/疏水接触和 decoy/negative-control 结果；只有一个最佳 score 不足以形成 claim。",
        "明确 docking score 的尺度和引擎版本，不能跨引擎直接排序；不要把 score 直接写成 binding affinity 或 ΔG。",
        "候选选择必须在 docking 前冻结，或使用独立的候选发现集和 test 集；否则空间表达高、文献强和 docking 好的候选会产生严重选择偏倚。",
        "结构缺失、低分辨率、配体质子化、柔性、口袋定义和蛋白构象都会影响结果，均应进入 uncertainty。",
    ])

    heading(doc, "4.3.3 空间 + docking 的主案例设计", 3)
    para(doc, "一个合格的主案例应是一个可审计的候选优先级工作流，而不是“空间图 + docking 图”的拼图。建议流程如下：")
    add_numbered(doc, [
        "从冻结的空间数据中定义样本、切片、区域和生物学重复，预先规定 QC、空间邻接和显著性阈值。",
        "从空间差异或空间变异结果得到候选 ligand/receptor/target；记录全部候选和排除原因，不只保存最后一个漂亮候选。",
        "将基因/蛋白/配体 ID 映射到物种一致的结构记录，记录 PDB/mmCIF/SDF 的版本、来源、结构质量和 preprocessing。",
        "在独立计算环境中运行 docking 或读取已锁定的 docking 结果，使用 decoys、negative controls、重复搜索或多引擎敏感性分析。",
        "把空间定位、表达/蛋白事实、结构接触和文献证据分别写成 typed claims，并由验证器检查 scope、ID、数值和 evidence role。",
        "输出候选优先级和限制，而不是“已证明结合/激活/治疗有效”；必要时提出原位杂交、蛋白验证、结合实验或扰动实验作为后续验证。",
    ])

    heading(doc, "4.3.4 必须加入的负对照和评价", 3)
    add_table(doc, ["挑战", "目的", "最低结果"], [
        ("随机空间标签/打乱区域", "检查空间共定位 claim 是否依赖真实结构", "空间统计显著性和 false-positive rate"),
        ("结构 decoy/非靶蛋白", "检查 docking score 是否仅反映结构或口袋偏好", "target-vs-decoy discrimination 和 CI"),
        ("无关配体/负配体", "检查模型是否把任何 ligand 都写成支持", "wrong-binding claim rate"),
        ("物种/组织错配证据", "检查 evidence scope", "wrong-species/wrong-tissue rejection rate"),
        ("表达高但 docking 差", "检查跨层冲突是否被保留", "conflict correctly reported，不强行整合"),
        ("docking 好但空间不表达", "检查结构证据是否被误写成组织机制", "dataset-support claim must fail or remain partial"),
        ("不同引擎/参数", "检查 score 的方法依赖", "sensitivity table，不合并成单一真值"),
    ], widths=[5.0, 7.0, 5.0], font_size=8.0)

    heading(doc, "4.4 修复内部序列化精度并固定 artifact manifest", 2)
    para(doc, "当前 R/utils.R 的 .json 使用 digits=10，而规模实验脚本另行使用 digits=17。这会导致包默认工作流和研究脚本的数值传递契约不一致。应把精度作为显式的内部协议，不依赖模型自行恢复小数。")
    add_bullets(doc, [
        "把用于模型和 manifest 的 JSON 精度固定为 17 位；显示文本可以继续使用 .fmt_num() 舍入。",
        "在 bio_context_payload() 中记录 serialization_version、precision_digits、payload_chars、input_summary_hash 和 facts_hash。",
        "对 PCA variance_explained、极小 p、零值、整数和字符串分别做回归测试。",
        "每次模型请求保存 prompt hash、context hash、config public hash、model label、provider host、时间、usage 和 response hash；不保存 key。",
        "研究评分程序读取同一 manifest，不再在不同脚本中重复定义精度和容差。",
    ])
    add_code(doc, "# 目标：同一 facts 在包、实验和报告中的数值来源一致\nmanifest <- bio_context_manifest(ctx, include_hashes = TRUE)\npayload <- bio_context_payload(ctx, precision_digits = 17)\nstopifnot(identical(manifest$facts_hash, hash_facts(payload$facts)))")

    heading(doc, "4.5 明确隐私和执行安全边界", 2)
    para(doc, "包目前做了很多有价值的脱敏，但它不是自动隐私认证系统；Shiny 和 callr 都与当前用户共享文件/网络权限。论文应该避免“安全沙箱”“隐私保证”等泛化措辞，改成可审计的发送字段和运行权限声明。")
    add_bullets(doc, [
        "增加 payload allowlist：只允许 facts、有限 top features、study metadata 的明确字段进入模型；对 feature/cluster/term 标签提供 hash 或用户选择。",
        "增加 secret/PII/样本名/条码注入测试；测试不仅检查常见 PATIENT_ 字符串，还检查用户自定义字段和路径。",
        "报告对象和原始代码分开：默认 report 不含值，但 bio_record 对象本身仍可能含返回值和完整代码，文档要明确这一点。",
        "研究基准中的代码代理运行必须通过 Docker/Apptainer 或等价 OS 层策略，禁用网络和写入项目之外的路径；BioExplainR 的 callr 结果不能代替该基准。",
        "将 host = 0.0.0.0 的工作台只作为高级部署说明，补充认证、TLS、反向代理和单租户约束。",
    ])

    heading(doc, "五、独立评价设计：从事实传递升级为科学效果", 1)
    para(doc, "现有 120 案例/480 facts 保留，但在新论文中必须明确归类为 protocol fidelity/serialization stress test。它可以证明 facts 被传递后没有被模型改写，不能证明模型独立完成统计解释。新的评价需要把“答案已给出”与“模型根据事实作出正确判断”分开。")

    heading(doc, "5.1 五层评价终点", 2)
    add_table(doc, ["层级", "问题", "输入是否含目标答案", "主要终点", "论文位置"], [
        ("A 软件接入", "adapter 是否提取正确 facts", "否；参考值独立计算", "exact agreement、缺失和错误率", "主文方法/结果"),
        ("B 协议保真", "模型是否复制已提供 facts", "是；明确标注为 transport", "coverage、exact match、precision", "主文次要结果/补充"),
        ("C 统计判断", "模型能否正确区分 raw/padj、效应、重复单位、batch、ORA/GSEA", "提供情境和 facts，不提供目标标签", "专家金标准 accuracy、unsupported rate", "主文主要结果"),
        ("D 证据支持", "文献/数据库是否真正支持 claim", "提供混合的支持、矛盾、定义-only 证据", "support precision/recall、专家一致性", "主文主要结果"),
        ("E 工作流收益", "研究者是否更快、更准、更校准", "实验组/对照组不同工作流", "任务正确率、时间、信心校准、负荷", "主文或独立用户研究"),
    ], widths=[2.0, 5.0, 4.0, 5.0, 3.2], font_size=8.2)

    heading(doc, "5.2 独立 claim-level benchmark", 2)
    para(doc, "建议最低规模：100–200 个 claim，来自 8–12 个真实公开数据/分析对象，至少覆盖 bulk RNA-seq、单细胞和一个跨模态案例。每个 claim 由两名以上领域专家在模型运行前标注。测试集不能参与 adapter、prompt 或规则开发。")
    add_table(doc, ["任务家族", "建议样本量", "错误类型", "最低报告"], [
        ("差异表达", "30–40 claims", "raw vs adjusted p、方向、阈值、NA padj、effect vs significance", "accuracy、false claim rate、per-case CI"),
        ("富集/GSEA", "20–30 claims", "ORA activation、NES direction、background、redundancy", "statistical interpretation accuracy"),
        ("单细胞设计", "20–30 claims", "cell vs donor、batch、cluster marker、UMAP distance", "design error detection"),
        ("图形映射", "20–30 claims", "raw p、axis transform、subset、color/size", "mapping pass/fail、false pass"),
        ("跨模态", "30–50 claims", "sample mismatch、discordant directions、derived evidence、causal overclaim", "joint claim accuracy and support"),
        ("空间 + 结构", "30–50 claims", "空间共定位、区域分母、ligand/receptor scope、docking score、decoy、结构证据过度解释", "spatial/structure claim accuracy、false mechanism rate、conflict handling"),
        ("错误诊断", "40–60 cases", "missing package/object/column、dimension、S4、ID、factor、counts、NA", "top-1/top-3 root cause accuracy"),
    ], widths=[3.2, 3.0, 7.0, 4.0], font_size=8.1)
    para(doc, "比较方法至少包括：普通模型 + 相同问题文本、BioExplainR 本地模板、BioExplainR facts-only、BioExplainR facts + evidence、facts + validator/revision。所有方法固定模型版本、温度、token、prompt 和成本记录。普通模型不能被故意给弱输入；如果输入预算不同，必须把它作为预注册的资源约束并报告。")

    heading(doc, "5.3 证据支持关系评价", 2)
    para(doc, "当前 evidence_used 只检查模型是否引用了供应集合中的 ID。下一步要为每个 claim 构造四类 evidence：支持、矛盾、仅定义、无关。专家在不看模型方法的情况下判断 claim 是否得到支持，并记录物种、组织、实验条件和结论强度。")
    add_bullets(doc, [
        "不要把 Europe PMC 的搜索命中当作 support；文献标题和全文结论必须分开。",
        "GO/Reactome/KEGG term definition 只能满足 definition_only，不能满足 pathway_activity 或 dataset_effect。",
        "UniProt annotation 只能支持 annotation/membership 类型 claim，不能直接支持处理组差异或因果。",
        "报告 evidence precision、support recall、wrong-species rate、definition-as-support rate 和 citation coverage。",
        "至少两名专家独立标注，报告 Krippendorff alpha 或适合的多评分者可靠性指标；exact agreement 作为补充。",
    ])

    heading(doc, "5.4 错误诊断、图形和修复评价", 2)
    add_table(doc, ["组件", "现有状态", "新增评价", "验收线"], [
        ("错误诊断", "离线规则单测", "60 个真实/盲法构造错误；根因由专家预标注；模型和规则盲评", "top-1 ≥ 0.75，top-3 ≥ 0.90；未知错误不强行归类"),
        ("图形核对", "9 个预设 challenge", "40–60 个真实图/注入图；包括 facet、coord、subset、multi-layer、label", "false-pass < 5%；not_checkable 单独报告"),
        ("修复验证", "显式不变量和 callr", "错误方向、行数、列、design、输出类型的候选修复集", "运行率和语义率分开；无 invariants 不计成功"),
        ("声明验证", "value/path/direction/citation", "typed claim ledger + 语义负例 + 未声明 prose", "对注入错误 recall ≥ 0.95；自由 prose 保持 unverified"),
    ], widths=[2.5, 4.0, 7.2, 3.5], font_size=8.1)

    heading(doc, "六、真实多组学案例：iMeta 送审的必要主结果", 1)
    para(doc, "主案例必须有一个真正的生物学问题，而不是把多个对象并排展示。推荐结构是“同一批样本的两个或三个模态，存在缺失样本、批次或方向不一致的现实困难；BioExplainR 识别这些困难并限制模型结论”。")
    add_table(doc, ["主案例章节", "必须展示", "避免的写法"], [
        ("数据与设计", "样本交集、模态缺失、donor/group/timepoint/batch、独立重复", "只报每个对象的 feature 数量"),
        ("对象接入", "MultiAssayExperiment/SCE/SE/结果表如何映射到 context", "声称支持所有上游版本"),
        ("模态内 facts", "每个模态的差异、QC、阈值、来源和不确定性", "把不同模态的 p 值简单合并"),
        ("跨模态整合", "agreement、discordance、derived edges、样本/feature 对齐", "方向一致就写成共同机制"),
        ("证据", "支持/矛盾/定义-only 证据和适用范围", "把数据库定义写成实验激活"),
        ("模型输出", "原始 prose、typed claims、验证状态和人工修订", "只展示完美示例"),
        ("结果价值", "错误率、时间、专家判断或用户任务收益", "用 100% facts transfer 代替生物学效果"),
    ], widths=[3.5, 8.2, 5.5], font_size=8.3)
    para(doc, "可接受的第一版不必覆盖所有 omics。一个严格完成的 transcriptome + methylation/CNV 或 transcriptome + proteome 案例，优于五个没有样本对齐和专家评价的模态 demo。若选择微生物组/代谢组/转录组，应提前确定 feature ID、物种、批次和 compositionality 的解释边界。")

    heading(doc, "七、正式人类评价设计", 1)
    para(doc, "当前 30 位功能评分和 3 位外部评分者的聚合数据只能作为 pilot 或可接受性信号。若要把用户收益写进一区论文，应在收集前完成伦理审查/豁免判断、知情同意、招募标准、补偿、数据保留和预注册。")
    add_table(doc, ["设计元素", "最低要求"], [
        ("人群", "初级或中级生物信息学研究者；记录 R/生信背景和先验能力"),
        ("分组", "BioExplainR 工作流 vs 原始 R 输出或普通 LLM；同一数据和时间预算"),
        ("任务", "差异表达、富集/单细胞设计、图形判断、一个跨模态任务；加入未见数据迁移"),
        ("终点", "主要：任务正确率；次要：完成时间、信心、信心-正确性校准、NASA-TLX/SUS"),
        ("盲法", "评分者不知道组别；答案密钥和分组表分开保管"),
        ("样本量", "由预注册效应和功效计算决定；不要用当前便利样本的 91.98 分倒推效果"),
        ("统计", "pre/post 或任务分数使用混合效应模型/预设分析；报告效应量和置信区间"),
        ("伦理", "真实审批或豁免编号、同意、招募、补偿、匿名化和数据管理"),
    ], widths=[4.0, 13.2])
    para(doc, "如果短期内无法完成正式用户研究，论文应明确不声称用户收益，把教学模块和当前功能评分放入补充材料或另一个后续研究。iMeta 的主论文可以只依靠独立专家 claim benchmark + 真实多组学案例，但不能同时使用“帮助初级研究者”作为已证实结论。")

    heading(doc, "八、统计分析与预注册方案", 1)
    add_table(doc, ["项目", "建议预注册内容"], [
        ("主要终点", "unsupported biological claim rate 或 statistical/design error rate；定义分母和 claim 单位"),
        ("主要比较", "BioExplainR facts+evidence+validator vs fair baseline；模型、版本、温度和 prompt 冻结"),
        ("事实终点", "adapter agreement 与 transport fidelity 单独报告，不能合并为模型 accuracy"),
        ("专家评分", "每个 claim 的 0/1 或 0–2 标准、双评者、冲突裁决和 alpha/ICC"),
        ("置信区间", "按 case/bootstrap 或 mixed-effects 结构计算，不把重复调用当独立生物样本"),
        ("多重比较", "主终点只设少数；其余为探索性并报告 BH 或不调整说明"),
        ("成本", "输入/输出 tokens、网络时间、失败重试、每个正确 claim 成本"),
        ("缺失", "模型不回答计为未完成，不填补；解析失败与事实错误分别统计"),
        ("开发隔离", "开发案例、提示调参、阈值选择和最终 test case 的时间线可审计"),
    ], widths=[4.0, 13.2])
    para(doc, "建议的主要效应表达：相对于公平 baseline，BioExplainR 将 unsupported/incorrect claim 的风险降低多少；同时报告 absolute risk difference、relative risk、95% CI 和每个任务家族结果。不要把 480 个共享程序生成事实看成 480 个独立样本进行显著性检验。")

    heading(doc, "九、论文重构方案", 1)
    heading(doc, "9.1 推荐中心题目", 2)
    add_callout(doc, "英文题目候选", "BioExplainR: deterministic, evidence-constrained and claim-verifiable interpretation of integrated omics results", "EAF3F5")
    para(doc, "中文对应：BioExplainR：面向整合组学结果的确定性事实、证据约束与可验证声明解释框架。题目中的 integrated omics 只有在 MultiAssayExperiment/真实跨模态案例完成后才能使用；否则改为 omics R analysis objects。")

    heading(doc, "9.2 主文结构", 2)
    add_table(doc, ["章节", "应写什么", "不应写什么"], [
        ("摘要", "问题、架构、跨模态案例、独立 benchmark、主要风险降低结果", "74 个函数清单、模型 100% 复制事实"),
        ("引言", "组学解释的事实/证据/因果错配；现有 LLM 工具缺口；本文假设", "泛泛介绍 AI 热点"),
        ("方法", "context contract、typed claims、evidence roles、验证政策、数据和盲法", "把结果数字写进方法"),
        ("结果 1", "软件接入和确定性事实 conformance", "把 R 事实当模型能力"),
        ("结果 2", "独立 claim benchmark 和公平 baseline", "只展示成功案例"),
        ("结果 3", "空间转录组 + 结构层案例和冲突识别", "把共定位或 docking 写成已证实机制"),
        ("结果 4", "专家/用户任务结果（若完成）", "用便利功能评分替代任务正确率"),
        ("讨论", "适用边界、未验证 prose、证据支持限制、模型漂移、隐私和 sandbox", "宣称临床、教学或因果收益"),
    ], widths=[3.0, 8.2, 6.0], font_size=8.3)

    heading(doc, "9.3 图表设计", 2)
    add_table(doc, ["图/表", "建议内容"], [
        ("Figure 1", "跨模态 R objects → contexts → fact/claim/evidence graph → validation/report；标明 local/model/verified 边界"),
        ("Figure 2", "公平 baseline 的 statistical/design/unsupported claim 错误率；显示独立 test set 和 CI"),
        ("Figure 3", "真实多组学案例：sample map、模态内结果、跨模态一致/冲突、验证状态"),
        ("Figure 4", "证据角色和 claim validation 的 confusion matrix / precision-recall；不是只画 100% pass"),
        ("Table 1", "与相关工具的设计差异，明确不是性能排名"),
        ("Table 2", "软件 conformance 和版本矩阵"),
        ("Table 3", "独立 claim benchmark 的逐任务结果"),
        ("Table 4", "人类专家/用户评价设计和效应估计"),
        ("Supplement", "所有 prompts、case manifest、逐 claim labels、审计日志、原始脱敏输出和复现环境"),
    ], widths=[3.0, 14.2], font_size=8.4)

    heading(doc, "十、逐文件改造清单", 1)
    add_table(doc, ["文件/目录", "改动", "完成标志"], [
        ("R/utils.R", "统一 JSON precision、hash、manifest、脱敏字段策略", "新增 precision/hash 回归测试；研究和包使用同一序列化"),
        ("R/context.R", "扩展 context provenance、modalities、sample map 和 manifest", "MultiAssay/context conformance 测试通过"),
        ("R/adapters-spatial.R（新增）", "SpatialExperiment/Seurat spatial、坐标/区域/QC/空间结果 adapter", "真实切片对象、空间结果表和脱敏 payload 通过"),
        ("R/docking.R（新增）", "读取 docking 输出、结构/引擎/参数 provenance、pose/decoy facts；不内置未经审计的引擎执行", "多引擎/负对照/结构 hash 的 conformance 测试通过"),
        ("R/explanation-validation.R", "typed claim ledger、claim policy、causal default unverified", "注入错误 recall 和自由 prose 状态正确"),
        ("R/evidence.R / evidence-live.R", "evidence role、support annotation、scope/organism matching", "support/contradict/definition-only 单测和专家数据接口"),
        ("R/project.R", "跨模态依赖、sample alignment、conflict/derived evidence", "真实多组学项目端到端报告"),
        ("R/validate.R", "保留 callr 警告；增加外部 sandbox backend 接口或明确研究端隔离", "安全边界文档与 sandbox 研究基线"),
        ("R/record.R / runtime.R", "artifact manifest、session/config/prompt/response hash", "报告可重建且无 key/PII"),
        ("R/plots.R / plot-cross-check.R", "扩展代表性 plot mappings；完整覆盖和 false-pass 统计", "自然/注入图 benchmark"),
        ("tests/testthat", "增加 multiomics、precision、claim、evidence support、privacy、manifest 测试", "R CMD check 和 CI 全绿"),
        ("research/", "新建 frozen benchmark、label schema、expert adjudication、independent test runner", "开发/测试集隔离并有 manifest"),
        ("research/audit_results.R", "改为审计 main_scale_480 和新 benchmark；历史结果单独入口", "不会再读取旧 225/54 结果"),
        ("README.md / docs/ROADMAP.md", "删除历史数字或明确 historical；更新当前 525/120/480 状态", "grep 不再发现无标签旧主实验结论"),
        ("DESCRIPTION", "补真实作者、维护者、URL、BugReports、版本和 DOI 发布信息", "可公开安装和引用"),
        ("GitHub/Zenodo/CI", "公开仓库、release、DOI、R CMD check matrix、artifact upload", "第三方按 README 从零复现"),
    ], widths=[4.1, 8.6, 4.5], font_size=8.0)

    heading(doc, "十一、按周执行计划", 1)
    add_table(doc, ["时间", "工作包", "交付物", "停/继续判断"], [
        ("第 0 周", "冻结 0.3.1、清理历史状态、公开发布准备、修复审计脚本", "clean branch、版本 manifest、单一结果目录", "若旧数字仍混入论文，停止投稿路线"),
        ("第 1–2 周", "precision + manifest + typed claim ledger skeleton", "新 API、schema、单测、示例报告", "P0 软件契约通过才进入 benchmark"),
        ("第 2–4 周", "MultiAssayExperiment/sample map/context 适配", "一个真实多组学对象和端到端 project report", "若数据不匹配，换数据，不拼接"),
        ("第 3–5 周", "SpatialExperiment/Seurat spatial 与 docking provenance adapter", "空间 context、结构 context、negative-control schema", "结构层只做结果审计，不自动运行不受控 docking"),
        ("第 3–6 周", "冻结 claim benchmark 和 evidence support labels", "case JSON、答案标准、专家标注模板、预注册", "标签一致性不足则先修 rubric"),
        ("第 5–8 周", "运行公平模型 baseline 和 BioExplainR 条件", "原始 prompts/responses、cost、latency、逐 claim scores", "任何调参后 test 集必须重冻"),
        ("第 6–9 周", "专家盲评、错误诊断/图形 benchmark", "匿名评分、alpha/ICC、false-pass 分析", "不能把方便样本结果写成用户收益"),
        ("第 8–10 周", "空间转录组 + 结构层主案例和冲突/证据图", "Figure 1/3、Table 2/3 草稿", "没有真实空间主案例则转投软件期刊路线"),
        ("第 9–11 周", "用户任务研究或正式伦理材料", "预注册、审批/豁免、任务数据和分析脚本", "伦理未完成则不报告 user benefit"),
        ("第 11–13 周", "统计分析、论文重写、补充材料和 artifact", "英文主稿、supplement、repo、DOI", "全部主张能回到原始记录"),
        ("第 13–14 周", "第三方复现、内部红队审稿、期刊格式", "reproducibility report、response-to-reviewer 预案", "P0/P1 未清零则不投 iMeta"),
    ], widths=[2.2, 6.2, 6.0, 3.0], font_size=8.0)

    heading(doc, "十二、iMeta 与一区备选投稿决策树", 1)
    para(doc, "期刊分区会随年份、数据库和学科类别变化，投稿前必须用当年 JCR/CAS/期刊官网核验。以下是基于主题匹配和方法学要求的路线选择，不是对分区或录用的保证。")
    add_table(doc, ["完成状态", "首选", "备选", "稿件叙事"], [
        ("多组学案例 + 独立 claim benchmark + evidence support + 正规用户/专家研究", "iMeta", "Briefings in Bioinformatics、PLOS Computational Biology（按当年分区和 scope 核验）", "integrated omics reliability layer"),
        ("多组学案例 + 强 benchmark，但没有正式用户实验", "iMeta 或 NAR Genomics and Bioinformatics", "Bioinformatics", "reproducible computational method with independent validation"),
        ("只有转录组/单细胞 + 强 claim benchmark", "Bioinformatics / NAR Genomics and Bioinformatics", "PLOS Computational Biology", "verifiable interpretation of bioinformatics objects"),
        ("只有软件 QA、transport fidelity 和机制 challenge", "Bioinformatics Advances / BMC Bioinformatics 等软件方法期刊", "不应把 iMeta 当作唯一目标", "software framework and reproducibility"),
    ], widths=[6.0, 3.2, 5.5, 4.0], font_size=8.0)
    add_callout(doc, "最低 iMeta 送审门槛", "至少同时满足：一个真实匹配多组学案例；独立 test set 的 claim-level 结果；公平 baseline；证据支持关系评价；公开 repo/DOI/环境；所有 P0 清零；当前人类评价不再被夸大；论文不把 transport fidelity 当模型推理能力。", "FFF2CC")

    heading(doc, "十三、80 分投稿准备度评分表", 1)
    para(doc, "这是团队内部的 go/no-go 评分，不是录用概率。若总分未达到 80，或者任一 P0 未完成，不建议投 iMeta。达到 80 也只能表示材料足够成熟，可以进入编辑筛选和同行评审。")
    add_table(doc, ["维度", "权重", "80 分门槛", "当前估计", "补齐动作"], [
        ("iMeta/组学主题匹配", "15", "真实跨模态主案例和明确科学问题", "6", "MultiAssay + matched omics case"),
        ("独立科学证据", "25", "冻结 test set；claim/evidence/diagnosis 指标和 CI", "8", "100–200 claims、专家标注、fair baselines"),
        ("方法创新与可解释边界", "15", "typed claims、证据角色、冲突和不确定性可验证", "10", "claim ledger + support policy"),
        ("软件和复现", "15", "公开 repo、DOI、CI、container、R CMD check、artifact manifest", "10", "发布和第三方复现"),
        ("人类评价", "10", "伦理/招募/任务正确率/盲法或正式专家评价", "2", "预注册和真实任务研究"),
        ("论文主张一致性", "10", "无历史数字、无过度 claim、主次终点分开", "5", "重写主稿和审计器"),
        ("数据与伦理/作者信息", "10", "作者、基金、COI、伦理、数据管理全部真实填写", "3", "投稿前核验并补齐"),
        ("合计", "100", "≥80 且 P0=0", "44", "预计 12–14 周集中改造"),
    ], widths=[4.0, 1.5, 6.0, 2.0, 4.0], font_size=8.0)

    heading(doc, "十四、最终投稿前验收清单", 1)
    check_groups = [
        ("软件", [
            "R CMD check --as-cran 无 ERROR/WARNING；testthat 和 optional-dependency matrix 有记录",
            "精度、隐私、claim schema、证据角色、MultiAssay 对象、图形和修复验证测试通过",
            "所有 exported functions 的 Rd、README、用户指南和示例一致",
            "没有把 callr 或 Shiny 描述为 OS sandbox；网络和文件权限边界写清楚",
        ]),
        ("研究", [
            "开发案例、调参案例和最终 test set 有时间戳和 hash；最终 test set 未参与开发",
            "transport fidelity、statistical interpretation、evidence support、user benefit 分开计分",
            "每个主结论有 reference facts、原始响应、评分记录、脚本和 manifest",
            "模型版本或可复现实验日期、prompt、token、temperature、cost 和失败重试全部记录",
        ]),
        ("多组学", [
            "样本交集、缺失模态、批次、donor、group 和 timepoint 都有数据表",
            "至少一个真实匹配的 SpatialExperiment/Seurat spatial 案例可在干净环境重建",
            "空间统计记录 spot/区域分母、邻接定义、切片/样本层级和空间 QC",
            "若接入 docking，结构文件、引擎、参数、pose、decoy 和 negative control 全部可追溯",
            "跨模态一致方向、空间共定位和 docking score 都不被写成因果或实验结合；派生证据和独立证据明确分开",
        ]),
        ("人类与伦理", [
            "真实招募、参与者背景、知情同意、补偿和伦理审批/豁免已核实",
            "专家评分保留匿名原始数据，报告适合的可靠性指标和裁决流程",
            "用户任务有对照、正确率、时间和预注册的主终点；没有就不写用户收益",
        ]),
        ("投稿材料", [
            "作者、单位、通讯邮箱、CRediT、基金、利益冲突和 AI 使用声明真实完成",
            "GitHub release、Zenodo DOI、LICENSE、container/renv.lock、CI 和复现说明公开",
            "主文、补充材料、图表、原始 records、代码和哈希一致；审计器使用当前冻结目录",
            "iMeta scope 和当年分区核验完成；同时准备 Bioinformatics/NAR G&B/PLOS Comp Bio 路线版本",
        ]),
    ]
    for group, items in check_groups:
        heading(doc, group, 2)
        add_bullets(doc, ["□ " + item for item in items])

    heading(doc, "十五、明确不要做的事", 1)
    add_bullets(doc, [
        "不要把 facts_only/facts_evidence 的 100% 写成模型具有 100% 生物信息学解释准确率。",
        "不要把程序生成的 120 个案例写成 120 个独立生物学样本，也不要把三轮重复当成独立样本推断。",
        "不要把当前 3 位评分者和 30 位功能评分的聚合结果写成正式用户实验、学习增益或临床效用。",
        "不要把数据库 term definition、Europe PMC 搜索命中或 UniProt annotation 自动写成实验通路激活/机制证据。",
        "不要把 iSEE、GPTCelltype、ellmer 的入口测试写成相近工具性能排名。",
        "不要在没有真实空间/多组学主案例前使用 integrated omics 作为论文核心标题。",
        "不要把空间共定位、ligand-receptor 表达或 docking score 单独写成真实分子结合、通路激活或治疗机制。",
        "不要把不同 docking 引擎的 score 放在同一尺度直接排名，也不要忽略 decoy、构象和口袋定义的敏感性。",
        "不要继续用功能数量替代科学问题和独立终点。",
        "不要上传 token、API key、原始患者/样本标识或未脱敏评分文件。",
    ])

    heading(doc, "附录 A：74 个导出函数的当前功能地图", 1)
    appendix_rows = [
        ("模型配置与调用（13）", "bio_config; bio_configure; bio_get_config; bio_reset_config; bio_presets; bio_register_preset; bio_unregister_preset; bio_endpoint_url; bio_llm_available; bio_build_request; bio_llm_chat; bio_parse_response; bio_redact"),
        ("对象解析与适配器（6）", "bio_context; new_bio_context; bio_context_payload; bio_adapters; bio_register_adapter; bio_unregister_adapter"),
        ("解释与声明（2）", "bio_explain; bio_validate_explanation"),
        ("证据与来源（13）", "bio_evidence; bio_evidence_register; bio_evidence_list; bio_evidence_clear; bio_evidence_lookup; bio_register_evidence_provider; bio_fetch_evidence; bio_evidence_for_context; bio_evidence_providers; bio_register_live_providers; bio_evidence_cache_clear; bio_go_annotations; bio_dependency_provenance"),
        ("图形与组成（7）", "bio_plot; bio_plot_pca; bio_plot_umap; bio_plot_enrichment; bio_plot_heatmap; bio_cross_check; bio_cell_proportions"),
        ("记录、诊断与修复（16）", "bio_record; bio_record_code; bio_records; bio_value; bio_last_error; bio_explain_error; bio_suggest_fix; bio_validate_fix; bio_invariant; inv_class; inv_equal; inv_factor_levels; inv_has_columns; inv_no_na; inv_nrow; inv_same_nrow_as_input"),
        ("运行环境、项目与报告（6）", "bio_runtime_info; bio_object_inventory; bio_package_usage; bio_project; bio_explain_project; bio_report"),
        ("教学与评估（10）", "bio_question_bank; bio_quiz; bio_grade; bio_review_rubric; bio_benchmark; bio_benchmark_fixes; bio_evaluate_experts; bio_evaluate_learning; bio_plan_learning_study; bio_randomize_learning_study"),
        ("工作台（1）", "bio_run_app"),
    ]
    add_table(doc, ["功能组", "函数"], appendix_rows, widths=[4.3, 13.0], font_size=8.0)

    heading(doc, "附录 B：建议的第一轮新目录", 1)
    add_code(doc, "research/\n  frozen_benchmark/\n    cases/                 # 开发集和 test 集分离\n    claim_schema.json\n    evidence_labels.csv\n    expert_adjudication.csv\n    benchmark_manifest.json\n  multiomics/\n    download_manifest.csv\n    sample_map.csv\n    run_multiomics_case.R\n  run_claim_benchmark.R\n  run_evidence_support_benchmark.R\n  run_error_diagnosis_benchmark.R\n  run_plot_audit_benchmark.R\n  audit_current_results.R\n  results/\n    frozen_v1/\n      prompts/ responses/ scores/ manifests/\ndocs/\n  PAPER_DRAFT.md\n  SUPPLEMENTARY_MATERIAL.md\n  DATA_DICTIONARY.md\n  ETHICS_AND_HUMAN_EVALUATION.md\n  REPRODUCTION.md")

    heading(doc, "结论", 1)
    para(doc, "BioExplainR 现在已经具备一篇扎实软件方法学论文的骨架，真正的短板不是代码量，而是主张与证据不在同一个层级：软件提供了声明和证据的检查接口，但当前研究主要验证接口机制；论文想进入 iMeta 或其他一区，需要把这些接口用于独立的真实多组学科学问题，并量化它们是否减少统计误读、无支持生物学表述和证据错配。")
    para(doc, "最推荐的路线是：先把 0.3.1 冻结为 reproducibility baseline，修复历史审计和 JSON 精度问题；然后完成 typed claim ledger、证据支持关系和一个真实匹配多组学案例；接着运行开发集完全隔离的 claim-level benchmark，并在条件允许时完成伦理完整的用户任务研究。只有当 P0 清零、内部投稿准备度达到 80/100、主文每个数字和主张都能回到独立记录时，再把 iMeta 作为首投，并准备 NAR Genomics and Bioinformatics、Bioinformatics 或 PLOS Computational Biology 的针对性版本。")

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
