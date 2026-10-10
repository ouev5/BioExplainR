"""Build manuscript figures, editable tables and an auditable record bundle."""
import csv
import hashlib
import json
import math
from pathlib import Path
import statistics
import zipfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "research/results"
MAIN_RESULTS = RESULTS / "main_scale_480"
OUT = ROOT / "docs/submission_assets"
METHODS = ["direct_bounded", "facts_only", "facts_evidence"]
LABELS = ["有界原始输出", "R 事实摘要", "R 事实加证据"]
FIGURE_LABELS = ["Bounded raw output", "R facts only", "R facts + evidence"]
COLORS = ["#BD5147", "#247A6B", "#3376A3"]
FAMILIES = ["cell_composition", "differential_expression", "enrichment", "generic", "pca", "single_cell"]
FAMILY_LABELS = ["Cell composition", "Differential expression", "Enrichment", "Generic summary", "PCA", "Single-cell"]


def read_csv(path):
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def normalize_cell(value):
    return "" if value is None else value


def save_figure(fig, name):
    fig.canvas.draw()
    for ax in fig.axes:
        if not ax.axison:
            ax.set_xticks([])
            ax.set_yticks([])
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    unrendered_ticks = set()
    for ax in fig.axes:
        for axis in (ax.xaxis, ax.yaxis):
            low, high = sorted(axis.get_view_interval())
            for tick in axis.get_major_ticks() + axis.get_minor_ticks():
                if not low <= tick.get_loc() <= high:
                    unrendered_ticks.update([tick.label1, tick.label2])
    text_objects = []
    for text in fig.findobj(matplotlib.text.Text):
        if text not in unrendered_ticks and text.get_visible() and text.get_text():
            bounds = text.get_window_extent(renderer)
            assert fig.bbox.contains(bounds.x0, bounds.y0) and fig.bbox.contains(bounds.x1, bounds.y1), (name, text.get_text())
            text_objects.append((text, bounds))
    overlaps = []
    for i, (left, a) in enumerate(text_objects):
        for right, b in text_objects[i + 1:]:
            if left.axes is not right.axes:
                continue
            width = max(0, min(a.x1, b.x1) - max(a.x0, b.x0))
            height = max(0, min(a.y1, b.y1) - max(a.y0, b.y0))
            intersection = width * height
            smaller = min(a.width * a.height, b.width * b.height)
            if smaller and intersection / smaller > .025:
                overlaps.append((left.get_text(), right.get_text(), round(intersection, 2)))
    assert not overlaps, (name, "text overlaps", overlaps[:8])
    fig.savefig(OUT / f"{name}.pdf", facecolor="white")
    fig.savefig(OUT / f"{name}.png", dpi=300, facecolor="white")
    fig.savefig(OUT / f"{name}.svg", facecolor="white")
    plt.close(fig)


def save_panel(fig, name):
    """Save one panel as an independently reusable vector figure."""
    save_figure(fig, name)


def panel_figure(name, draw, figsize=(7.6, 5.5)):
    fig, ax = plt.subplots(figsize=figsize)
    draw(fig, ax)
    save_panel(fig, name)


def finish_axis(ax):
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(color="#E5E5E5", linewidth=.5)
    ax.set_axisbelow(True)


def panel_title(ax, label, title, subtitle=None):
    ax.text(.01, .99, label, transform=ax.transAxes, ha="left", va="top",
            fontsize=14, weight="bold", color="#222222")
    ax.text(.075, .985, title, transform=ax.transAxes, ha="left", va="top",
            fontsize=11.5, weight="bold", color="#222222")
    if subtitle:
        ax.text(.075, .89, subtitle, transform=ax.transAxes, ha="left", va="top",
                fontsize=8.5, color="#5A6266")


def architecture():
    def draw_pipeline(fig, ax):
        ax.axis("off")
        panel_title(ax, "A", "Evidence-constrained workflow", "R computes facts; optional models organize bounded language")
        stages = [
            ("1", "R objects", "DESeq2 / limma / edgeR\nSCE / Seurat / ggplot", "#247A6B"),
            ("2", "Context", "Facts + provenance\nbounded preview + scope", "#247A6B"),
            ("3", "Language", "Local templates\nor optional LLM", "#3376A3"),
            ("4", "Checks", "Claims / plots / fixes\npass, fail, partial", "#6C535E"),
            ("5", "Outputs", "Reports + Shiny\nsource states retained", "#454545"),
        ]
        ys = [.70, .56, .42, .28, .14]
        for (number, title, detail, color), y in zip(stages, ys):
            ax.add_patch(FancyBboxPatch((.08, y), .84, .108, boxstyle="round,pad=.012,rounding_size=.015",
                                        transform=ax.transAxes, facecolor="white", edgecolor=color, linewidth=1.4))
            ax.text(.12, y + .048, number, transform=ax.transAxes, fontsize=10, weight="bold", color=color, va="center")
            ax.text(.19, y + .084, title, transform=ax.transAxes, fontsize=9.8, weight="bold", color=color, va="center")
            ax.text(.19, y + .014, detail, transform=ax.transAxes, fontsize=7.2, color="#252525", va="center")
            if y > .18:
                ax.add_patch(FancyArrowPatch((.5, y), (.5, y - .026), transform=ax.transAxes,
                                             arrowstyle="-|>", mutation_scale=10, linewidth=1, color="#687176"))
        ax.text(.08, .055, "Full matrices stay in R; only bounded summaries cross the model interface.",
                transform=ax.transAxes, fontsize=7.4, color="#5A6266")

    def draw_contract(fig, ax):
        ax.axis("off")
        panel_title(ax, "B", "Data contract", "What is computed, what is sent, and what is never delegated")
        boxes = [
            (.05, .47, .26, .27, "Full input", "counts / result table\nmetadata / plot source", "#247A6B"),
            (.37, .47, .26, .27, "R-computed\ncontext", "scalar facts\nsource + uncertainty", "#247A6B"),
            (.69, .47, .26, .27, "Model\nrequest", "bounded JSON\nwording only", "#3376A3"),
        ]
        for x, y, w, h, title, detail, color in boxes:
            ax.add_patch(FancyBboxPatch((x, y), w, h, transform=ax.transAxes,
                                        boxstyle="round,pad=.015,rounding_size=.02", facecolor="white",
                                        edgecolor=color, linewidth=1.4))
            ax.text(x + .02, y + h - .045, title, transform=ax.transAxes, fontsize=9.2, weight="bold", color=color,
                    va="top", linespacing=1.0)
            ax.text(x + .02, y + .075, detail, transform=ax.transAxes, fontsize=7.6, color="#252525", va="bottom",
                    linespacing=1.0)
        for x in (.31, .63):
            ax.add_patch(FancyArrowPatch((x, .605), (x + .055, .605), transform=ax.transAxes,
                                         arrowstyle="-|>", mutation_scale=11, linewidth=1.2, color="#4D575B"))
        ax.add_patch(FancyBboxPatch((.13, .17), .74, .14, transform=ax.transAxes,
                                    boxstyle="round,pad=.015,rounding_size=.02", facecolor="#FFF7F5",
                                    edgecolor=COLORS[0], linewidth=1.2))
        ax.text(.5, .24, "Not delegated: full-table counting, statistical testing, raw-matrix access",
                transform=ax.transAxes, ha="center", va="center", fontsize=9, weight="bold", color="#8D3B34")
        ax.text(.5, .08, "Evidence retrieval is explicit opt-in and does not establish claim support.",
                transform=ax.transAxes, ha="center", va="center", fontsize=8.4, color="#5A6266")

    def draw_validation(fig, ax):
        ax.axis("off")
        panel_title(ax, "C", "Validation states", "Each check reports its scope; no global reliability score is implied")
        states = [("pass", "specified value / mapping agrees", "#247A6B"),
                  ("fail", "specified value / mapping disagrees", "#BD5147"),
                  ("partial", "some checks available; value absent", "#C18A2B"),
                  ("unverified", "no explicit machine-checkable claim", "#6C7377")]
        for i, (status, detail, color) in enumerate(states):
            y = .73 - i * .14
            ax.add_patch(FancyBboxPatch((.12, y), .22, .085, transform=ax.transAxes,
                                        boxstyle="round,pad=.012,rounding_size=.02", facecolor=color,
                                        edgecolor=color))
            ax.text(.23, y + .042, status, transform=ax.transAxes, ha="center", va="center",
                    fontsize=9.5, weight="bold", color="white")
            ax.text(.39, y + .042, detail, transform=ax.transAxes, va="center", fontsize=8.8, color="#252525")
        ax.add_patch(FancyBboxPatch((.12, .11), .76, .09, transform=ax.transAxes,
                                    boxstyle="round,pad=.012,rounding_size=.02", facecolor="#F4F7F6",
                                    edgecolor="#8B9699"))
        ax.text(.5, .155, "Facts, provenance, model wording and check status stay separate in reports.",
                transform=ax.transAxes, ha="center", va="center", fontsize=8.6, color="#30383B")

    def draw_session(fig, ax):
        ax.axis("off")
        panel_title(ax, "D", "Local workbench", "A reproducible session boundary for code, objects, explanations and reports")
        tabs = [("Settings", "model / language / limits"), ("Runtime", "packages / objects / records"),
                ("Analysis", "facts / wording / claims"), ("Plots", "PCA / UMAP / enrichment")]
        for i, (title, detail) in enumerate(tabs):
            x = .08 + (i % 2) * .45
            y = .62 - (i // 2) * .25
            ax.add_patch(FancyBboxPatch((x, y), .37, .14, transform=ax.transAxes,
                                        boxstyle="round,pad=.012,rounding_size=.02", facecolor="white",
                                        edgecolor="#247A6B", linewidth=1.2))
            ax.text(x + .02, y + .092, title, transform=ax.transAxes, fontsize=9.8, weight="bold", color="#247A6B")
            ax.text(x + .02, y + .042, detail, transform=ax.transAxes, fontsize=7.7, color="#30383B")
        ax.add_patch(FancyBboxPatch((.12, .12), .76, .11, transform=ax.transAxes,
                                    boxstyle="round,pad=.012,rounding_size=.02", facecolor="#F6FAF9",
                                    edgecolor="#247A6B"))
        ax.text(.5, .175, "Per-browser session: state is retained locally; model never executes code.",
                transform=ax.transAxes, ha="center", va="center", fontsize=8.4, color="#30383B")

    draws = [draw_pipeline, draw_contract, draw_validation, draw_session]
    for label, draw in zip(("A", "B", "C", "D"), draws):
        def wrapped(fig, ax, draw=draw):
            draw(fig, ax)
        panel_figure(f"Figure_1{label}_framework", wrapped, figsize=(6.0, 4.4))
    fig, axes = plt.subplots(2, 2, figsize=(12.0, 8.6))
    for draw, ax in zip(draws, axes.flat):
        draw(fig, ax)
    fig.subplots_adjust(left=.025, right=.975, top=.96, bottom=.035, wspace=.05, hspace=.08)
    save_figure(fig, "Figure_1_architecture")


def ablation(scores):
    rounds = []
    summary = {}
    for index, method in enumerate(METHODS):
        counts = [sum(row["correct"] == "TRUE" for row in scores
                      if row["method"] == method and int(row["replicate"]) == r) for r in (1, 2, 3)]
        rates = [100*n/480 for n in counts]
        mean, sd = statistics.mean(rates), statistics.stdev(rates)
        summary[method] = dict(counts=counts, rates=rates, mean=mean, sd=sd)
        for replicate, count in enumerate(counts, 1):
            rounds.append(dict(method=method, replicate=replicate, correct=count, target=480,
                               matching_rate=count/480, mean_rate=mean/100, sample_sd=sd/100))

    def draw_accuracy(fig, ax):
        for index, method in enumerate(METHODS):
            item = summary[method]
            ax.bar(index, item["mean"], width=.53, color=COLORS[index], alpha=.9)
            ax.errorbar(index, item["mean"], yerr=item["sd"], capsize=6, color="#151515",
                        fmt="none", linewidth=1.3, zorder=4)
            ax.scatter([index-.12, index, index+.12], item["rates"], marker="o", s=40,
                       facecolor="white", edgecolor="#151515", linewidth=1, zorder=5)
            ax.text(index, item["mean"] + item["sd"] + 6.5,
                     f'{sum(item["counts"])}/1440\n({item["mean"]:.1f}%)',
                     ha="center", fontsize=8.5, weight="bold", linespacing=1.0)
        ax.set_xticks(range(3), ["Bounded raw\noutput", "R facts\nonly", "R facts +\nevidence"])
        ax.set_ylabel("Requested-fact matching rate (%)")
        ax.set_ylim(0, 116)
        ax.set_yticks(range(0, 101, 20))
        ax.spines[["top", "right"]].set_visible(False)
        ax.grid(axis="y", color="#E2E2E2", linewidth=.6)
        ax.set_axisbelow(True)
        ax.set_title("Overall matching", loc="left", weight="bold")

    def draw_rounds(fig, ax):
        x = [1, 2, 3]
        for index, method in enumerate(METHODS):
            item = summary[method]
            ax.plot(x, item["rates"], marker="o", linewidth=2, markersize=6,
                    color=COLORS[index], label=FIGURE_LABELS[index])
            for replicate, rate, count in zip(x, item["rates"], item["counts"]):
                x_offset = (-20, 0, 20)[index]
                y_offset = (6, 14, 22)[index]
                ax.annotate(f'{count}/480', (replicate, rate), textcoords="offset points",
                            xytext=(x_offset, y_offset), ha="center", fontsize=6.8, color=COLORS[index])
        ax.set_xticks(x, ["Round 1", "Round 2", "Round 3"])
        ax.set_ylim(0, 122)
        ax.set_yticks(range(0, 101, 20))
        ax.set_ylabel("Matching rate (%)")
        ax.set_title("Repeated rounds", loc="left", weight="bold", pad=2)
        ax.legend(frameon=False, fontsize=7.2, loc="lower right")
        ax.grid(axis="y", color="#E2E2E2", linewidth=.6)
        ax.set_axisbelow(True)
        ax.spines[["top", "right"]].set_visible(False)

    family_rows = read_csv(MAIN_RESULTS / "family_summary.csv")
    call_rows = read_csv(MAIN_RESULTS / "model_calls.csv")

    def draw_case_heatmap(fig, ax):
        fig.subplots_adjust(left=.34, right=.86, bottom=.28, top=.88)
        ordered_families = FAMILIES
        values = []
        labels = []
        for family in ordered_families:
            row_values = []
            row_labels = []
            for method in METHODS:
                row = next(r for r in family_rows if r["family"] == family and r["method"] == method)
                row_values.append(float(row["accuracy"]) * 100)
                row_labels.append(f'{row["correct"]}/{row["expected"]}')
            values.append(row_values)
            labels.append(row_labels)
        image = ax.imshow(values, cmap="YlGn", vmin=0, vmax=100, aspect="auto")
        for i, row in enumerate(labels):
            for j, text in enumerate(row):
                ax.text(j, i, text, ha="center", va="center", fontsize=8.5,
                        color="#102820" if values[i][j] > 50 else "#252525", weight="bold")
        ax.set_xticks(range(3), ["Bounded raw\noutput", "R facts\nonly", "R facts +\nevidence"])
        ax.tick_params(axis="x", labelsize=7.2)
        ax.set_yticks(range(len(ordered_families)), FAMILY_LABELS)
        ax.set_title("Functional-family matching", loc="left", weight="bold")
        ax.set_xlabel("Correct / target" if ax.get_position().height > .35 else "")
        ax.set_ylabel("")
        fig.colorbar(image, ax=ax, fraction=.046, pad=.04, label="Matching rate (%)")

    def draw_latency(fig, ax):
        medians = []
        for method in METHODS:
            elapsed = [float(r["elapsed_sec"]) for r in call_rows if r["method"] == method]
            medians.append(statistics.median(elapsed))
        bars = ax.bar(range(3), medians, color=COLORS, width=.56)
        for bar, value in zip(bars, medians):
            ax.text(bar.get_x() + bar.get_width()/2, value + max(medians)*.04, f"{value:.2f}",
                    ha="center", fontsize=9, weight="bold")
        ax.set_xticks(range(3), ["Bounded raw", "R facts", "R facts +\nevidence"])
        ax.set_ylabel("Median response time (s)")
        ax.set_title("Request cost in time", loc="left", weight="bold")
        finish_axis(ax)

    panel_figure("Figure_2A_overall", draw_accuracy, figsize=(5.4, 4.8))
    panel_figure("Figure_2B_rounds", draw_rounds, figsize=(5.4, 4.8))
    panel_figure("Figure_2C_case_heatmap", draw_case_heatmap, figsize=(6.4, 4.8))
    panel_figure("Figure_2D_latency", draw_latency, figsize=(5.4, 4.8))
    fig, axes = plt.subplots(2, 2, figsize=(12.0, 8.7), gridspec_kw={"height_ratios": [1, 1.05]})
    draw_accuracy(fig, axes[0, 0])
    draw_rounds(fig, axes[0, 1])
    draw_case_heatmap(fig, axes[1, 0])
    draw_latency(fig, axes[1, 1])
    fig.subplots_adjust(left=.25, right=.96, bottom=.1, top=.96, wspace=.36, hspace=.42)
    save_figure(fig, "Figure_2_ablation")
    with (OUT / "Figure_2_source.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=rounds[0].keys())
        writer.writeheader()
        writer.writerows(rounds)
    with (OUT / "Figure_2_family_source.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=family_rows[0].keys())
        writer.writeheader()
        writer.writerows(family_rows)
    return rounds


def biological_figures():
    source = read_csv(OUT / "Figure_3_source.csv")
    donor_codes = {donor: f"D{i + 1:02d}" for i, donor in enumerate(sorted({row["donor"] for row in source}))}
    donors = list(donor_codes.values())
    for row in source:
        row["donor_code"] = donor_codes[row["donor"]]
    donor_rows = {donor: {row["treatment"]: row for row in source if row["donor_code"] == donor} for donor in donors}

    def draw_pca(fig, ax):
        panel_title(ax, "A", "Sample scores", "Paired airway libraries; colour = treatment, shape = donor")
        for row in source:
            ax.scatter(float(row["PC1"]), float(row["PC2"]), s=65,
                       marker=["o", "s", "^", "D"][donors.index(row["donor_code"])],
                       facecolor=COLORS[1] if row["treatment"] == "untrt" else COLORS[0],
                       edgecolor="#202020", linewidth=.6)
        handles = [plt.Line2D([], [], color=color, marker="o", linestyle="none", label=label, markersize=7)
                   for color, label in [(COLORS[1], "untrt"), (COLORS[0], "trt")]]
        ax.legend(handles=handles, frameon=False, loc="lower right", fontsize=7)
        ax.set_xlabel("PC1 (47.8429%)")
        ax.set_ylabel("PC2 (23.3834%)")
        finish_axis(ax)

    def draw_pair_shift(fig, ax):
        panel_title(ax, "B", "Within-donor shifts", "Each arrow connects untrt to trt")
        for i, donor in enumerate(donors):
            before = donor_rows[donor]["untrt"]
            after = donor_rows[donor]["trt"]
            x0, y0 = float(before["PC1"]), float(before["PC2"])
            x1, y1 = float(after["PC1"]), float(after["PC2"])
            ax.annotate("", xy=(x1, y1), xytext=(x0, y0),
                        arrowprops={"arrowstyle": "-|>", "color": COLORS[0], "lw": 1.2, "alpha": .75})
            ax.scatter([x0], [y0], color=COLORS[1], s=45, edgecolor="#202020", zorder=3)
            ax.scatter([x1], [y1], color=COLORS[0], s=45, edgecolor="#202020", zorder=3)
            ax.text(x0, y0, donor, fontsize=7, ha="right", va="bottom")
        ax.set_xlabel("PC1")
        ax.set_ylabel("PC2")
        finish_axis(ax)

    def draw_variance(fig, ax):
        panel_title(ax, "C", "Variance explained", "PCA display axes retained most of the measured variation")
        values = [47.8429, 23.3834]
        bars = ax.bar(["PC1", "PC2"], values, color=[COLORS[1], COLORS[2]], width=.55)
        for bar, value in zip(bars, values):
            ax.text(bar.get_x() + bar.get_width()/2, value + 1.5, f"{value:.1f}%", ha="center", fontsize=9, weight="bold")
        ax.set_ylim(0, 60)
        ax.set_ylabel("Variance explained (%)")
        finish_axis(ax)

    def draw_design(fig, ax):
        ax.axis("off")
        panel_title(ax, "D", "Analysis design", "Observed compatibility example, not a treatment efficacy test")
        facts = [("8", "libraries"), ("4", "donors"), ("2", "treatment levels"), ("~ cell + dex", "paired design")]
        for i, (value, label) in enumerate(facts):
            x = .08 + (i % 2) * .48
            y = .63 - (i // 2) * .25
            ax.add_patch(FancyBboxPatch((x, y), .38, .14, transform=ax.transAxes,
                                        boxstyle="round,pad=.012,rounding_size=.02", facecolor="white",
                                        edgecolor="#247A6B", linewidth=1.2))
            ax.text(x + .19, y + .085, value, transform=ax.transAxes, ha="center", fontsize=12, weight="bold", color="#247A6B")
            ax.text(x + .19, y + .035, label, transform=ax.transAxes, ha="center", fontsize=8, color="#30383B")
        ax.text(.08, .16, "No inference from visual separation is claimed.", transform=ax.transAxes, fontsize=8.5, color="#8D3B34")

    for name, draw in [("Figure_3A_scores", draw_pca), ("Figure_3B_pair_shift", draw_pair_shift),
                       ("Figure_3C_variance", draw_variance), ("Figure_3D_design", draw_design)]:
        panel_figure(name, draw, figsize=(5.6, 4.5))
    fig, axes = plt.subplots(2, 2, figsize=(11.8, 8.8))
    for draw, ax in zip((draw_pca, draw_pair_shift, draw_variance, draw_design), axes.flat):
        draw(fig, ax)
    fig.subplots_adjust(left=.12, right=.96, bottom=.09, top=.96, wspace=.25, hspace=.27)
    save_figure(fig, "Figure_3_airway_PCA")

    source = read_csv(OUT / "Figure_4_source.csv")
    comp = read_json(RESULTS / "compatibility.json")["SingleCellExperiment"]["facts"]

    def draw_umap(fig, ax):
        panel_title(ax, "A", "Cell embedding", "Colour indicates mitochondrial QC grouping, not cell type")
        for group, color in [("mt <= 5%", COLORS[1]), ("mt > 5%", COLORS[0])]:
            selected = [row for row in source if row["mito_qc"] == group]
            ax.scatter([float(row["UMAP1"]) for row in selected], [float(row["UMAP2"]) for row in selected],
                       s=5, color=color, alpha=.8, linewidth=0, label=group)
        ax.legend(frameon=False, markerscale=3, fontsize=7, loc="lower right")
        ax.set_xlabel("UMAP 1")
        ax.set_ylabel("UMAP 2")
        finish_axis(ax)

    def draw_umap_qc(fig, ax):
        fig.subplots_adjust(left=.2, right=.96, bottom=.25, top=.88)
        panel_title(ax, "B", "QC grouping", "The threshold is descriptive and did not exclude cells")
        counts = {group: sum(row["mito_qc"] == group for row in source) for group in ("mt <= 5%", "mt > 5%")}
        bars = ax.bar(["mt <= 5%", "mt > 5%"], list(counts.values()), color=[COLORS[1], COLORS[0]], width=.55)
        for bar, value in zip(bars, counts.values()):
            label_y = value - max(15, value * .08)
            ax.text(bar.get_x() + bar.get_width()/2, label_y, f"{value:,}", ha="center", va="top",
                    fontsize=8.5, weight="bold", color="white")
        ax.set_ylabel("Cells")
        ax.tick_params(axis="x", rotation=12)
        ax.set_ylim(0, max(counts.values()) * 1.16)
        finish_axis(ax)

    def draw_qc_intervals(fig, ax):
        fig.subplots_adjust(left=.27, right=.96, bottom=.2, top=.88)
        panel_title(ax, "C", "QC distributions", "Quantiles from the saved SCE summary")
        metrics = [
            ("nCount_RNA", comp["qc"]["n_counts"]["quantiles"], "UMIs"),
            ("nFeature_RNA", comp["qc"]["n_features"]["quantiles"], "genes"),
            ("percent.mt", comp["qc"]["pct_mt"]["quantiles"], "%"),
        ]
        for i, (label, q, unit) in enumerate(metrics):
            y = len(metrics) - 1 - i
            ax.hlines(y, q["min"], q["max"], color="#9AA6A8", linewidth=2)
            ax.plot([q["q25"], q["q75"]], [y, y], color=COLORS[2], linewidth=8, solid_capstyle="butt")
            ax.scatter([q["median"]], [y], s=45, color=COLORS[0], edgecolor="#202020", zorder=3)
            ax.text(q["max"], y + .13, f"median {q['median']:g} {unit}", fontsize=7.5, ha="right")
        ax.set_yticks(range(2, -1, -1), ["UMIs", "Genes", "Mito %"])
        ax.set_xlabel("Observed range; thick line = IQR")
        ax.set_ylim(-.6, 2.6)
        finish_axis(ax)

    def draw_scope(fig, ax):
        ax.axis("off")
        panel_title(ax, "D", "Scope of the example", "Compatibility and summarisation, not cell-type validation")
        cards = [("2,700", "cells"), ("32,738", "features"), ("1", "donor"), ("0", "cell-type labels")]
        for i, (value, label) in enumerate(cards):
            x = .08 + (i % 2) * .48
            y = .62 - (i // 2) * .25
            ax.add_patch(FancyBboxPatch((x, y), .37, .14, transform=ax.transAxes,
                                        boxstyle="round,pad=.012,rounding_size=.02", facecolor="white",
                                        edgecolor="#3376A3" if label != "cell-type labels" else "#BD5147", linewidth=1.2))
            ax.text(x + .185, y + .084, value, transform=ax.transAxes, ha="center", fontsize=12, weight="bold", color="#3376A3" if label != "cell-type labels" else "#BD5147")
            ax.text(x + .185, y + .035, label, transform=ax.transAxes, ha="center", fontsize=8, color="#30383B")
        ax.text(.08, .16, "UMAP is an exploratory embedding; cells are not independent donors.", transform=ax.transAxes, fontsize=8.1, color="#8D3B34")

    for name, draw in [("Figure_4A_embedding", draw_umap), ("Figure_4B_qc_counts", draw_umap_qc),
                       ("Figure_4C_qc_intervals", draw_qc_intervals), ("Figure_4D_scope", draw_scope)]:
        panel_figure(name, draw, figsize=(5.6, 4.5))
    fig, axes = plt.subplots(2, 2, figsize=(11.8, 8.8))
    for draw, ax in zip((draw_umap, draw_umap_qc, draw_qc_intervals, draw_scope), axes.flat):
        draw(fig, ax)
    fig.subplots_adjust(left=.07, right=.96, bottom=.07, top=.96, wspace=.25, hspace=.27)
    save_figure(fig, "Figure_4_PBMC3k_UMAP")

    screenshots = [
        ("settings", plt.imread(RESULTS / "browser/desktop-settings.png")),
        ("plot", plt.imread(RESULTS / "browser/desktop-plot.png")),
        ("analysis", plt.imread(RESULTS / "browser/desktop-analysis.png")),
    ]
    def draw_screenshot(fig, ax, image, title):
        ax.imshow(image)
        ax.axis("off")
        panel_title(ax, title[0], title[1], title[2])

    draw_functions = []
    for label, image, title, subtitle in [("A", screenshots[0][1], "Settings", "Model and language controls"),
                                          ("B", screenshots[1][1], "Plot", "Generated PCA saved to session"),
                                          ("C", screenshots[2][1], "Analysis", "Facts, wording and validation states")]:
        draw_functions.append(lambda fig, ax, image=image, label=label, title=title, subtitle=subtitle: draw_screenshot(fig, ax, image, (label, title, subtitle)))
    def draw_workbench_steps(fig, ax):
        ax.axis("off")
        panel_title(ax, "D", "Workflow", "The automated path used for the screenshot record")
        steps = ["Load settings", "Run and record R code", "Explain object", "Generate PCA", "Cross-check source", "Download report"]
        for i, step in enumerate(steps):
            y = .79 - i * .105
            ax.add_patch(FancyBboxPatch((.1, y), .8, .065, transform=ax.transAxes,
                                        boxstyle="round,pad=.008,rounding_size=.015", linewidth=1,
                                        edgecolor="#247A6B", facecolor="#F6FAF9"))
            ax.text(.15, y + .032, f"{i+1}", transform=ax.transAxes, fontsize=8, weight="bold", color="#247A6B", va="center")
            ax.text(.22, y + .032, step, transform=ax.transAxes, fontsize=8.4, color="#252525", va="center")
        ax.text(.1, .08, "Workflow record; not a user-study usability score.", transform=ax.transAxes, fontsize=7.8, color="#5A6266")
    draw_functions.append(draw_workbench_steps)
    for label, draw in zip(("A", "B", "C", "D"), draw_functions):
        panel_figure(f"Figure_S1{label}_workbench", draw, figsize=(6.0, 4.5))
    fig, axes = plt.subplots(2, 2, figsize=(12.0, 8.8))
    for draw, ax in zip(draw_functions, axes.flat):
        draw(fig, ax)
    fig.subplots_adjust(left=.02, right=.98, bottom=.03, top=.97, wspace=.04, hspace=.08)
    save_figure(fig, "Figure_S1_workbench")

    # Independent pasilla validation figure.  The source CSV files are exported
    # by run_related_tool_tests.R and contain no sample identifiers.
    pca = read_csv(RESULTS / "related_tools/pasilla_pca.csv")
    volcano = read_csv(RESULTS / "related_tools/pasilla_volcano.csv")
    colors = {"untreated": COLORS[1], "treated": COLORS[0]}
    def draw_pasilla_pca(fig, ax):
        for condition in ("untreated", "treated"):
            rows = [row for row in pca if row["condition"] == condition]
            ax.scatter([float(row["PC1"]) for row in rows], [float(row["PC2"]) for row in rows],
                       s=58, color=colors[condition], edgecolor="#202020", linewidth=.5,
                       label=condition)
        ax.set_xlabel("PC1")
        ax.set_ylabel("PC2")
        ax.set_title("VST expression PCA", loc="left", weight="bold")
        ax.legend(frameon=False)
        finish_axis(ax)

    def draw_pasilla_volcano(fig, ax):
        padj = [float(row["padj"]) if row["padj"] not in ("", "NA") else 1.0 for row in volcano]
        lfc = [float(row["log2FC"]) if row["log2FC"] not in ("", "NA") else 0.0 for row in volcano]
        significant = [p < .05 and abs(fc) > 1 for p, fc in zip(padj, lfc)]
        ax.scatter(lfc, [-math.log10(max(p, 1e-300)) for p in padj], s=4,
                   color=[COLORS[0] if hit else "#B8C2C1" for hit in significant], alpha=.75,
                   linewidth=0)
        ax.axvline(-1, color="#555555", linestyle="--", linewidth=.7)
        ax.axvline(1, color="#555555", linestyle="--", linewidth=.7)
        ax.axhline(-math.log10(.05), color="#555555", linestyle="--", linewidth=.7)
        ax.set_xlabel("log2 fold change")
        ax.set_ylabel("−log10 adjusted P")
        ax.set_title("DESeq2 treated vs untreated", loc="left", weight="bold")
        finish_axis(ax)

    def draw_pasilla_counts(fig, ax):
        fig.subplots_adjust(left=.2, right=.98, bottom=.28, top=.88)
        panel_title(ax, "C", "Independent fact benchmark", "P_adj < 0.05 and |log2FC| > 1")
        ax.bar(["Features", "Significant", "Up", "Down"], [14599, 218, 108, 110],
               color=["#8AA6A0", COLORS[0], COLORS[1], "#6C535E"], width=.58)
        ax.set_ylabel("Count")
        ax.tick_params(axis="x", rotation=22)
        finish_axis(ax)

    fact_checks = read_csv(RESULTS / "related_tools/pasilla_fact_checks.csv")
    def draw_pasilla_thresholds(fig, ax):
        fig.subplots_adjust(left=.2, right=.96, bottom=.2, top=.88)
        panel_title(ax, "D", "Threshold sensitivity", "All 30 BioExplainR-to-source fact checks passed")
        for lfc, color in [("0", COLORS[2]), ("1", COLORS[0])]:
            rows = [r for r in fact_checks if r["lfc"] == lfc and r["fact"] == "n_significant"]
            rows = sorted(rows, key=lambda r: float(r["alpha"]))
            xs = [float(r["alpha"]) for r in rows]
            ys = [float(r["actual"]) for r in rows]
            ax.plot(xs, ys, marker="o", linewidth=2, color=color, label=f"|log2FC| = {lfc}")
        ax.set_xticks([.01, .05, .1], ["0.01", "0.05", "0.10"])
        ax.set_xlabel("Adjusted P threshold")
        ax.set_ylabel("Significant features")
        ax.legend(frameon=False, fontsize=8)
        finish_axis(ax)

    panel_figure("Figure_S2A_pasilla_PCA", draw_pasilla_pca, figsize=(5.2, 4.6))
    panel_figure("Figure_S2B_pasilla_volcano", draw_pasilla_volcano, figsize=(5.2, 4.6))
    panel_figure("Figure_S2C_pasilla_counts", draw_pasilla_counts, figsize=(5.2, 4.6))
    panel_figure("Figure_S2D_pasilla_thresholds", draw_pasilla_thresholds, figsize=(5.2, 4.6))
    fig, axes = plt.subplots(2, 2, figsize=(12.0, 8.8))
    for ax, draw in zip(axes.flat, [draw_pasilla_pca, draw_pasilla_volcano, draw_pasilla_counts, draw_pasilla_thresholds]):
        draw(fig, ax)
    fig.subplots_adjust(left=.07, right=.96, bottom=.08, top=.96, wspace=.27, hspace=.28)
    save_figure(fig, "Figure_S2_pasilla")

    expert = read_csv(RESULTS / "human_evaluation/expert_dimension_summary.csv")
    users = read_csv(RESULTS / "human_evaluation/user_feature_summary.csv")
    agreement = read_csv(RESULTS / "human_evaluation/expert_agreement_summary.csv")[0]

    def draw_expert_means(fig, ax):
        fig.subplots_adjust(left=.34, right=.96, bottom=.18, top=.88)
        panel_title(ax, "A", "External rating dimensions", "Three raters × 30 cases; score range 0–2")
        rows = [r for r in expert if r["dimension"] != "overall"]
        rows = sorted(rows, key=lambda r: float(r["mean"]))
        labels = [r["dimension"].replace("_", " ") for r in rows]
        means = [float(r["mean"]) for r in rows]
        ax.barh(range(len(rows)), means, color=COLORS[2], height=.62)
        ax.set_yticks(range(len(rows)), labels, fontsize=7.5)
        ax.set_xlim(0, 2.1)
        ax.set_xlabel("Mean rating")
        ax.axvline(1.0, color="#A4AEB0", linewidth=.8, linestyle="--")
        finish_axis(ax)

    def draw_agreement(fig, ax):
        panel_title(ax, "B", "Inter-rater agreement", "Case–dimension units; descriptive agreement only")
        labels = ["Exact", "Within 1 point"]
        values = [float(agreement["exact_agreement_percent"]), float(agreement["within_one_point_percent"])]
        bars = ax.bar(labels, values, color=[COLORS[0], COLORS[1]], width=.55)
        for bar, value in zip(bars, values):
            ax.text(bar.get_x() + bar.get_width()/2, value - 10, f"{value:.1f}%", ha="center", va="top",
                    fontsize=8.5, weight="bold", color="white")
        ax.set_ylim(0, 112)
        ax.set_ylabel("Agreement (%)")
        ax.tick_params(axis="x", rotation=12)
        finish_axis(ax)

    def draw_student_scores(fig, ax):
        fig.subplots_adjust(left=.34, right=.96, bottom=.18, top=.88)
        panel_title(ax, "C", "Student-rated feature scores", "30 participants; score range 0–100")
        rows = [r for r in users if r["feature_id"] != "overall"]
        rows = sorted(rows, key=lambda r: float(r["mean_score"]))
        labels = [r["feature"].replace("运行并记录 R 代码", "Run and record R code")
                  .replace("解释 R 生物信息学对象", "Explain R objects")
                  .replace("生成常用分析图形", "Generate plots")
                  .replace("证据与来源追踪", "Evidence trace")
                  .replace("生成分析报告", "Generate report") for r in rows]
        means = [float(r["mean_score"]) for r in rows]
        ax.barh(range(len(rows)), means, color=COLORS[1], height=.62)
        ax.set_yticks(range(len(rows)), labels, fontsize=7.5)
        ax.set_xlim(0, 105)
        ax.set_xlabel("Mean score")
        finish_axis(ax)

    def draw_student_useful(fig, ax):
        fig.subplots_adjust(left=.34, right=.96, bottom=.18, top=.88)
        panel_title(ax, "D", "Perceived usefulness", "1 = useful; 2 = not useful")
        rows = [r for r in users if r["feature_id"] != "overall"]
        rows = sorted(rows, key=lambda r: float(r["useful_percent"]))
        labels = [r["feature"].replace("运行并记录 R 代码", "Run and record R code")
                  .replace("解释 R 生物信息学对象", "Explain R objects")
                  .replace("生成常用分析图形", "Generate plots")
                  .replace("证据与来源追踪", "Evidence trace")
                  .replace("生成分析报告", "Generate report") for r in rows]
        values = [float(r["useful_percent"]) for r in rows]
        bars = ax.barh(range(len(rows)), values, color=COLORS[0], height=.62)
        for bar, value in zip(bars, values):
            ax.text(value + 1.2, bar.get_y() + bar.get_height()/2, f"{value:.1f}%", va="center", fontsize=8)
        ax.set_yticks(range(len(rows)), labels, fontsize=7.5)
        ax.set_xlim(0, 112)
        ax.set_xlabel("Useful responses (%)")
        finish_axis(ax)

    human_draws = [draw_expert_means, draw_agreement, draw_student_scores, draw_student_useful]
    for label, draw in zip(("A", "B", "C", "D"), human_draws):
        panel_figure(f"Figure_S3{label}_human_evaluation", draw, figsize=(5.8, 4.5))
    fig, axes = plt.subplots(2, 2, figsize=(12.0, 8.8))
    for draw, ax in zip(human_draws, axes.flat):
        draw(fig, ax)
    fig.subplots_adjust(left=.13, right=.96, bottom=.08, top=.96, wspace=.3, hspace=.3)
    save_figure(fig, "Figure_S3_human_evaluation")


def hydrate_tables():
    specs = read_json(ROOT / "docs/paper_tables.json")["tables"]
    supplementary_specs = hydrate_supplementary_tables(read_json(ROOT / "docs/supplementary_tables.json")["tables"])
    specs.extend(supplementary_specs)
    main = {r["method"]: r for r in read_csv(MAIN_RESULTS / "ablation_summary.csv")}
    families = {(r["family"], r["method"]): r for r in read_csv(MAIN_RESULTS / "family_summary.csv")}
    calls = read_csv(MAIN_RESULTS / "model_calls.csv")
    for spec in specs:
        if spec["id"] == "2":
            spec["rows"] = [[label, f'{r["correct"]}/{r["expected"]}（{100*float(r["accuracy"]):.1f}%）',
                             f'{100*float(r["coverage"]):.1f}%', f'{100*float(r["answered_accuracy"]):.1f}%',
                             f'{float(r["median_elapsed_sec"]):.3f}'] for label, m in zip(LABELS, METHODS) for r in [main[m]]]
        elif spec["id"] == "4":
            spec["rows"] = [[label] + [f'{families[(family, m)]["correct"]}/{families[(family, m)]["expected"]}' for m in METHODS]
                             for family, label in zip(FAMILIES, FAMILY_LABELS)]
        elif spec["id"] == "5":
            names = {"wrong_number": "错误数字", "unknown_path": "未知事实路径", "vector_value": "向量冒充标量",
                     "fabricated_citation": "未提供引用", "reference_only": "仅路径、无数值"}
            data = {r["mutation"]: r for r in read_csv(RESULTS / "validation_summary.csv")}
            spec["rows"] = [[names[key], int(data[key]["n"]), f'{data[key]["n"]} {data[key]["status"]}', "unchecked"] for key in names]
        elif spec["id"] == "6":
            names = {"volcano": "正确火山图", "reversed_direction": "火山图方向反转", "raw_p_as_adjusted": "原始 P 值替代校正值",
                     "pca": "正确 PCA", "pca_wrong_coordinate": "PCA 坐标改变", "pca_wrong_label": "PCA 方差轴标签改变",
                     "enrichment": "正确富集图", "enrichment_wrong_size": "富集点大小改变", "enrichment_wrong_colour": "富集点颜色改变"}
            spec["rows"] = [[names[r["case"]], r["expected"], r["actual"], "unchecked"] for r in read_csv(RESULTS / "plot_challenges.csv")]
            fixes = read_json(RESULTS / "fix_validation.json")
            spec["rows"] += [["保持 untrt 参考水平", "invariants_passed", fixes["correct_reference"]["status"], "unchecked"],
                             ["反转为 trt 参考水平", "invariants_failed", fixes["reversed_reference"]["status"], "unchecked"]]
        elif spec["id"] == "7":
            spec["rows"] = human_evaluation_rows()
        elif spec["id"] == "S1":
            spec["rows"] = scaled_cost_rows(calls)
        elif spec["id"] == "8":
            spec["rows"] = full_fact_rows()
    return specs


def scaled_cost_rows(calls):
    rows = []
    for label, method in zip(LABELS, METHODS):
        selected = [r for r in calls if r["method"] == method]
        prompt = sum(float(r["prompt_tokens"]) for r in selected)
        completion = sum(float(r["completion_tokens"]) for r in selected)
        citations = sum(int(r["citations"]) for r in selected)
        unknown = sum(int(r.get("unknown_citations", 0) or 0) for r in selected)
        cost = (prompt * .15 + completion * .6) / 1e6
        rows.append([label, citations, unknown, int(prompt), int(completion), f"{cost:.6f}"])
    return rows


def full_fact_rows():
    scores = read_csv(MAIN_RESULTS / "fact_scores.csv")
    grouped = {}
    for row in scores:
        key = (row["case"], row["family"], row["fact"], row["expected"])
        grouped.setdefault(key, {})
        grouped[key].setdefault(row["method"], {"correct": 0, "answered": 0, "n": 0})
        item = grouped[key][row["method"]]
        item["correct"] += row["correct"] == "TRUE"
        item["answered"] += row["answered"] == "TRUE"
        item["n"] += 1
    rows = []
    for (case, family, fact, expected), methods in sorted(grouped.items()):
        cells = [case, family, fact, expected]
        for method in METHODS:
            item = methods[method]
            cells.append(f'{item["correct"]}/{item["n"]}')
            cells.append(f'{item["answered"]}/{item["n"]}')
        rows.append(cells)
    assert len(rows) == 480
    return rows


def human_evaluation_rows():
    expert = read_csv(RESULTS / "human_evaluation/expert_dimension_summary.csv")
    users = read_csv(RESULTS / "human_evaluation/user_feature_summary.csv")
    agreement = read_csv(RESULTS / "human_evaluation/expert_agreement_summary.csv")[0]
    rows = []
    for row in expert:
        if row["dimension"] == "overall":
            rows.append(["外部评分者", "总体 0-2 评分", int(row["n"]), f'{float(row["mean"]):.3f}',
                         f'{float(row["sd"]):.3f}', "", "3 人 × 30 案例 × 8 维度"])
        else:
            rows.append(["外部评分者", row["dimension"], int(row["n"]), f'{float(row["mean"]):.3f}',
                          f'{float(row["sd"]):.3f}', "", "允许 0-2 分；本批无 0 分记录"])
    rows.append(["外部评分者", "case-dimension 一致性", int(agreement["n"]), "", "",
                 f'{float(agreement["exact_agreement_percent"]):.1f}%', "完全一致；相差不超过1分为100%"])
    for row in users:
        if row["feature_id"] == "overall":
            rows.append(["功能感知", "全部功能", int(row["n"]), f'{float(row["mean_score"]):.2f}',
                         f'{float(row["sd_score"]):.2f}', f'{float(row["useful_percent"]):.1f}%',
                          "0-100 功能感知评分；1=有用，2=没用"])
        else:
            rows.append(["功能感知", row["feature"], int(row["n"]), f'{float(row["mean_score"]):.2f}',
                         f'{float(row["sd_score"]):.2f}', f'{float(row["useful_percent"]):.1f}%',
                          "0-100 功能感知评分；1=有用，2=没用"])
    return rows


def hydrate_supplementary_tables(specs):
    summary = {row["method"]: row for row in read_csv(RESULTS / "common_task_scale_480/common_task_summary.csv")}
    rows = []
    labels = {
        "configured_models:step-5-preview": ("StepFun step-5-preview", "OpenAI-compatible; response_format disabled"),
        "configured_models:deepseek-flash": ("DeepSeek deepseek-flash", "OpenAI-compatible; response_format disabled"),
    }
    for method, (label, protocol) in labels.items():
        row = summary[method]
        note = "120 个案例、480 个目标事实；完整小数提示修复后评分"
        rows.append([label, protocol, int(row["expected"]), int(row["correct"]), int(row["answered"]),
                     f'{100 * float(row["accuracy"]):.1f}%', f'{100 * float(row["coverage"]):.1f}%', note])
    for spec in specs:
        if spec["id"] == "S4":
            spec["rows"] = rows
        elif spec["id"] == "S5":
            spec["rows"] = human_evaluation_supplement_rows()
    return specs


def human_evaluation_supplement_rows():
    expert = read_csv(RESULTS / "human_evaluation/expert_dimension_summary.csv")
    users = read_csv(RESULTS / "human_evaluation/user_feature_summary.csv")
    agreement = read_csv(RESULTS / "human_evaluation/expert_agreement_summary.csv")[0]
    rows = []
    for row in expert:
        rows.append(["外部评分者", row["dimension"], int(row["n"]), f'{float(row["mean"]):.3f}',
                     f'{float(row["sd"]):.3f}',
                     f'{float(agreement["exact_agreement_percent"]):.1f}%' if row["dimension"] == "overall" else "",
                      "允许 0-2 分；3 位外部评分者；30 个匿名案例；本批无 0 分记录"])
    for row in users:
        rows.append(["功能感知", row["feature"], int(row["n"]), f'{float(row["mean_score"]):.2f}',
                     f'{float(row["sd_score"]):.2f}', f'{float(row["useful_percent"]):.1f}%',
                      "0-100 功能感知评分；1=有用，2=没用"])
    return rows


def style_sheet(ws, widths=None):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="247A6B")
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=Side(style="hair", color="DDDDDD"))
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3F7F6")
        ws.row_dimensions[row[0].row].height = 45 if ws.title == "Display" else 30
    ws.row_dimensions[1].height = 36
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths[col-1] if widths else 22
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def export_tables(specs):
    csv_sources = {"1": [], "2": ["main_scale_480/ablation_summary.csv", "main_scale_480/model_calls.csv"], "4": ["main_scale_480/family_summary.csv", "main_scale_480/fact_scores.csv"],
                   "5": ["validation_summary.csv", "validation_ablation.csv"], "6": ["plot_challenges.csv"], "7": ["human_evaluation/expert_dimension_summary.csv", "human_evaluation/user_feature_summary.csv", "human_evaluation/expert_agreement_summary.csv"],
                   "8": ["main_scale_480/fact_scores.csv", "main_scale_480/case_manifest.csv"],
                   "S1": ["main_scale_480/model_calls.csv", "main_scale_480/ablation_summary.csv"],
                   "S2": ["related_tools/related_tool_tests.csv", "related_tools/transport_replay.csv",
                           "related_tools/final_install_status.csv"],
                   "S3": ["related_tools/pasilla_fact_checks.csv"],
                   "S4": ["common_task_scale_480/common_task_summary.csv", "common_task_scale_480/common_task_calls.csv"],
                   "S5": ["human_evaluation/expert_dimension_summary.csv", "human_evaluation/user_feature_summary.csv", "human_evaluation/expert_agreement_summary.csv"]}
    for spec in specs:
        wb = Workbook()
        ws = wb.active
        ws.title = "Display"
        ws.append(spec["headers"])
        for row in spec["rows"]:
            ws.append(row)
        style_sheet(ws, [max(17, width*7) for width in spec["widths_cm"]])
        notes = wb.create_sheet("Caption_Notes")
        notes.append(["字段", "内容"])
        for key, label in [("title", "表题"), ("note", "表注"), ("source", "数据来源")]:
            notes.append([label, spec[key]])
        notes.append(["版本", "中文稿 v0.7 / BioExplainR 0.3.1 / 2026-10-05"])
        style_sheet(notes, [18, 130])
        notes.row_dimensions[3].height = 220 if spec["id"] == "1" else 120
        for number, filename in enumerate(csv_sources.get(spec["id"], []), 1):
            data = read_csv(RESULTS / filename)
            raw = wb.create_sheet(f"Source_{number}")
            raw.append(list(data[0].keys()))
            for row in data:
                values = []
                for value in row.values():
                    if value in ("TRUE", "FALSE"):
                        value = value == "TRUE"
                    else:
                        try:
                            value = float(value) if any(c in value for c in ".eE") else int(value)
                        except ValueError:
                            pass
                    values.append(value)
                raw.append(values)
            style_sheet(raw)
        path = OUT / spec["file"]
        wb.save(path)
        reloaded = load_workbook(path, data_only=True)
    actual = [tuple(normalize_cell(value) for value in row) for row in reloaded["Display"].values]
    assert actual == [tuple(spec["headers"])] + [tuple(row) for row in spec["rows"]], path
    output_name = "tables_rendered.json" if any(spec["id"] == "1" for spec in specs) else "supplementary_tables_rendered.json"
    (OUT / output_name).write_text(json.dumps({"tables": specs}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def supplementary_bundle():
    paths = []
    for path in sorted(RESULTS.rglob("*")):
        if not path.is_file():
            continue
        relative = path.relative_to(RESULTS)
        if path.suffix in (".json", ".csv", ".txt") or (relative.parts[0] in ("browser", "related_tools") and path.suffix == ".png" and not path.name.startswith("paper-")):
            if path.name not in {"submission-audit.json", "reference_audit.json"}:
                paths.append(path)
    paths += [ROOT / f"research/{name}" for name in [
        "prepare_real_data.R", "run_model_experiments.R", "summarize_results.R", "audit_results.R",
        "run_main_scale_experiment.R", "generate_common_task_scale_benchmark.R",
        "browser_smoke.cjs", "run_related_tool_tests.R", "related_browser.cjs", "related_mock.cjs",
        "install_related_packages.R", "generate_study_materials.py", "run_common_task_benchmark.R",
        "run_compatibility_matrix.R", "run_complex_plot_checks.R", "common_task_methods.json"
        , "summarize_human_evaluation.py", "test_public_datasets.R"
    ]]
    paths += [ROOT / "docs/SUPPLEMENTARY_MATERIAL.md"]
    main = [p for p in (RESULTS / "model").glob("*.json") if not p.name.endswith("__prompt.json")]
    pilot = [p for p in (RESULTS / "pilot_precision/model").glob("*.json") if not p.name.endswith("__prompt.json")]
    assert len(main) == len(pilot) == 54, (len(main), len(pilot))
    assert len(list((RESULTS / "model").glob("*__prompt.json"))) == 54
    assert len(list((RESULTS / "pilot_precision/model").glob("*__prompt.json"))) == 54
    records = []
    with zipfile.ZipFile(OUT / "supplementary_records.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in paths:
            data = path.read_bytes()
            name = path.relative_to(ROOT).as_posix()
            archive.writestr(name, data)
            records.append({"path": name, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()})
    with (OUT / "supplementary_manifest.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["path", "bytes", "sha256"])
        writer.writeheader()
        writer.writerows(records)
    return len(records)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    # Remove only generated figure files so obsolete panel names cannot remain in
    # the delivery directory after a figure redesign.
    for pattern in ("Figure*.pdf", "Figure*.png", "Figure*.svg"):
        for path in OUT.glob(pattern):
            path.unlink()
    for obsolete in ("Table_1a_input_comparison.xlsx", "Table_1b_check_comparison.xlsx"):
        path = OUT / obsolete
        if path.exists():
            path.unlink()
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11, "pdf.fonttype": 42,
                         "svg.fonttype": "none", "axes.labelsize": 11, "axes.unicode_minus": False})
    scores = read_csv(MAIN_RESULTS / "fact_scores.csv")
    assert len(scores) == 4320
    architecture()
    rounds = ablation(scores)
    biological_figures()
    tables = hydrate_tables()
    export_tables(tables)
    count = supplementary_bundle()
    supplementary = json.loads((ROOT / "docs/supplementary_tables.json").read_text(encoding="utf-8"))["tables"]
    supplementary = hydrate_supplementary_tables(supplementary)
    export_tables(supplementary)
    panel_pdfs = sorted(path.name for path in OUT.glob("Figure_*[ABCD]_*.pdf"))
    print(json.dumps({"figure_pdfs": len(list(OUT.glob("Figure*.pdf"))), "panel_pdfs": panel_pdfs,
                      "editable_tables": len(tables) + len(supplementary), "bundled_records": count,
                      "repeat_rates": rounds}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
