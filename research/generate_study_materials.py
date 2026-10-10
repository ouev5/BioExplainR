"""Create blinded expert-review and user-study collection materials.

The outputs are empty collection templates plus anonymized case prompts.  They
contain no fabricated participant or expert scores.  Case answers remain in a
separate investigator-only key so that the review package can be distributed
without revealing the expected interpretation.
"""
import csv
import hashlib
import json
import random
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "research/study_materials"
RESULTS = ROOT / "research/results"
CASE_NAMES = {
    "C01": "airway_case_a",
    "C02": "airway_case_b",
    "C03": "airway_case_c",
    "C04": "airway_case_d",
    "C05": "airway_case_e",
    "C06": "airway_case_f",
    "C07": "slice_case_a",
    "C08": "slice_case_b",
    "C09": "slice_case_c",
    "C10": "slice_case_d",
    "C11": "slice_case_e",
    "C12": "slice_case_f",
    "C13": "ora_case_a",
    "C14": "ora_case_b",
    "C15": "ora_case_c",
    "C16": "pca_case_a",
    "C17": "pca_case_b",
    "C18": "pca_case_c",
    "C19": "sce_case_a",
    "C20": "sce_case_b",
    "C21": "sce_case_c",
    "C22": "diagnostic_case_a",
    "C23": "diagnostic_case_b",
    "C24": "diagnostic_case_c",
    "C25": "plot_case_a",
    "C26": "plot_case_b",
    "C27": "plot_case_c",
    "C28": "fix_case_a",
    "C29": "fix_case_b",
    "C30": "fix_case_c",
}


def facts():
    return json.loads((RESULTS / "reference_facts.json").read_text(encoding="utf-8"))


def case_definition(case_id, key, variant):
    if key.startswith("airway"):
        threshold = 0 if variant % 2 else 1
        f = facts()[f"airway_lfc{threshold}"]
        return {
            "case_id": case_id, "case_type": "差异表达阈值解释", "source": "匿名数据集 A",
            "prompt": f"请解释一个差异表达结果。阈值为 P_adj < 0.05 且 |log2FC| > {threshold}。请填写：特征总数、显著特征数、上调数、下调数、P_adj 缺失数；说明这些数字能支持和不能支持什么结论。",
            "facts": f, "expected": "只使用提供的统计事实；区分 P_adj、log2FC 方向和生物学假设。"
        }
    if key.startswith("slice"):
        f = facts()["complete_slice120"]
        return {
            "case_id": case_id, "case_type": "有界结果子表解释", "source": "匿名数据集 B",
            "prompt": "请解释一份已经筛选并按显著性排序的 120 行差异表达子表。填写：行数、满足 P_adj < 0.05 的行数、上调数、下调数、P_adj 缺失数；说明这些数字不能代表完整实验结果的哪些方面。",
            "facts": f, "expected": "明确这是 120 行子集，不能把子表分布当作完整结果。"
        }
    if key.startswith("ora"):
        f = facts()["reactome_ora"]
        return {
            "case_id": case_id, "case_type": "富集结果解释", "source": "匿名数据集 C",
            "prompt": "请解释一个过度表示分析结果。填写：检验条目数、显著条目数、分析方法；说明显著富集是否等于通路激活或因果机制。",
            "facts": f, "expected": "区分 ORA 的统计结果、通路注释与活性/因果结论。"
        }
    if key.startswith("pca"):
        f = facts()["airway_pca"]
        return {
            "case_id": case_id, "case_type": "PCA 图形解释", "source": "匿名数据集 D",
            "prompt": "请解释一个 PCA 摘要。填写：观测数、变量数、PC1 和 PC2 方差解释比例；说明 PCA 分离趋势不能单独证明什么。",
            "facts": f, "expected": "百分比应保持足够精度，并说明 PCA 是探索性表示。"
        }
    if key.startswith("sce"):
        f = facts()["pbmc3k_sce"]
        return {
            "case_id": case_id, "case_type": "单细胞对象摘要", "source": "匿名数据集 E",
            "prompt": "请解释一个单细胞对象摘要。填写：细胞数、特征数、样本数；说明单个供体中的细胞不能作为独立生物学重复。",
            "facts": f, "expected": "区分细胞数与样本/供体数，不能进行供体层面泛化。"
        }
    if key.startswith("diagnostic"):
        diagnostic = [
            ("P_adj 被误当作 raw P", "说明原始 P 与校正 P 的区别，并指出需要检查输入列。"),
            ("参考水平被反转", "指出比较方向和因子参考水平必须与研究问题一致。"),
            ("把细胞当作重复", "指出单细胞数量不等于生物学重复数量。"),
        ][variant % 3]
        return {
            "case_id": case_id, "case_type": "错误诊断", "source": "匿名错误现场",
            "prompt": f"错误现场提示：{diagnostic[0]}。请写出可能原因、需要检查的对象或设计信息，以及一项不能未经验证就执行的修复建议。",
            "facts": {"issue": diagnostic[0]}, "expected": diagnostic[1]
        }
    if key.startswith("plot"):
        plot = [
            "火山图的比较方向与源结果不一致",
            "PCA 轴标签与真实方差解释率不一致",
            "富集气泡图点大小不再对应 Count",
        ][variant % 3]
        return {
            "case_id": case_id, "case_type": "图形来源核对", "source": "匿名图形现场",
            "prompt": f"图形核对提示：{plot}。请说明需要对照哪些源数据映射，哪些视觉观察不能单独推出生物学机制。",
            "facts": {"issue": plot}, "expected": "把图形映射核对与生物学结论分开。"
        }
    fix = [
        "修复代码可以运行，但比较方向与原设计不一致",
        "修复代码输出行数改变",
        "修复代码改变了因子水平顺序",
    ][variant % 3]
    return {
        "case_id": case_id, "case_type": "修复语义保持", "source": "匿名修复现场",
        "prompt": f"候选修复提示：{fix}。请判断应声明哪些不变量，并解释为什么‘代码运行成功’不足以证明修复正确。",
        "facts": {"issue": fix}, "expected": "至少检查输出结构、比较方向、设计与研究目标。"
    }


def write_csv(path, rows, fields):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        if rows and isinstance(rows[0], dict):
            writer.writerows(rows)
        else:
            csv.writer(handle).writerows(rows)


def workbook(path, sheets):
    book = Workbook()
    book.remove(book.active)
    for title, headers, rows in sheets:
        sheet = book.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="247A6B")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[chr(64 + min(index, 26))].width = 24
    book.save(path)


def anonymize_response(text):
    """Remove source/package identifiers while retaining the explanation text."""
    replacements = [
        (r"PBMC3k|PBMC", "单细胞数据集"),
        (r"airway", "数据集 A"),
        (r"Reactome", "通路数据库"),
        (r"DESeq2", "差异表达分析"),
        (r"SingleCellExperiment", "单细胞对象"),
        (r"SeuratObject", "单细胞对象"),
        (r"clusterProfiler|DOSE", "富集分析工具"),
        (r"GSE\d+", "公开数据编号"),
        (r"EuropePMC:[A-Za-z0-9:_-]+", "外部文献记录"),
        (r"R-HSA-\d+", "通路条目"),
        (r"ENSG\d+", "基因编号"),
        (r"Homo sapiens|Drosophila melanogaster", "研究物种"),
        (r"dexamethasone", "处理条件"),
        (r"地塞米松", "处理条件"),
        (r"\btrt\b|\buntrt\b", "处理组"),
    ]
    for pattern, replacement in replacements:
        text = re.sub(pattern, replacement, text, flags=re.IGNORECASE)
    return text


def build_blinded_explanations(cases):
    """Select 30 saved responses and randomize their visible review IDs."""
    family_keys = ["airway_lfc0", "airway_lfc1", "complete_slice120", "reactome_ora", "airway_pca", "pbmc3k_sce"]
    methods = ["direct_bounded", "facts_only", "facts_evidence"]
    candidates = []
    for family in family_keys:
        for method in methods:
            for replicate in (1, 2, 3):
                path = RESULTS / "model" / f"{family}__{method}__{replicate}.json"
                if not path.exists():
                    continue
                record = json.loads(path.read_text(encoding="utf-8"))
                text = record.get("response", {}).get("text")
                if text:
                    candidates.append({"family": family, "method": method, "replicate": replicate,
                                       "text": anonymize_response(text)})
    selected = []
    # Five saved responses per family gives 30 cases while retaining both
    # bounded and facts-driven explanations in every family.
    for family in family_keys:
        family_rows = [row for row in candidates if row["family"] == family]
        selected.extend(family_rows[:5])
    assert len(selected) == 30, len(selected)
    random.Random(20261004).shuffle(selected)
    visible = []
    key = []
    for index, row in enumerate(selected, 1):
        review_id = f"R{index:03d}"
        case = next(item for item in cases if item["case_id"].startswith("C"))
        family_to_type = {
            "airway_lfc0": "差异表达阈值解释", "airway_lfc1": "差异表达阈值解释",
            "complete_slice120": "有界结果子表解释", "reactome_ora": "富集结果解释",
            "airway_pca": "PCA 图形解释", "pbmc3k_sce": "单细胞对象摘要"
        }
        visible_case = {
            "review_id": review_id,
            "task_type": family_to_type[row["family"]],
            "task": "请评价以下匿名分析解释在数字、统计含义、生物学支持、不确定性和不能推出的结论方面的质量。请指出具体错误或可追溯之处。",
            "explanation": row["text"]
        }
        visible.append(visible_case)
        key.append({"review_id": review_id, "family": row["family"], "method": row["method"],
                    "replicate": row["replicate"]})
        (OUT / "expert_cases" / f"{review_id}.json").write_text(
            json.dumps(visible_case, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return visible, key


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "expert_cases").mkdir(parents=True, exist_ok=True)
    cases = []
    answer_key = []
    for index, (case_id, key) in enumerate(CASE_NAMES.items()):
        item = case_definition(case_id, key, index)
        cases.append({k: item[k] for k in ("case_id", "case_type", "source", "prompt")})
        answer_key.append(item)
        (OUT / f"{case_id}.json").write_text(json.dumps(cases[-1], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    review_cases, review_key = build_blinded_explanations(cases)
    (OUT / "investigator_answer_key.json").write_text(json.dumps(
        {"case_facts": answer_key, "review_source_key": review_key}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_csv(OUT / "expert_case_index.csv", cases, ["case_id", "case_type", "source", "prompt"])
    write_csv(OUT / "expert_review_index.csv",
              [{"review_id": case["review_id"], "task_type": case["task_type"], "task": case["task"]}
               for case in review_cases],
              ["review_id", "task_type", "task"])

    dimensions = ["numeric_accuracy", "statistical_concepts", "biological_support", "traceability",
                  "uncertainty", "readability", "error_localization", "semantic_preservation"]
    expert_rows = [[case["review_id"], "", dimension, "", "" if dimension else ""]
                   for case in review_cases for dimension in dimensions]
    write_csv(OUT / "expert_ratings_template.csv", expert_rows, ["case_id", "rater_id", "dimension", "score_0_to_2", "comment"])
    workbook(OUT / "expert_review_workbook.xlsx", [
        ("Instructions", ["field", "value"], [
            ["purpose", "Blind expert review; leave scores blank until an independent expert rates the anonymized response."],
            ["scale", "0 = incorrect/missing; 1 = partial; 2 = correct and traceable."],
            ["dimensions", "Use the eight dimensions in the Ratings sheet; comments should cite the case text, not the hidden source."],
            ["blinding", "Do not infer or record the original method, dataset name, model, or author identity."],
        ]),
        ("Case_Index", ["review_id", "task_type", "task"],
         [[c[x] for x in ("review_id", "task_type", "task")] for c in review_cases]),
        ("Ratings", ["case_id", "rater_id", "dimension", "score_0_to_2", "comment"], expert_rows),
    ])

    feature_rows = []
    features = [
        ("F1", "运行并记录 R 代码", "执行用户提交的 R 代码，保留运行记录、警告和错误现场"),
        ("F2", "解释 R 生物信息学对象", "解释差异表达、富集、PCA、单细胞等已有 R 对象"),
        ("F3", "生成常用分析图形", "生成 PCA、UMAP、富集图和热图等常用图形"),
        ("F4", "证据与来源追踪", "保存事实来源、证据 ID 和引用追踪信息"),
        ("F5", "生成分析报告", "输出包含事实、来源和检查状态的 HTML/Markdown 报告"),
    ]
    for participant in range(1, 31):
        for feature_id, feature_name, feature_description in features:
            feature_rows.append([f"P{participant:02d}", feature_id, feature_name,
                                  feature_description, "", "", ""])
    workbook(OUT / "user_study_workbook.xlsx", [
        ("Feature_Ratings", ["student_id", "feature_id", "feature", "feature_description",
                             "score_0_to_100", "useful_1_yes_2_no", "comment"], feature_rows),
    ])
    (OUT / "user_task_cards.json").write_text(json.dumps([
        {"task_id": "T1", "title": "解释差异表达阈值", "instructions": "根据分配给你的分析输出，填写显著数、方向数和校正值缺失数，并说明不能推出的生物学结论。"},
        {"task_id": "T2", "title": "识别结果范围", "instructions": "判断输入是完整结果还是有界子表，指出可以报告的范围以及不能泛化的内容。"},
        {"task_id": "T3", "title": "解释富集结果", "instructions": "区分富集显著性、数据库注释、通路活性和因果机制。"},
        {"task_id": "T4", "title": "解释 PCA/单细胞图形", "instructions": "填写可由图形或对象直接确认的维度信息，并说明细胞数和样本数的区别。"},
        {"task_id": "T5", "title": "诊断并约束修复", "instructions": "定位错误，提出需要检查的对象和不变量，并判断候选修复是否仅仅‘能运行’。"},
    ], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    randomizer = {
        "purpose": "Generate a reproducible group assignment after consented participant IDs are entered.",
        "seed": 20261004,
        "strata": "pre_score_quartile_or_prespecified_baseline_band",
        "command": "bio_randomize_learning_study(participants, strata, seed = 20261004)",
        "warning": "Do not place the assignment output in the blinded expert package. This file contains no participant result."
    }
    (OUT / "randomization_plan.json").write_text(json.dumps(randomizer, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUT / "README.md").write_text("""# BioExplainR 外部评价材料\n\n这些文件是外部评分者案例评分和功能感知评分的采集模板。案例评分请使用 `expert_review_workbook.xlsx` 的 `Ratings` 表和 `expert_ratings_template.csv`，其中 `R001-R030` 对应 `expert_cases/R001-R030.json`；`expert_review_index.csv` 是案例索引。`investigator_answer_key.json` 只供研究负责人保存，不应发给评分者或参与者。功能感知评分请使用 `user_study_workbook.xlsx` 的唯一工作表 `Feature_Ratings`。\n\n功能感知表预置 30 位评分参与者 × 5 项 BioExplainR 功能 = 150 行。每行只填功能评分（0-100）、是否有用（1=有用，2=没用）和可选备注。实际采集前需完成伦理与知情同意要求；不要填写姓名、邮箱或临床身份信息。所有模板评分字段均为空；项目中的聚合结果来自外部填写表，不代表正式用户实验。\n""", encoding="utf-8")

    files = sorted(p for p in OUT.rglob("*") if p.is_file())
    with zipfile.ZipFile(OUT / "blinded_case_package.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            if path.name in {"investigator_answer_key.json", "blinded_case_package.zip"}:
                continue
            archive.write(path, path.relative_to(OUT).as_posix())
    with zipfile.ZipFile(OUT / "blinded_case_package.zip") as archive:
        names = archive.namelist()
        assert sum(name.startswith("expert_cases/") and name.endswith(".json") for name in names) == 30
        assert "investigator_answer_key.json" not in names
    manifest = []
    for path in sorted(OUT.rglob("*")):
        if path.is_file() and path.name != "blinded_case_package.zip":
            manifest.append({"file": path.relative_to(OUT).as_posix(), "bytes": path.stat().st_size,
                             "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    (OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"cases": len(cases), "expert_ratings_rows": len(expert_rows), "participants": 30,
                      "package": str(OUT / "blinded_case_package.zip")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
