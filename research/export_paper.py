"""Export native DOCX elements and standalone HTML from the same Pandoc AST."""
import argparse
import json
from pathlib import Path
import re
import subprocess
import tempfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def typeset(text):
    return text.replace("P_adj", "P~adj~").replace("10^-6", "10⁻⁶").replace("10^-5", "10⁻⁵")


def expand_tables(source, manifest="tables_rendered.json"):
    tables = json.loads((ROOT / "docs/submission_assets" / manifest).read_text(encoding="utf-8"))["tables"]
    text = source.read_text(encoding="utf-8")
    for spec in tables:
        placeholder = f'::: {{.paper-table #table-{spec["id"]}}}\n:::'
        header = "| " + " | ".join(spec["headers"]) + " |"
        lines = [f'**{spec["title"]}**', "", header, "|" + "---|" * len(spec["headers"])]
        lines += ["| " + " | ".join(str(value).replace("|", "\\|") for value in row) + " |" for row in spec["rows"]]
        lines += ["", f'**表注：** {spec["note"]}', "", f'**来源：** {spec["source"]}']
        assert text.count(placeholder) == 1, placeholder
        text = text.replace(placeholder, "\n".join(lines))
    return typeset(text), tables


def plain(inlines):
    result = []
    for inline in inlines:
        kind, value = inline["t"], inline.get("c")
        if kind == "Str":
            result.append(value)
        elif kind in ("Space", "SoftBreak", "LineBreak"):
            result.append(" ")
        elif kind in ("Code", "Math"):
            result.append(value[1])
        elif kind in ("Strong", "Emph", "Subscript", "Superscript"):
            result.append(plain(value))
        elif kind == "Link":
            result.append(plain(value[1]))
        elif kind == "Span":
            result.append(plain(value[1]))
    return "".join(result)


def configure_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in [("Normal", 10.5), ("Title", 17), ("Subtitle", 11),
                       ("Heading 1", 14), ("Heading 2", 11.5), ("Heading 3", 11)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("202020")
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        for key in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
            fonts.attrib.pop(qn(f"w:{key}"), None)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.2
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(10)
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def add_inlines(paragraph, inlines, bold=False, italic=False, sub=False, sup=False):
    for inline in inlines:
        kind, value = inline["t"], inline.get("c")
        if kind in ("Strong", "Emph", "Subscript", "Superscript"):
            add_inlines(paragraph, value, bold or kind == "Strong", italic or kind == "Emph",
                        sub or kind == "Subscript", sup or kind == "Superscript")
            continue
        if kind == "Link":
            add_inlines(paragraph, value[1], bold, italic, sub, sup)
            continue
        if kind == "Span":
            add_inlines(paragraph, value[1], bold, italic, sub, sup)
            continue
        if kind == "Quoted":
            paragraph.add_run('"')
            add_inlines(paragraph, value[1], bold, italic, sub, sup)
            paragraph.add_run('"')
            continue
        if kind == "Str":
            text = value
        elif kind in ("Space", "SoftBreak"):
            text = " "
        elif kind == "LineBreak":
            text = "\n"
        elif kind == "Code":
            text = value[1]
        else:
            raise ValueError(f"Unsupported inline: {kind}")
        run = paragraph.add_run(text)
        run.bold, run.italic = bold, italic
        run.font.subscript, run.font.superscript = sub, sup
        if kind == "Code":
            run.font.name = "DejaVu Sans Mono"
            run.font.size = Pt(9)


def add_table(doc, spec):
    table = doc.add_table(rows=1, cols=len(spec["headers"]))
    table.autofit = False
    table.alignment = 1
    widths = spec["widths_cm"]
    # Scale portrait tables to the actual 17 cm text area.
    landscape = spec["id"] in ("1", "8")
    scale = min(1, 17/sum(widths)) if not landscape else min(1, 25.7/sum(widths))
    widths = [width*scale for width in widths]
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for cell, text in zip(table.rows[0].cells, spec["headers"]):
        cell.text = text
    for values in spec["rows"]:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            cell.text = str(value)
    tblw = table._tbl.tblPr.find(qn("w:tblW"))
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(Cm(sum(widths)).twips))
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "bottom", "insideH"]:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6" if edge != "insideH" else "2")
        element.set(qn("w:color"), "404040" if edge != "insideH" else "DDDDDD")
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for index, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9 if spec["id"] == "1" else 9.5)
                    run.bold = index == 0
            if index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EDF3F1")
                cell._tc.get_or_add_tcPr().append(shading)


def native_docx(ast, path, tables=None, legacy_layout=True):
    doc = configure_document()
    metadata = ast["meta"]
    for key, style in [("title", "Title"), ("subtitle", "Subtitle"), ("date", "Normal")]:
        if key in metadata:
            paragraph = doc.add_paragraph(style=style)
            add_inlines(paragraph, metadata[key]["c"])
    table_index = 0
    appendix = False

    def render(blocks, prefix=None):
        nonlocal table_index, appendix
        for block in blocks:
            kind, value = block["t"], block.get("c")
            if kind == "Header":
                level, attributes, inlines = value
                text = plain(inlines)
                if text.startswith("附录 A"):
                    appendix = True
                if legacy_layout and text.startswith("表 1 "):
                    section = doc.add_section(WD_SECTION.NEW_PAGE)
                    section.orientation = WD_ORIENT.LANDSCAPE
                    section.page_width, section.page_height = Cm(29.7), Cm(21)
                if legacy_layout and text.startswith("表 8"):
                    section = doc.add_section(WD_SECTION.NEW_PAGE)
                    section.orientation = WD_ORIENT.LANDSCAPE
                    section.page_width, section.page_height = Cm(29.7), Cm(21)
                if legacy_layout and text.startswith("表 2 主实验"):
                    section = doc.add_section(WD_SECTION.NEW_PAGE)
                    section.orientation = WD_ORIENT.PORTRAIT
                    section.page_width, section.page_height = Cm(21), Cm(29.7)
                if legacy_layout and text.startswith("附录 B"):
                    section = doc.add_section(WD_SECTION.NEW_PAGE)
                    section.orientation = WD_ORIENT.PORTRAIT
                    section.page_width, section.page_height = Cm(21), Cm(29.7)
                paragraph = doc.add_paragraph(style=f"Heading {min(level, 3)}")
                add_inlines(paragraph, inlines)
                if legacy_layout and appendix and text.startswith(("附录 A", "附录 B", "图 2 ", "图 3 ", "图 4 ", "表 1 ", "表 3 ", "表 4 ", "表 5 ", "表 6 ", "表 7 ", "表 8 ", "补充表 S1", "补充表 S2", "补充表 S3", "补充表 S4", "补充表 S5")):
                    paragraph.paragraph_format.page_break_before = True
            elif kind in ("Para", "Plain"):
                paragraph = doc.add_paragraph()
                if prefix:
                    paragraph.add_run(prefix)
                    prefix = None
                image = next((item for item in value if item["t"] == "Image"), None)
                if image:
                    attributes, caption, target = image["c"]
                    width = dict(attributes[2]).get("width", "16cm")
                    assert width.endswith("cm")
                    paragraph.alignment = 1
                    paragraph.paragraph_format.keep_with_next = True
                    paragraph.add_run().add_picture(str(ROOT / "docs" / target[0]), width=Cm(float(width[:-2])))
                else:
                    add_inlines(paragraph, value)
                    text = paragraph.text
                    if appendix and re.match(r"^(图 [1-4]\. |补充图 S[1-3]\. |表 [1-8][ab]?\. |补充表 S[1-5]\.)", text):
                        for run in paragraph.runs:
                            run.font.size = Pt(9.5)
                        paragraph.paragraph_format.keep_with_next = text.startswith(("表 ", "补充表"))
                    if text.startswith(("表注：", "来源：")):
                        for run in paragraph.runs:
                            run.font.size = Pt(8.5 if table_index == 1 else 9)
            elif kind == "Table":
                assert tables and table_index < len(tables)
                add_table(doc, tables[table_index])
                table_index += 1
            elif kind == "OrderedList":
                start, number_style, delimiter = value[0]
                for number, item in enumerate(value[1], start):
                    render(item, prefix=f"{number}. ")
            elif kind == "BulletList":
                for item in value:
                    render(item, prefix="- ")
            elif kind == "BlockQuote":
                render(value)
                for paragraph in doc.paragraphs[-len(value):]:
                    paragraph.paragraph_format.left_indent = Cm(.6)
            elif kind == "CodeBlock":
                paragraph = doc.add_paragraph(value[1])
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "DejaVu Sans Mono"
                    run.font.size = Pt(8.5)
            elif kind == "Div":
                render(value[1])
            elif kind == "Figure":
                render(value[2])
            else:
                raise ValueError(f"Unsupported block: {kind}")

    render(ast["blocks"])
    assert table_index == len(tables or [])
    doc.save(path)
    return doc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--pandoc", default="pandoc")
    parser.add_argument("--data-dir")
    args = parser.parse_args()
    expanded, tables = expand_tables(ROOT / "docs/PAPER_DRAFT.md")
    supplementary, supplementary_tables = expand_tables(ROOT / "docs/SUPPLEMENTARY_MATERIAL.md",
                                                        "supplementary_tables_rendered.json")
    with tempfile.TemporaryDirectory(prefix="bioexplainr-paper-", dir="/tmp/opencode") as tmp:
        for name, content, specs in [("BioExplainR-paper-draft", expanded, tables),
                                      ("BioExplainR-supplementary", supplementary, supplementary_tables)]:
            source = Path(tmp) / f"{name}.md"
            source.write_text(content, encoding="utf-8")
            base = [args.pandoc, str(source), "--standalone", f"--resource-path={ROOT / 'docs'}:{ROOT}"]
            if args.data_dir:
                base.append(f"--data-dir={args.data_dir}")
            ast = json.loads(subprocess.check_output(base + ["-t", "json"]))
            output = ROOT / "docs" / name
            doc = native_docx(ast, output.with_suffix(".docx"), specs)
            subprocess.run(base + ["--embed-resources", f"--css={ROOT / 'docs/paper.css'}", "-o", str(output.with_suffix(".html"))], check=True)
            print(f'{name}: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(doc.inline_shapes)} figures')
    exported = Document(ROOT / "docs/BioExplainR-paper-draft.docx")
    assert len(exported.inline_shapes) == 7 and len(exported.tables) == len(tables)
    assert len(exported.sections) == 5
    for table, spec in zip(exported.tables, tables):
        assert [[cell.text for cell in row.cells] for row in table.rows] == [spec["headers"]] + [[str(value) for value in row] for row in spec["rows"]]
    supplement = Document(ROOT / "docs/BioExplainR-supplementary.docx")
    assert len(supplement.tables) == len(supplementary_tables)
    assert len(supplement.inline_shapes) == 2


if __name__ == "__main__":
    main()
