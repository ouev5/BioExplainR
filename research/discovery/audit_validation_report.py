"""Audit new numeric exports, source arrays, report tables and historical manifests."""
import argparse
import ast
import csv
import gzip
import hashlib
import json
import zipfile
from pathlib import Path

import numpy as np
from scipy import stats
import fitz
import openpyxl
from docx import Document

from validation_sources import ROOT, HERE, record, utc
from validation_integrity import old_integrity


def rows(path):
    with path.open(newline="", encoding="utf-8-sig") as stream:
        return list(csv.DictReader(stream))


def close(a, b):
    assert np.isclose(float(a), float(b), rtol=1e-7, atol=1e-9), (a, b)


def audit(source, dest):
    historical = old_integrity(source)
    for path in [source / "r_expression_audit.json", source / "asm_early/r_early_audit.json",
                 source / "gr_public/python_numeric_audit.json", source / "protein_public/r_protein_audit.json",
                 source / "human_protein/r_human_protein_audit.json", source / "bioexplainr_validation_workflow.json"]:
        assert json.loads(path.read_text())["status"] == "pass", path
    # Independent exact ASM count-model arithmetic (R's fitted moderation retained).
    design = rows(source / "voom_exact_design.csv")
    cols = [k for k in design[0] if k != "sample"]
    x = np.asarray([[float(r[k]) for k in cols] for r in design])
    contrast = rows(source / "voom_exact_contrasts.csv")
    assert [r["coefficient"] for r in contrast] == cols and np.linalg.matrix_rank(x) == 9
    wls = rows(source / "voom_exact_wls_inputs.csv")
    checks = 0
    for r in rows(source / "voom_exact_results.csv"):
        subset = {s["sample"]: s for s in wls if s["symbol"] == r["symbol"]}
        y = np.asarray([float(subset[s["sample"]]["logCPM"]) for s in design])
        w = np.asarray([float(subset[s["sample"]]["weight"]) for s in design])
        cov = np.linalg.inv(x.T @ (x*w[:, None])); beta = cov @ x.T @ (w*y)
        c = np.asarray([float(s[r["contrast"]]) for s in contrast])
        s2 = float(next(iter(subset.values()))["s2_post"]); df = float(next(iter(subset.values()))["df_total"])
        effect = c @ beta; se = np.sqrt(c @ cov @ c * s2); width = stats.t.ppf(.975, df)*se
        for key, value in zip(["effect", "se", "ci_low", "ci_high", "pvalue"], [effect, se, effect-width, effect+width, 2*stats.t.sf(abs(effect/se), df)]):
            close(value, r[key]); checks += 1
    rr = rows(source / "voom_exact_results.csv"); p = np.asarray([float(r["pvalue"]) for r in rr]); order = np.argsort(p)
    adjust = np.ones(44); adjust[order] = np.minimum(1, np.maximum.accumulate(p[order]*np.arange(44, 0, -1)))
    for r, hp in zip(rr, adjust): close(r["holm_44"], hp)
    # Verify selected human XLSX export fidelity and all 144 original cells.
    human = source / "human_protein"
    wb = openpyxl.load_workbook(human / "Table S2.xlsx", read_only=True, data_only=True)
    selected_checks = 0
    for sheet in ["S2_24hProteome_DiffAbundant", "S4_72hProteome_DiffAbundant", "S6_24hSecretome_DiffAbundant", "S8_72hSecretome_DiffAbundant"]:
        exported = rows(human / (sheet + ".csv"))
        original = list(wb[sheet].iter_rows(values_only=True))
        assert len(exported) == len(original)
        for a, b in zip(exported, original):
            assert list(a.values()) == ["" if v is None else str(v) for v in b]
    for r in rows(human / "human_selected_measurements.csv"):
        original = list(wb[r["source_sheet"]].iter_rows(values_only=True))
        close(original[int(r["source_row"])-1][int(r["source_column"])-1], r["value"])
        selected_checks += 1
    # Preserve the actual (not fabricated raw) adult time statistics.
    time = rows(source / "published_adult_ASM_time_tests.csv")
    assert len(time) == 33
    assert next(r for r in time if r["symbol"] == "PER2" and r["comparison1"] == "t24 Veh")["published_Tukey_pvalue_text"] == "0.7831"
    pair_checks = 0
    csv_tables = []
    for path in sorted((dest / "tables").glob("*.csv")):
        data = rows(path); csv_tables.append(data)
        w = openpyxl.load_workbook(path.with_suffix(".xlsx"), data_only=True)
        values = list(w["数据"].values)
        assert [[str(v) for v in row] for row in values] == [list(data[0])] + [list(r.values()) for r in data]
        pair_checks += 1
    asm = rows(source / "expression_results.csv"); voom = rows(source / "voom_exact_results.csv")
    for r in csv_tables[0]:
        a = next(a for a in asm if a["symbol"] == r["基因"] and a["study"] == "GSE130715" and a["contrast"] == "GC_control")
        v = next(v for v in voom if v["symbol"] == r["基因"] and v["contrast"] == "GC_control")
        assert r["供体配对效应"] == f'{float(a["effect"]):.3f}' and r["Holm44 P"] == f'{float(a["holm_44"]):.3g}'
        assert r["计数模型效应"] == f'{float(v["effect"]):.3f}'
    for r in csv_tables[2]:
        study = "GSE342887" if r["测量"] == "成熟RNA" else "GSE342771"
        contrast = "CORT_alpha" if r["对比"] == "WT rescue药物效应" else "CORT_x_AF2"
        a = next(a for a in rows(source / "gr_public/gr_results.csv") if a["study"] == study and a["symbol"] == r["基因"] and a["contrast"] == contrast)
        assert r["效应"] == f'{float(a["effect"]):.3f}' and r["Holm55 P"] == f'{float(a["holm_55"]):.3g}'
    for r in csv_tables[3]:
        a = next(a for a in rows(human / "human_selected_results.csv") if a["symbol"] == r["基因"] and a["time_h"] + "h" == r["暴露"])
        assert r["蛋白配对效应"] == f'{float(a["effect"]):.3f}' and r["Holm22 P"] == f'{float(a["holm_22"]):.3g}'
    basename = "BioExplainR-公开数据验证-结果增补-v1"
    doc = Document(dest / (basename + ".docx"))
    assert len(doc.tables) == 4 and len(doc.inline_shapes) == 3
    for table, data in zip(doc.tables, csv_tables):
        assert [[c.text for c in row.cells] for row in table.rows] == [list(data[0])] + [list(r.values()) for r in data]
    with fitz.open(dest / (basename + ".pdf")) as pdf:
        n_pages = len(pdf); spans = 0
        text = "".join(page.get_text() for page in pdf)
        for term in ["0.7831", "METTL7A", "PER2", "未确认", "Holm", "六成人"]:
            assert term in text, term
        for page in pdf:
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        assert (page.rect + (-1, -1, 1, 1)).contains(fitz.Rect(span["bbox"])), span["text"]
                        spans += 1
    figures = []
    for path in sorted((dest / "figures").glob("*.pdf")):
        with fitz.open(path) as pdf:
            assert len(pdf) == 1 and len(pdf[0].get_drawings()) > 0 and not pdf[0].get_images()
            figures.append(path.name)
    html_text = (dest / (basename + ".html")).read_text()
    assert html_text.count("data:image/png;base64,") == 3 and "<script" not in html_text
    cross = rows(source / "cross_geo_series.csv")
    # Count actual results, never use an expected number in prose unchecked.
    source_counts = {"initial_unique_series": len(rows(source / "geo_series.csv")), "expanded_unique_series": len(cross),
                     "expanded_literature_records": len(rows(source / "cross_literature.csv"))}
    md = (dest / "results.md").read_text()
    assert f'{len(cross)}个唯一series' in md, source_counts
    scripts = sorted(set(HERE.glob("*validation*.py")))
    for path in scripts:
        ast.parse(path.read_text(), filename=str(path))
        assert all(line == line.rstrip() for line in path.read_text().splitlines()), path
    result = {"status": "pass", "utc": utc(), "old_integrity": historical, "source_counts": source_counts,
              "exact_ASM_WLS_numeric_checks": checks, "human_original_XLSX_measurements": selected_checks,
              "csv_xlsx_pairs": pair_checks, "report_pages": n_pages, "report_tables": 4, "report_figures": 3,
              "PDF_spans": spans, "vector_original_figures": figures,
              "scope": "Fresh public-data audit; published adult Ct absent, human protein selected-array limitations retained; not global project acceptance"}
    if not (dest / "validation_manifest.json").exists():
        path = dest / "audit.json"
        if not path.exists():
            path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__); p.add_argument("source", type=Path); p.add_argument("report", type=Path)
    a = p.parse_args(); audit(a.source.resolve(), a.report.resolve())
