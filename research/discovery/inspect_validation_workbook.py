"""Describe supplemental workbook structure before numerical protein contrasts."""
import sys
from pathlib import Path
import openpyxl
from validation_sources import json_out,utc
dest=Path(sys.argv[1]).resolve()/"human_protein"
w=openpyxl.load_workbook(dest/"Table S2.xlsx",data_only=True,read_only=True)
summary=[]
for sheet in w:
    r=list(sheet.iter_rows(values_only=True))
    summary.append({"sheet":sheet.title,"rows":sheet.max_row,"columns":sheet.max_column,"first_rows":r[:3]})
    print(sheet.title,sheet.max_row,sheet.max_column,flush=True)
    for line in r[:2]:print(line[:55],flush=True)
json_out(dest/"workbook_design_inspection.json",{"utc":utc(),"sheets":summary,"timing":"Headers/layout inspected; no candidate quantitative contrast."})
