"""Check exported numeric tables, document content/geometry and real browser layouts."""
import argparse
import csv
import json
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import fitz
import openpyxl
from docx import Document
from playwright.sync_api import sync_playwright

from build_secretome_provenance_report import NAME
from secretome_provenance_complete import rows
from validation_sources import ROOT, utc, record, json_out


def main(source,report):
    assert not (report / "audit.json").exists()
    complete=source / "complete"
    pairs=[]
    for path in sorted((report / "tables").glob("*.csv")):
        with path.open(newline="") as stream:data=list(csv.reader(stream))
        wb=openpyxl.load_workbook(path.with_suffix(".xlsx"),read_only=True,data_only=True)
        xlsx=[["" if v is None else str(v) for v in r] for r in wb["数据"].iter_rows(values_only=True)]
        assert data==xlsx,path
        pairs.append(dict(name=path.name,rows=len(data)-1,columns=len(data[0]),csv=record(path),xlsx=record(path.with_suffix(".xlsx"))))
    assert len(pairs)==12
    source_summary=rows(complete / "cross_array_summary.csv")
    assert rows(report / "tables/cross_array_summary.csv")==source_summary
    assert rows(report / "tables/EV_selected_glycolysis.csv")==rows(complete / "EV_selected_glycolysis.csv")
    assert all(float(r["published_pvalue"])>.05 for r in rows(report / "tables/EV_selected_glycolysis.csv"))
    doc=Document(report / (NAME+".docx"));assert len(doc.tables)==3 and len(doc.inline_shapes)==2
    for table,name in zip(doc.tables,["Table_1_four_arrays","Table_2_public_sources","Table_3_EV_selected_glycolysis"]):
        with (report / "tables" / (name+".csv")).open(newline="") as stream:data=list(csv.reader(stream))
        assert [[c.text for c in r.cells] for r in table.rows]==data,name
    pdf=fitz.open(report / (NAME+".pdf"));text="\n".join(p.get_text() for p in pdf)
    compact="".join(text.split())
    for marker in ["0/803","4,543/4,543","14,829/14,829","2,128/2,128","281,232","31,228","23,436","168", "BA184", ".0912", "doi:10.1042/BSR20171090"]:
        assert marker in compact,marker
    bad=[]
    for pnum,page in enumerate(pdf,1):
        for block in page.get_text("dict")["blocks"]:
            if block.get("type")!=0:continue
            for line in block["lines"]:
                for span in line["spans"]:
                    r=fitz.Rect(span["bbox"])
                    if not fitz.Rect(page.rect.x0-1,page.rect.y0-1,page.rect.x1+1,page.rect.y1+1).contains(r):
                        bad.append(dict(page=pnum,text=span["text"],bbox=list(r)))
    assert not bad,bad
    figure_checks=[]
    for path in sorted((report / "figures").glob("*.pdf")):
        f=fitz.open(path);assert len(f)==1 and not f[0].get_images()
        assert f[0].get_text().strip()
        figure_checks.append(dict(path=record(path),pages=len(f),raster_images=len(f[0].get_images())))
    browser_dir=report / "browser_final";browser_dir.mkdir()
    final=complete / "report_final"
    pages=[report / (NAME+".html"),final / "BioExplainR-provenance-zh.html",final / "BioExplainR-provenance-en.html"]
    browser_checks=[]
    with sync_playwright() as playwright:
        browser=playwright.chromium.launch(headless=True,args=["--no-sandbox"])
        for path in pages:
            for width in [1440,390]:
                page=browser.new_page(viewport=dict(width=width,height=1000));errors=[]
                page.on("pageerror",lambda err:errors.append(str(err)))
                page.goto(path.as_uri(),wait_until="networkidle")
                stats=page.evaluate("""() => ({width:innerWidth,scroll:document.documentElement.scrollWidth,
                    images:[...document.images].map(x=>({complete:x.complete,width:x.naturalWidth})),
                    tables:document.querySelectorAll('table').length,text:document.body.innerText.length})""")
                assert stats["scroll"]<=width+1,(path,width,stats)
                assert not errors and stats["text"]>200
                assert all(r["complete"] and r["width"]>0 for r in stats["images"])
                if path==pages[0]:assert len(stats["images"])==2 and stats["tables"]==3
                else:
                    visible=page.inner_text("body")
                    assert "scope_limitations" in visible and "14829" in visible
                screenshot=browser_dir / (path.stem+"-"+str(width)+".png")
                page.screenshot(path=str(screenshot),full_page=True)
                browser_checks.append(dict(path=record(path),viewport=width,stats=stats,errors=errors,screenshot=record(screenshot)))
                page.close()
        browser.close()
    json_out(browser_dir / "audit.json",dict(status="pass",utc=utc(),checks=browser_checks))
    checks=[json.loads((complete / name).read_text()) for name in ["r_audit.json","independent_audit.json","EV_selected_table_audit.json"]]
    checks.append(json.loads((final / "bioexplainr_workflow.json").read_text()))
    assert all(r["status"]=="pass" for r in checks)
    acquisition=json.loads((complete / "source_acquisition_resume.json").read_text())
    requested=json.loads((complete / "fulltext_requests.json").read_text())["selected"]
    parsed=0
    for pmc in requested:
        path=source / f"sources/{pmc}.xml"
        if path.exists():
            assert ET.parse(path).getroot().tag=="article";parsed+=1
    json_out(complete / "scientific_ledger.json",dict(status="completed_and_audited",utc=utc(),
        exact_original_input_records=json.loads((complete / "freeze.json").read_text())["inputs"],
        arrays=4,decoded_run_columns=168,original_published_measurements=23436,
        cross_measurement_comparisons=281232,total_intensity_comparisons=31228,scale_formula_comparisons=4968,
        secretome24_positive_total_intensity_matches=0,secretome24_nonzero_LFQ_matches=0,
        public_literature_records=180,fulltext_requested_bounded=len(requested),fulltext_parsed=parsed,
        public_pride_unique=8,ASM_EV_patient_labels=6,ASM_EV_runs=18,GC_candidate_validation_arrays_obtained=0,
        EV_selected_glycolysis_proteins=8,EV_selected_reported_P_under_05=0,
        candidate_original_results="unchanged; no new GC candidate significance tests",
        causal_function_confirmed=False,online_model_calls=0,publication="paused",
        audit_reports=[record(complete / name) for name in ["r_audit.json","independent_audit.json","EV_selected_table_audit.json"]]+[record(final / "bioexplainr_workflow.json")]))
    import numpy,scipy,docx,matplotlib
    json_out(complete / "python_environment.json",dict(python=sys.version,versions=dict(numpy=numpy.__version__,scipy=scipy.__version__,
        openpyxl=openpyxl.__version__,docx=docx.__version__,matplotlib=matplotlib.__version__,PyMuPDF=fitz.VersionBind)))
    json_out(report / "audit.json",dict(status="pass",utc=utc(),script=record(Path(__file__)),pages=len(pdf),
        tables=3,figures=2,table_pairs=pairs,pdf_boundary_failures=bad,figure_checks=figure_checks,
        browser=record(browser_dir / "audit.json"),
        limitations="Automated content/geometry and real browser checks, not human visual or scientific signoff"))
    print(json.dumps(dict(status="pass",pages=len(pdf),table_pairs=len(pairs),browser_checks=len(browser_checks),fulltext_parsed=parsed)))


if __name__=="__main__":
    p=argparse.ArgumentParser();p.add_argument("source",type=Path);p.add_argument("report",type=Path)
    a=p.parse_args();main(a.source.resolve(),a.report.resolve())
