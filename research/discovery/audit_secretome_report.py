"""Check source-to-CSV/XLSX/Word/PDF/HTML correspondence and extracted input integrity."""
import argparse
import csv
import hashlib
import json
import re
import zlib
from pathlib import Path

import fitz
import openpyxl
from docx import Document

from build_secretome_report import NAME
from set_protein_extract import LocalRanges
from set_validation import rows
from validation_sources import ROOT, record, utc, json_out


def audit(source,report):
    for time in [24,72]:
        p=source/f"archives/{time}h/extraction.json"
        extracted=json.loads(p.read_text())
        assert extracted["status"] == "pass" and extracted["whole_archive_SHA1_verified"]
        inspected=json.loads((source/f"archives/{time}h/inspection.json").read_text())
        assert extracted["container_sha1"] == inspected["upstream_sha1"]
        segments=[]
        for r in extracted["range_records"]:
            segments.append((r["start"],r["end"],ROOT/r["body"]["path"]))
        header=json.loads((source/f"archives/{time}h/header_ranges.json").read_text())
        for r in header["range_records"]:
            assert record(ROOT/r["body"]["path"]) == r["body"] and record(ROOT/r["headers"]["path"]) == r["headers"]
            segments.append((r["start"],r["end"],ROOT/r["body"]["path"]))
        stream=LocalRanges(extracted["container_bytes"],segments)
        digest=hashlib.sha1(); remaining=extracted["container_bytes"]
        while remaining:
            blob=stream.read(min(1024*1024,remaining)); assert blob
            digest.update(blob); remaining-=len(blob)
        assert digest.hexdigest() == inspected["upstream_sha1"]
        for r in extracted["outputs"]:
            path=ROOT/r["path"]
            assert record(path) == {k:r[k] for k in ["path","bytes","sha256"]}
            assert zlib.crc32(path.read_bytes()) == r["crc32"]
        for r in extracted["range_records"]:
            assert r["status"] == "pass" and record(ROOT/r["body"]["path"]) == r["body"]
            h=(ROOT/r["attempts"][-1]["header"]["path"]).read_text()
            status=re.findall(r"(?m)^HTTP/\S+\s+(\d+)",h)
            ranges=re.findall(r"(?im)^content-range:\s*bytes\s+(\d+)-(\d+)/(\d+)",h)
            assert status[-1] == "206" and tuple(map(int,ranges[-1])) == (r["start"],r["end"],extracted["container_bytes"])
    table_checks=[]
    for path in sorted((report/"tables").glob("*.csv")):
        with path.open(newline="",encoding="utf-8-sig") as stream:
            data=list(csv.reader(stream))
        book=openpyxl.load_workbook(path.with_suffix(".xlsx"),data_only=True,read_only=True)
        cells=list(book["数据"].iter_rows(values_only=True))
        assert len(cells)==len(data)
        for a,b in zip(data,cells):
            assert a == ["" if v is None else str(v) for v in b],path
        if not path.name.startswith("Table_"):
            assert path.read_bytes() == (source/path.name).read_bytes()
        table_checks.append(dict(name=path.name,rows=len(data)-1))
    document=Document(report/(NAME+".docx")); assert len(document.tables)==3
    for t,path in zip(document.tables,sorted((report/"tables").glob("Table_*.csv"))):
        with path.open(newline="",encoding="utf-8-sig") as stream:
            data=list(csv.reader(stream))
        assert [[c.text for c in row.cells] for row in t.rows] == data,path
    export=json.loads((report/"export.json").read_text())
    with fitz.open(report/(NAME+".pdf")) as pdf:
        text="\n".join(p.get_text() for p in pdf)
        assert len(pdf)==export["pages"]
        for term in ["ANGPTL2","FSTL3","112","缺测","分泌速率"]:
            assert term in text,term
        for p in pdf:
            for block in p.get_text("blocks"):
                x0,y0,x1,y1,*_=block
                assert x0>=-1 and y0>=-1 and x1<=p.rect.width+1 and y1<=p.rect.height+1
    figures=[]
    for path in sorted((report/"figures").glob("*.pdf")):
        with fitz.open(path) as pdf:
            figures.append(dict(name=path.name,pages=len(pdf),raster_images=sum(len(p.get_images()) for p in pdf)))
            assert sum(len(p.get_images()) for p in pdf)==0
    html=(report/(NAME+".html")).read_text()
    assert html.count("<table>")==3 and html.count("data:image/png;base64,")==2
    for path in [source/"r_numeric_audit.json",source/"bioexplainr_workflow.json"]:
        assert json.loads(path.read_text())["status"] == "pass"
    result=dict(status="pass",utc=utc(),script=record(Path(__file__)),table_checks=table_checks,Word_tables=3,
                PDF_pages=export["pages"],PDF_text_boundary="pass",figures=figures,
                CRC_members=8,scope="Saved original inputs/hash/range/CRC and all CSV-XLSX/native Word correspondence; numeric independent R audit separate")
    json_out(report/"audit.json",result)
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=="__main__":
    p=argparse.ArgumentParser();p.add_argument("source",type=Path);p.add_argument("report",type=Path)
    a=p.parse_args();audit(a.source.resolve(),a.report.resolve())
