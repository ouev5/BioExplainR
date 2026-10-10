"""Chinese source-provenance follow-up report, complete traceable numeric attachments."""
import argparse
import base64
import csv
import html
import json
import shutil
import subprocess
from pathlib import Path

import fitz
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from secretome_provenance_complete import rows
from validation_sources import ROOT, utc, record, csv_out, json_out

NAME="分泌蛋白组-来源诊断与ASM公开证据续接"


def build(source,dest):
    dest.mkdir();fig=dest / "figures";fig.mkdir();tab=dest / "tables";tab.mkdir()
    complete=source / "complete"
    for name in ["r_audit.json","independent_audit.json","bioexplainr_workflow.json","supplement_review.json","EV_selected_table_audit.json"]:
        assert json.loads((complete / name).read_text())["status"]=="pass",name
    summary=rows(complete / "cross_array_summary.csv")
    own=[next(r for r in summary if r["array"]==r["expected_array"]==key and r["measure"]=="LFQ intensity")
        for key in ["proteome24","proteome72","secretome24","secretome72"]]
    table1=[{"数组":r["array"],"ALL一致/全部":r["all_intensity_equal"]+"/"+r["all_rows"],
        "非零ALL一致":r["all_positive_equal"],"非零LFQ一致/可比":r["matches_rounding"]+"/"+r["original_nonzero"],
        "原零/发表有限":r["original_zero"]} for r in own]
    table2=[{"来源":"PXD055968 / ASM EV","实际材料":"六患者/18run；54RAW/MSF/pdResult＋FASTA/映射；13页补充","可用范围":"真实ASM EV来源；未筛定量导出未取得；无GC/vehicle"},
        {"来源":"PMC6706008 / PTX3","实际材料":"ASM DEX分时上清ELISA；RU486；原图注/方法","可用范围":"已有PTX3分泌/药理GR支持；无逐供体重算，年龄/时点冲突"},
        {"来源":"PMC5686396 / cough","实际材料":"11健康/6咳嗽；FP＋polyIC；32因子Luminex","可用范围":"真实成人ASM上清；聚焦面板，非固定候选未筛蛋白组"}]
    selected=rows(complete / "EV_selected_glycolysis.csv")
    table3=[{"蛋白":r["symbol"],"发表fold change":r["published_fold_change"],"发表Expr P":r["published_pvalue"],"发表Expected":r["published_expected"]} for r in selected]
    for name,data in [("Table_1_four_arrays",table1),("Table_2_public_sources",table2),("Table_3_EV_selected_glycolysis",table3)]:
        csv_out(tab / (name+".csv"),data)
    for name in ["cross_array_summary.csv","sample_map.csv","secretome24_protein_offsets.csv","scale_formula_summary.csv",
        "project_feasibility.csv","public_evidence_scope.csv","EV_file_linkage.csv","EV_supplement_scope.csv","EV_selected_glycolysis.csv"]:
        shutil.copyfile(complete / name,tab / name)
    for path in sorted(tab.glob("*.csv")):
        data=rows(path);wb=openpyxl.Workbook();ws=wb.active;ws.title="数据";ws.append(list(data[0]))
        for r in data:ws.append(list(r.values()))
        ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
        wb.create_sheet("来源").append([str(source.relative_to(ROOT)),"来源诊断/原发表值；不是新GC检验、无候选效应校正"])
        wb.save(path.with_suffix(".xlsx"))
    plt.rcParams.update({"font.family":"DejaVu Sans","font.size":9,"pdf.fonttype":42})
    f,ax=plt.subplots(figsize=(8,3))
    counts=[int(r["original_nonzero"]) for r in own];matched=[int(r["matches_rounding"]) for r in own]
    ax.barh(range(4),[100*m/n for m,n in zip(matched,counts)],color=["#287f84","#287f84","#bd6846","#287f84"])
    for i,(m,n) in enumerate(zip(matched,counts)):
        ax.text(2,i,f"{m:,}/{n:,}",va="center",color="white" if m else "#bd6846")
    ax.set_yticks(range(4),["Proteome24h","Proteome72h","Secretome24h","Secretome72h"])
    ax.set_xlim(0,100);ax.set_xlabel("Original nonzero LFQ values matching published rounding (%)")
    ax.invert_yaxis();f.tight_layout();f.savefig(fig / "Figure_1_match.pdf");f.savefig(fig / "Figure_1_match.png",dpi=180);plt.close(f)
    offsets=rows(complete / "secretome24_protein_offsets.csv")
    f,ax=plt.subplots(figsize=(8,3.6));offsets=sorted(offsets,key=lambda r:float(r["delta_mean"]))
    ax.errorbar(range(len(offsets)),[float(r["delta_mean"]) for r in offsets],
        yerr=[float(r["delta_sd"]) for r in offsets],fmt="o",color="#287f84",capsize=2)
    ax.set_xticks(range(len(offsets)),[r["symbol"] for r in offsets],rotation=65,ha="right",fontsize=7)
    ax.set_ylabel("Original log2 LFQ - published value\nmean and within-protein SD (not CI)")
    ax.set_title("Secretome24h: protein-specific offsets; no fitted correction")
    f.tight_layout();f.savefig(fig / "Figure_2_offsets.pdf");f.savefig(fig / "Figure_2_offsets.png",dpi=180);plt.close(f)
    doc=Document()
    for style in ["Normal","Title","Heading 1","Heading 2"]:
        doc.styles[style].font.name="Noto Sans CJK SC"
        doc.styles[style]._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Noto Sans CJK SC")
    doc.styles["Normal"].font.size=Pt(10)
    sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.65)
    md=[];h=[]
    def heading(text,level=1,page=False):
        if page:doc.add_page_break()
        doc.add_heading(text,level);md.append("\n"+"#"*(level+1)+" "+text);h.append(f"<h{level+1}>{html.escape(text)}</h{level+1}>")
    def para(text):
        doc.add_paragraph(text);md.append(text);h.append("<p>"+html.escape(text)+"</p>")
    def table(data,caption):
        para(caption);t=doc.add_table(rows=1,cols=len(data[0]));t.style="Table Grid"
        for c,k in zip(t.rows[0].cells,data[0]):c.text=k
        for r in data:
            for c,v in zip(t.add_row().cells,r.values()):c.text=str(v)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:run.font.size=Pt(8)
        md.extend(["|"+"|".join(data[0])+"|","|"+"|".join(["---"]*len(data[0]))+"|"])
        md.extend("|"+"|".join(str(v) for v in r.values())+"|" for r in data)
        h.append("<div class='table-wrap'><table><tr>"+"".join("<th>"+html.escape(k)+"</th>" for k in data[0])+"</tr>"+
            "".join("<tr>"+"".join("<td>"+html.escape(str(v))+"</td>" for v in r.values())+"</tr>" for r in data)+"</table></div>")
    def figure(name,caption):
        doc.add_picture(str(fig / (name+".png")),width=Inches(6.7));para(caption)
        md.append(f"![{caption}](figures/{name}.png)")
        h.append("<figure><img alt='"+html.escape(caption,quote=True)+"' src='data:image/png;base64,"+base64.b64encode((fig / (name+".png")).read_bytes()).decode()+"'></figure>")
    title="分泌蛋白组来源诊断与成人ASM公开材料续接"
    doc.add_heading(title,0);md.append("# "+title);h.append("<h1>"+title+"</h1>")
    para("2026-10-10｜结果后来源诊断｜实际BioExplainR 0.4.0.9000｜在线模型调用0")
    heading("1 本轮结论与固定比较范围")
    para("24h存档secretome与发表表的差异仍未得到处理原因确认；已排查同来源四数组、原样本解码、LFQ/Intensity及肽数归一化。本轮取得新的真实ASM细胞外囊泡来源与六患者映射，但没有取得可用于固定候选GC/vehicle验证的未筛定量数组。ANGPTL2/FSTL3尚未确认GC分泌响应，旧FBN2范围敏感线索保留。")
    para("四个原MaxQuant proteinGroups与原解码、Table S2按SHA256冻结；4,043/3,097胞内原组、600/531上清原组，168个实际技术/重复列。按Protein IDs精确优先、唯一Majority IDs回退，不按基因选最优组。ALL总Intensity精确比较，DIFF相同日期/条件/重复编号比较非零log2 LFQ（容差0.000051）；零保持缺测。诊断不产生新蛋白效应P值。")
    table(table1,"表1 四数组与原发表对应工作表。ALL一致包含零；非零ALL列单独排除零零。原零/发表有限不当作原实测。")
    figure("Figure_1_match","图1 非零LFQ逐格一致比例。三个对应数组全部舍入一致，24h secretome为0/803；非参数效应/显著性未在本图评价。")
    heading("2 24h差异能排除什么",page=True)
    para("24h secretome ALL467位置中466可唯一对应，9个总Intensity相等全部是双方为零：正Intensity一致0。其他对应数组正值ALL一致3,870/3,943、3,002/3,017、377/380。三个数组非零LFQ分别4,543/4,543、14,829/14,829、2,128/2,128一致，表明本轮解析/日期解码可以正确恢复其尺度。")
    para("24h secretome发表23个选中蛋白共828技术值，原803非零LFQ均不同；另25原零却发表有限。交叉对应secretome72/胞内24/胞内72非零LFQ分别690/467/418，精确舍入匹配均0；不能解释成简单整表来源互换。未归一Intensity非零812也0匹配；四数组均无逐样本iBAQ列，iBAQ不可对应而不是测量0。")
    para("差值原log2 LFQ−发表值2.027–5.446；每个蛋白平均偏移2.577–5.408，显示蛋白特异结构。六固定诊断公式（LFQ或Intensity分别除以总/razor+unique/unique肽数）4,968位置均未匹配。蛋白内偏移相对小或高相关不能证明同一处理结果，拟合偏移不用于修正发表数值；更复杂MaxQuant/Perseus版本差异仍未排除。")
    figure("Figure_2_offsets","图2 23个发表选中蛋白的非零测量偏移均值与蛋白内SD；SD不是推断CI，蛋白由旧发表差异筛选，不能代表全数组。")
    para("原文明确四组分别处理、MaxQuant1.6.11/Perseus1.6.5及六供体、三个技术重复。24h来源处理差异、72h样本范围、旧proteome1.6.15冲突继续保存。本轮不能宣布作者错误或原Perseus复现成功。")
    heading("3 新成人ASM来源与已有候选蛋白证据",page=True)
    table(table2,"表2 实际取得的重点来源。文献证据与本轮原数组重算计数分别注明。")
    para("范围检索六Europe PMC查询分别63/60/30/5/1/25条，去重180文献；按公开顺序规则最多30篇请求全文，部分HTTP500，未说全部180篇均审读。PRIDE正确检索1/7/0/0条、唯一8项目，无关检索0；最初普通/projects接口忽略词返回相同100条，已明确排除其相关命中意义。八项目逐项适用性表公开；未取得合格数组不等于数组不存在。")
    para("PXD055968来源原文Celle等2025 iScience：患者原代BSM培养上清72h分离EV。原映射：非哮喘595/614/615，哮喘BA184/BA185/BA186，每标签3run共18；原表写biological replicates，同患者重复不可称18独立供体，43为全研究受试者，不是蛋白n。无GC/vehicle因子，临床所有哮喘者ICS暴露也不可改成体外GC因果比较。个体年龄与未知别名未完全核实。")
    para("54RAW/MSF/pdResult＋FASTA/文件映射共56清单项取得；未下载全部大RAW/搜索数据库，没有未筛定量文本导出。PRIDE写ProteomeDiscoverer2.5/Mascot2.5/Proline、2020数据库、5%背景推断；原文写PD3.0/CHIMERYS、2023数据库及低5%随机填补，冲突不能猜同一处理版本。真实13页出版社补充通过题名/作者/PII及PDF解析核对。")
    para("PTX3原ASM文献（2019 PLoS ONE）：肺切除来源原代ASM、DEX分时上清ELISA及RU486；属于已有PTX3蛋白/药理GR线索，不是本轮首次生物发现。缺逐供体原ELISA、年龄未列，五independent experiments不作五已核实供体；正文48h/图注24h与promoter叙述冲突保留，非NR3C1遗传因果。旧HTPC PTX3 Holm112=.592未通过不改变，不替代PER2/METTL7A/ANGPTL2/FSTL3验证。")
    heading("4 原文/补充冲突、核验与下一步",page=True)
    table(table3,"表3 PXD055968原补充选中糖酵解表原值（13页末页），不是本轮新统计。")
    para("原文称8个糖酵解蛋白significative upregulated，实际补充Expr P为.0912–.590，8/8大于.05。表中Expected Up为原IPA栏，不能把通路P或Z-score写成八蛋白显著；也不推断作者全数组无显著蛋白。以线性文本和PDF坐标双解析核对16个fold/P值。选中/通路表未见ANGPTL2/FSTL3/FBN2不表示完整数组未检出。")
    para("独立R从四原数组与原解码重建281,232逐测量比较、31,228总Intensity比较、168样本列、4,968公式与48摘要；独立Python不用openpyxl而解包Excel原XML/styling核对23,436发表测量、31,228ALL位置及EV六患者/18run。实际BioExplainR两个generic context/项目/中英文报告保存完整计数与范围，不把来源诊断包装为差异表达或P值；在线调用0。")
    para("优先后续是取得PXD055968实际Proline/PD未筛定量导出并确认版本/个体年龄/重复层级，再另冻结完整候选识别检查；即使识别成功也只回答ASM EV身份，GC响应仍需有药物配对数组。24h差异需原Perseus项目/处理版本材料。本轮没有人工作者联络、没有新湿实验；既有候选结果继续按原有效配对及全family解释。")
    para("复现与原始数据：research/results/discovery_secretome_provenance_v1/complete为最终输出，根首轮不完整比较保留；EXECUTION_NOTES逐项记录下载中断、解码/Excel日期/布尔类型、补充URL及分母修订。新provenance_manifest仅对旧secretome的根rgo依赖使用开始精确快照，其他旧科学原路径/hash；2,905保护文件核验。GitHub和外部上传继续暂停。")
    para("来源：[1] Stepanov等，Cells2022;11:3164，doi:10.3390/cells11193164，TableS2/PXD033504/033514/033534。[2] Celle等，iScience2025;28:112546，doi:10.1016/j.isci.2025.112546，PMC12141081/PXD055968。[3] Glucocorticoids regulate pentraxin-3 expression in human airway smooth muscle cells，PLoS ONE2019，doi:10.1371/journal.pone.0220772。[4] Impaired innate immune gene profiling in airway smooth muscle cells from chronic cough patients，BiosciRep2017，doi:10.1042/BSR20171090。全部原响应、方法/图表段落和获取hash保存。")
    docx=dest / (NAME+".docx");doc.save(docx)
    (dest / "results.md").write_text("\n\n".join(md)+"\n")
    css="body{font:16px/1.65 sans-serif;margin:auto;padding:24px;max-width:960px;color:#17343b}h1,h2{line-height:1.35}p{overflow-wrap:anywhere}.table-wrap{overflow-x:auto;max-width:100%}table{border-collapse:collapse;min-width:600px;width:100%;font-size:14px}td,th{border:1px solid #ccd8db;padding:8px;text-align:left}th{background:#e9f2f1}img{max-width:100%;height:auto}figure{margin:16px 0}@media(max-width:600px){body{padding:14px}h1{font-size:25px}}"
    (dest / (NAME+".html")).write_text("<!doctype html><html lang='zh'><meta charset='utf-8'><meta name='viewport' content='width=device-width, initial-scale=1'><title>"+title+"</title><style>"+css+"</style><body>"+"\n".join(h)+"</body></html>")
    cmd=["libreoffice","-env:UserInstallation=file:///tmp/opencode/provenance-lo","--headless","--convert-to","pdf","--outdir",str(dest),str(docx)]
    run=subprocess.run(cmd,capture_output=True,text=True,timeout=120)
    assert run.returncode==0 and (dest / (NAME+".pdf")).exists(),run.stdout+run.stderr
    pdf=fitz.open(dest / (NAME+".pdf"))
    json_out(dest / "export.json",dict(status="completed",utc=utc(),script=record(Path(__file__)),pages=len(pdf),
        tables=3,figures=2,csv_xlsx_pairs=len(list(tab.glob("*.csv"))),
        command=cmd,stdout=run.stdout,stderr=run.stderr,
        figure_scope="ReportWord/PDF containPNG previews, original two vector PDFs provided separately"))
    print(json.dumps(dict(pages=len(pdf),csv_xlsx_pairs=len(list(tab.glob('*.csv'))))))


if __name__=="__main__":
    p=argparse.ArgumentParser();p.add_argument("source",type=Path);p.add_argument("destination",type=Path)
    a=p.parse_args();build(a.source.resolve(),a.destination.resolve())
