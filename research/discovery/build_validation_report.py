"""Build a Chinese, result-driven public validation report with traceable tables."""
import argparse
import base64
import csv
import html
import json
import subprocess
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
import fitz
import openpyxl

from validation_sources import ROOT, csv_out, json_out, record, utc

GENES = ["PER2", "PER3", "METTL7A", "MAOA", "ANGPTL2", "FSTL3"]


def rows(path):
    with path.open(newline="", encoding="utf-8-sig") as stream:
        return list(csv.DictReader(stream))


def num(value):
    return f"{float(value):.3f}" if value != "" else "未取得"


def pf(value):
    return f"{float(value):.3g}"


def build(source, dest):
    dest.mkdir(exist_ok=False)
    fig = dest / "figures"; fig.mkdir()
    tab = dest / "tables"; tab.mkdir()
    asm = rows(source / "expression_results.csv")
    voom = rows(source / "voom_exact_results.csv")
    early = rows(source / "asm_early/early_results.csv")
    gr = rows(source / "gr_public/gr_results.csv")
    human = rows(source / "human_protein/human_selected_results.csv")
    direct = rows(source / "human_protein/human_direct_results.csv")
    protein = rows(source / "protein_public/protein_results.csv")
    table1 = []
    for gene in GENES:
        r = next(r for r in asm if r["study"] == "GSE130715" and r["symbol"] == gene and r["contrast"] == "GC_control")
        v = next(r for r in voom if r["symbol"] == gene and r["contrast"] == "GC_control")
        table1.append({"基因": gene, "供体配对效应": num(r["effect"]), "点态95%CI": f'{num(r["ci_low"])}至{num(r["ci_high"])}',
                       "Holm44 P": pf(r["holm_44"]), "方向": f'{r["positive_donors"]}↑/{r["negative_donors"]}↓',
                       "计数模型效应": num(v["effect"]), "计数Holm44 P": pf(v["holm_44"])})
    table2 = [
        {"来源／估计目标": "GSE135730，4h DEX−UT", "PER2效应": num(early[0]["effect"]), "校正P": pf(early[0]["holm_33"]), "独立单位与限制": "3培养/组；年龄及供体ID未知"},
        {"来源／估计目标": "GSE13168，8h FP−vehicle", "PER2效应": "0.691", "校正P": "1", "独立单位与限制": "3配对供体；年龄/别名未知"},
        {"来源／估计目标": "Bartman2021，24h DEX脉冲−同时间vehicle", "PER2效应": "原个体Ct未取得", "校正P": "0.7831（原Tukey）", "独立单位与限制": "成人ASM；30min脉冲；n4–6/条件"},
        {"来源／估计目标": "Bartman2021，DEX组24h−12h", "PER2效应": "原个体Ct未取得", "校正P": "0.0498（原Tukey）", "独立单位与限制": "时间差；不是药物×时间交互"},
        {"来源／估计目标": "GSE130715，24h BUD−vehicle", "PER2效应": "−0.934", "校正P": "0.000157（Holm44）", "独立单位与限制": "六成人配对；单晚期终点"},
    ]
    table3 = []
    for study in ["GSE342887", "GSE342771"]:
        for gene in ["PER2", "PER1"]:
            for contrast in ["CORT_alpha", "CORT_x_AF2"]:
                r = next(r for r in gr if r["study"] == study and r["symbol"] == gene and r["contrast"] == contrast)
                table3.append({"测量": "成熟RNA" if study == "GSE342887" else "新生RNA", "基因": gene,
                               "对比": "WT rescue药物效应" if contrast == "CORT_alpha" else "AF2 mutant×药物", "效应": num(r["effect"]),
                               "点态95%CI": f'{num(r["ci_low"])}至{num(r["ci_high"])}', "Holm55 P": pf(r["holm_55"])})
    table4 = []
    for r in human:
        if r["status"] == "estimable_selected_array":
            table4.append({"基因": r["symbol"], "暴露": f'{r["time_h"]}h', "供体": r["n_donors"], "蛋白配对效应": num(r["effect"]),
                           "点态95%CI": f'{num(r["ci_low"])}至{num(r["ci_high"])}', "Holm22 P": pf(r["holm_22"]), "精确Holm22 P": pf(r["exact_holm_22"])})
    tables = [("Table_1", table1), ("Table_2", table2), ("Table_3", table3), ("Table_4", table4)]
    for name, data in tables:
        csv_out(tab / (name + ".csv"), data)
        wb = openpyxl.Workbook(); sheet = wb.active; sheet.title = "数据"
        sheet.append(list(data[0]));
        for r in data: sheet.append(list(r.values()))
        wb.create_sheet("来源说明").append(["本报告原值来源见results.md、research/results/discovery_validation_v2/；CI点态，校正family逐项说明。"])
        wb.save(tab / (name + ".xlsx"))
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "pdf.fonttype": 42})

    def save(name, figure):
        figure.savefig(fig / (name + ".pdf"), bbox_inches="tight")
        figure.savefig(fig / (name + ".png"), bbox_inches="tight", dpi=180)
        plt.close(figure)

    f, ax = plt.subplots(figsize=(8, 3.5))
    data = [next(r for r in asm if r["study"] == "GSE130715" and r["symbol"] == g and r["contrast"] == "GC_control") for g in GENES]
    for i, r in enumerate(data):
        ax.errorbar(float(r["effect"]), i, xerr=[[float(r["effect"])-float(r["ci_low"])], [float(r["ci_high"])-float(r["effect"])]],
                    fmt="o", color="#227b86", capsize=3)
    ax.set_yticks(range(6), GENES); ax.invert_yaxis(); ax.axvline(0, color="grey", linestyle="--")
    ax.set_xlabel("Mean donor change in log2(TPM+1), pointwise 95% CI")
    ax.set_title("Six adult ASM donors: 24h budesonide response")
    f.tight_layout(); save("Figure_1_adult_ASM", f)
    f, axs = plt.subplots(1, 2, figsize=(9, 3.5), sharey=True)
    for ax, study, title in zip(axs, ["GSE342887", "GSE342771"], ["Mature RNA", "Nascent RNA"]):
        for i, (gene, contrast) in enumerate([(g, c) for g in ["PER2", "PER1"] for c in ["CORT_alpha", "CORT_x_AF2"]]):
            r = next(r for r in gr if r["study"] == study and r["symbol"] == gene and r["contrast"] == contrast)
            ax.errorbar(float(r["effect"]), i, xerr=[[float(r["effect"])-float(r["ci_low"])], [float(r["ci_high"])-float(r["effect"])]],
                        fmt="o", color="#bb7150" if float(r["holm_55"]) < .05 else "#7b8b95", capsize=3)
        ax.set_title(title); ax.axvline(0, color="grey", linestyle="--"); ax.set_xlabel("Exact WLS logCPM contrast")
    axs[0].set_yticks(range(4), ["PER2 CORT WT", "PER2 CORT x AF2", "PER1 CORT WT", "PER1 CORT x AF2"]); axs[0].invert_yaxis()
    f.suptitle("GR KO-rescue variants: PER2 unsupported; PER1 control responds", fontsize=11)
    f.tight_layout(); save("Figure_2_GR_genetics", f)
    donor_data = rows(source / "human_protein/human_selected_donor_effects.csv")
    f, axs = plt.subplots(1, 2, figsize=(8.5, 3.5))
    for ax, gene in zip(axs, ["METTL7A", "FKBP5"]):
        for x, time in enumerate(["24", "72"]):
            values = [float(r["effect"]) for r in donor_data if r["symbol"] == gene and r["time_h"] == time]
            ax.scatter(x + np.linspace(-.10, .10, len(values)), values, color="#227b86", alpha=.85)
            r = next(r for r in human if r["symbol"] == gene and r["time_h"] == time)
            ax.errorbar(x, float(r["effect"]), yerr=[[float(r["effect"])-float(r["ci_low"])], [float(r["ci_high"])-float(r["effect"])]], fmt="s", color="#c56c46", capsize=4)
        ax.set_xticks([0, 1], ["24h", "72h"]); ax.set_title(gene); ax.axhline(0, color="grey", linestyle="--")
        ax.set_ylabel("Within-donor log2 protein change")
    f.suptitle("Human peritubular cells: 6 donors per time, selected arrays", fontsize=11)
    f.tight_layout(); save("Figure_3_human_protein", f)
    title = "公开数据验证增补：成人ASM响应、GR遗传交叉检验与人源蛋白证据"
    doc = Document()
    for style in ["Normal", "Title", "Heading 1", "Heading 2"]:
        doc.styles[style].font.name = "Noto Sans CJK SC"
        doc.styles[style]._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    doc.styles["Normal"].font.size = Pt(10)
    section = doc.sections[0]; section.top_margin = Inches(.65); section.bottom_margin = Inches(.6)
    section.left_margin = Inches(.6); section.right_margin = Inches(.6)
    md = ["# " + title, "2026-10-09｜结果后探索｜全部采用公开数据｜新增3个ASM表达实验、GR两类RNA及蛋白数据核对"]
    h = ["<h1>" + title + "</h1><p>2026-10-09 · 公开数据验证增补</p>"]
    doc.add_heading(title, 0); doc.add_paragraph(md[1])

    def heading(text, newpage=False):
        if newpage: doc.add_page_break()
        doc.add_heading(text, 1); md.append("\n## " + text); h.append("<h2>" + text + "</h2>")

    def para(text):
        doc.add_paragraph(text); md.append(text); h.append("<p>" + html.escape(text) + "</p>")

    def table(data, caption):
        para(caption)
        t = doc.add_table(rows=1, cols=len(data[0])); t.style = "Table Grid"
        for c, k in zip(t.rows[0].cells, data[0]): c.text = k
        for row in data:
            for c, v in zip(t.add_row().cells, row.values()): c.text = str(v)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs: run.font.size = Pt(7.5)
        md.extend(["|" + "|".join(data[0]) + "|", "|" + "|".join(["---"]*len(data[0])) + "|"])
        md.extend("|" + "|".join(str(v) for v in r.values()) + "|" for r in data)
        h.append("<div class='table-wrap'><table><thead><tr>" + "".join("<th>" + html.escape(k) + "</th>" for k in data[0]) +
                 "</tr></thead><tbody>" + "".join("<tr>" + "".join("<td>" + html.escape(str(v)) + "</td>" for v in r.values()) + "</tr>" for r in data) + "</tbody></table></div>")

    def figure(name, caption, width=6.8):
        doc.add_picture(str(fig / (name + ".png")), width=Inches(width)); para(caption)
        md.append(f"![{caption}](figures/{name}.png)")
        img = base64.b64encode((fig / (name + ".png")).read_bytes()).decode()
        h.append(f"<figure><img alt='{html.escape(caption)}' src='data:image/png;base64,{img}'></figure>")

    heading("1 本轮结论")
    para("完成了新增公开数据分析，而不仅是方案：GSE130715六名成人供体、GSE13168三名可配对供体及GSE135730早期ASM培养实验；另分析GSE342887/342771的GR敲除背景下野生型与突变体rescue，以及PXD081613蛋白组和人源HTPC蛋白补充表。")
    para("PER2晚期下调在新增六成人ASM供体中重复；早期ASM诱导、统一成人逐时方向反转及PER2受体遗传依赖尚未得到确认。新的蛋白层面线索集中在METTL7A：成人睾丸管周平滑肌样细胞的24/72h DEX处理均6/6供体上调，但这属于跨组织、原文选中数组的探索性支持。")
    para("六重点候选在新增成人ASM的24h BUD响应全部方向一致并通过配对t的44项Holm；同供体计数模型亦支持。精确小样本符号翻转校正未确认，完整保留。所有GNAS/PKI交互均未通过各自Holm，不能据此声称受体机制或等效。")
    heading("2 资料与检索范围")
    para(f"初始4个GEO查询得到{len(rows(source / 'geo_series.csv'))}个唯一series，6个EuropePMC查询完整分页保存；扩大受体遗传/蛋白检索又保存{len(rows(source / 'cross_geo_series.csv'))}个唯一series汇总及{len(rows(source / 'cross_literature.csv'))}条文献命中记录。检索只是范围限定的可用性检查，不证明其他资料不存在。原计数/TPM/SOFT/RPKM、原文HTML/XML、PRIDE数据和失败响应均保存。")
    para("成人时间证据来自Bartman等2021[1]。六新ASM供体年龄19–54岁由Nuñez等2020表1与GEO双源确认[2]；六label与旧GSE94335已核对label无exact重叠，但未知别名未完全排除，不累加供体数。其他匿名培养系或年龄未公开的数据只写实验重复。")
    para("本文是公开数据再分析；所有候选基于既知v0.11主线固定，新增比较方案在各新表达读取前冻结，人体蛋白选中数组重算及ortholog映射增补明确标记结果后探索。统计完成范围与待确认机制分开记录，负责人无需新增湿实验。")
    heading("3 新增成人ASM晚期响应", True)
    para("GSE130715包含六成人×GNAS siRNA/对照×1µM BUD/vehicle，共24文库、28,395基因。配对主估计量是log2(TPM+1)差；11候选/控制×4比较，44项Holm。供体是n=6，文库不是独立人。新TMM/voom过滤保留14,692基因，~donor+kd*gc，秩9、残差df15；精确逐基因WLS作为敏感性。")
    table(table1, "表1 六重点基因的新成人ASM响应。两列效应单位不同：TPM对数差与加权logCPM；同源两模型不增加独立证据数。")
    figure("Figure_1_adult_ASM", "图1 六成人逐供体差的均值与点态95%CI。所有六候选均6/6方向一致；精确符号翻转原P=.03125、44项Holm=1。")
    para("PER2 −0.934（CI −1.044至−0.824，Holm=.000157）；GNAS×BUD交互+0.224（CI −0.166至+0.615，Holm=1）。原计数WLS的PER2响应−1.405、Holm=3.54e−8，交互+0.377、Holm=1。敲低靶点GNAS RNA −2.462、6/6降低，表明干预发生，但不能把GNAS改称GR。")
    heading("4 成人时间序列与早期ASM反证", True)
    para("补回原文中真正的成人ASM部分：3女/3男、33–77岁来源池；DEX100nM或vehicle脉冲30min后洗去，12/16/20/24h匹配采样[1]。作者文字描述PER2升高，但表2只有DEX组24h−12h为Tukey P=.0498；同24h DEX−vehicle为.7831。时间变化不等于药物×时间交互；没有原个体Ct/协方差，未重算新效应、CI或交互。")
    table(table2, "表2 PER2时间条件核对。不同研究药物、剂量、培养、同步化和测量尺度不同，逐行读，不能拼接为同一供体时间曲线。")
    para("GSE13168 8h FP的PER2 +0.691、3/3↑，原P=.0426但44项Holm=1；PKI×FP −0.439、Holm=1。GSE135730 4h DEX的PER2 +0.111、Holm33=1，IL17背景+0.196、Holm=.928，均未确认早期诱导。该数据human标签与mm9摘要冲突保留，使用人Ensembl标识，不猜坐标、不猜供体配对。")
    para("早期ASM数据对METTL7A和MAOA更有支持：无IL17时分别+2.014（Holm33=.000266）、+1.586（.00925）；IL17背景也通过校正，直接DEX×IL17交互未通过。当前PER2结论收束为可重复的成人晚期抑制与条件相关响应，不能写成成人ASM已证实先升后降。")
    heading("5 GR遗传交叉验证", True)
    para("2026年公开预印本对应U2OS GR-KO背景的WTα、AF2/E755R与IDRΔ1–398 rescue构建体，100nM corticosterone90min；mRNA18文库、新生TT17文库。各自全基因TMM/voom，再提取11基因×5比较＝55项Holm；精确基因权重协方差。只有rescue变体，没有empty-KO转录组组，不能称为WT rescue相对空KO效应。两测量共用平行样本[4]。")
    table(table3, "表3 PER2与PER1控制的受体遗传交叉检查。AF2×药物为(药物−vehicle)AF2−(药物−vehicle)WT。")
    figure("Figure_2_GR_genetics", "图2 受体变体直接交互及点态CI；橙色为Holm55通过，灰色未通过。PER1阳性控制不替代PER2。")
    para("PER2成熟RNA WT效应+.177、Holm=1，AF2交互−.337、Holm=1；新生RNA WT+.313、Holm=.482，AF2交互−.424、Holm=.348。PER1成熟RNA WT+.782、Holm=2.53e−9，AF2交互−.776、Holm=9.05e−6；新生RNA交互−1.569、Holm=.00383。PER2受体变体依赖未被该独立背景确认。")
    heading("6 人源蛋白层面与供体重算", True)
    para("Stepanov等2022[5]的TableS2在出版社静态接口取得并核对ZIP CRC。HTPC来自29–55岁男性，六供体、每供体/条件3技术重复，DEX1µM24/72h。完整检测列表只含总intensity；逐样本LFQ只公开于原文预筛差异列表。六候选中可重算蛋白数组只有METTL7A，控制为FKBP5；其他缺失数组不当阴性。")
    table(table4, "表4 人源HTPC选中数组的供体层面重算。先平均三技术重复的log2LFQ，再每供体DEX−control；n=6。11基因×2时点Holm22保留缺测占位，选中条件与探索边界仍存在。")
    figure("Figure_3_human_protein", "图3 每点是一名供体，方块/误差线是均值/点态CI；两时点供体组成不完全相同，不按位置连线。")
    para("METTL7A24h +0.921（Holm22=.0437）、72h+1.504（.00261），均6/6↑；原文效应相同但原P明显更小，故本轮使用供体层面P。精确符号翻转Holm22=.688，两时点不作稳健小样本确认。五共享供体的72h−24h响应差+.688、Holm11=.412，未确认蛋白进一步增幅。")
    heading("7 蛋白缺测、功能与结论边界", True)
    para("PXD081613完整7,274行蛋白表经72段HTTP Range取得。PRIDE写human，搜索协议却写Mus musculus，7,218条实际mouse注释；保留此冲突，另作mouse注释支持性计算。PER2/PER3/PER1等未在固定规则下检出；MAOA/TMT1A无Holm44通过，GR/Nr3c1仅phase31/3h为−.185、Holm=.0211。既不是人ASM蛋白验证，也不能把PER2缺测写成无变化。")
    para("真实ASM GRα遗传/蛋白/功能文献确实存在[3]：GRα siRNA后蛋白接近消失，但rapid cAMP对BUD的响应差异很小，n4–5培养系。没有PER2测量或rescue；原文仍讨论残余GR解释，不证明GR等效无作用。Bartman[1]的Ca²⁺/SOCE干预是胎儿ASM PER1/ARNTL，不是成人PER2。")
    para("人源HTPC原文[5]还报告GR siRNA、GR蛋白敲低及FKBP5/IL6响应受阻、IL6 ELISA；这些为跨组织受体/功能先例，不能归因于METTL7A或PER2。FSTL3/ANGPTL2在部分secretome检测到，只能证实可测，不能推导激素响应或功能联系。")
    para("综合结论：①新增成人ASM支持PER2/PER3晚期抑制及六候选方向重复；②早期ASM诱导与成人逐时反转未确认；③GR变体遗传数据未确认PER2，保留PER1对照；④METTL7A增加了跨组织人源蛋白支持；⑤候选成人ASM GR遗传因果、蛋白及Ca²⁺/收缩功能仍是证据缺口。在无法新增湿实验的约束下，本轮已完成可取得数据的验证，不将缺口改写为成功。")
    heading("8 数值复核与可复现结果")
    para("R从原TPM/SOFT核对396基因测量及88检验、原RPKM核对132测量/33检验；Python从原GR counts核对385测量及550WLS数值；R从mouse蛋白表核对160测量/40可测检验，人蛋白选中表核对144测量/24供体响应/6可测统计。CSV/XLSX、原数组及报告来源另有封存审核。TMM/voom/moderation依赖原R拟合，未称全模型跨语言重新拟合。")
    para("BioExplainR0.4.0.9000实际源码载入，8context项目、双语离线HTML和事实计数核对保存，在线模型调用0。完整数据/方案/脚本/失败在research/results/discovery_validation_v2/；复现命令见本目录README。")
    heading("参考资料", True)
    refs = [
        "[1] Bartman CM et al. Cellular clocks in hyperoxia effects on [Ca2+]i regulation in developing human airway smooth muscle. 2021. doi:10.1152/ajplung.00406.2020. PMC8294620（成人部分及表2全文已核对）。",
        "[2] Nuñez FJ et al. Glucocorticoids rapidly activate cAMP production via Gαs to initiate non-genomic signaling... 2020. doi:10.1096/fj.201902521R. PMC7027561；GSE130715。",
        "[3] Koziol-White C et al. Budesonide enhances agonist-induced bronchodilation in human small airways... 2020. doi:10.1152/ajplung.00393.2019. PMC7052668。",
        "[4] Edmondson A et al. Post-transcriptional glucocorticoid receptor signalling synchronises circadian rhythms. bioRxiv 2026，v2（未同行评审）. doi:10.64898/2026.08.12.743743；GSE342887、GSE342771、PXD081613。全文403，只核对公开数据与API元数据。",
        "[5] Stepanov YK et al. Profound Effects of Dexamethasone on the Immunological State, Synthesis and Secretion Capacity of Human Testicular Peritubular Cells. Cells 2022;11:3164. doi:10.3390/cells11193164. PMC9562650；TableS2，PXD033514/033534。",
        "[6] Misior AM et al. Glucocorticoid- and protein kinase A-dependent transcriptome regulation in airway smooth muscle. 2009. PMC2701960；GSE13168。",
        "[7] GSE135730 GEO原发布记录，2020公开；4hASM IL17A/DEX RPKM，GPL10999；供体年龄未披露，human/mm9摘要冲突保留。",
    ]
    for ref in refs: para(ref)
    paths = ["expression_results.csv", "voom_exact_results.csv", "published_adult_ASM_time_tests.csv", "asm_early/early_results.csv",
             "gr_public/gr_results.csv", "human_protein/human_selected_results.csv", "protein_public/protein_results.csv"]
    for p in paths: para("数值入口：research/results/discovery_validation_v2/" + p)
    basename = "BioExplainR-公开数据验证-结果增补-v1"
    docx = dest / (basename + ".docx"); doc.save(docx)
    (dest / "results.md").write_text("\n\n".join(md) + "\n")
    css = "body{font-family:system-ui,'Noto Sans CJK SC',sans-serif;max-width:1080px;margin:2em auto;padding:0 1em;line-height:1.7;color:#26343b}h1,h2{color:#236477}table{border-collapse:collapse;width:100%;font-size:.85em}td,th{border:1px solid #ccd;padding:.4em;text-align:left}th{background:#edf4f6}.table-wrap{overflow-x:auto}img{max-width:100%;height:auto}p{overflow-wrap:anywhere}figure{margin:1em 0}"
    (dest / (basename + ".html")).write_text("<!doctype html><html lang='zh-CN'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>" + title + "</title><style>" + css + "</style><body>" + "\n".join(h) + "</body></html>")
    cmd = ["libreoffice", "-env:UserInstallation=file:///tmp/opencode/validation-report-lo", "--headless", "--convert-to", "pdf", "--outdir", str(dest), str(docx)]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    (dest / "export_log.txt").write_text(proc.stdout + proc.stderr)
    assert proc.returncode == 0 and docx.with_suffix(".pdf").exists()
    with fitz.open(docx.with_suffix(".pdf")) as pdf:
        pages = len(pdf)
    json_out(dest / "export.json", {"status": "completed", "utc": utc(), "pages": pages, "tables": 4, "figures": 3,
             "command": cmd, "returncode": proc.returncode, "pdf_figure_mode": "PNG previews from Word, separate original vector PDFs provided",
             "source": str(source.relative_to(ROOT)), "script": record(Path(__file__))})
    print("Chinese validation report exported", pages, "pages", flush=True)


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__); p.add_argument("source", type=Path); p.add_argument("destination", type=Path)
    a = p.parse_args(); build(a.source.resolve(), a.destination.resolve())
