"""Extract complete detection lists and selected published protein statistics honestly."""
import csv
import json
import re
import sys
from pathlib import Path
import openpyxl
from validation_sources import csv_out,json_out,record,utc

dest=Path(sys.argv[1]).resolve()/"human_protein"
json_out(dest/"workbook_scope_amendment.json",{
    "utc":utc(),"script":record(Path(__file__)),
    "reason":"ALL sheets expose total intensity, not per-sample arrays; per-sample LFQ and statistics exist only in preselected differential sheets.",
    "analysis":"Complete frozen candidate detection inventory; extract published effects/p/q only when candidate occurs in selected sheet. No new unselected family test or missing P=1 significance claim.",
    "selection_limit":"No-hit in differential list is unavailable effect/test, not a null result. Published significance conditional on original all-protein processing; selected re-tests not independent replication."})
genes=["PER2","PER3","METTL7A","MAOA","ANGPTL2","FSTL3","PER1","FKBP5","TSC22D3","DUSP1","NR3C1"]
w=openpyxl.load_workbook(dest/"Table S2.xlsx",data_only=True,read_only=True)
detected=[]; selected=[];headers=[]
all_sheets=["S1_24hProteome_ALL","S3_72hProteome_All","S5_24hSecretome_All","S7_72hSecretome_All"]
diff_sheets=["S2_24hProteome_DiffAbundant","S4_72hProteome_DiffAbundant","S6_24hSecretome_DiffAbundant","S8_72hSecretome_DiffAbundant"]
for all_name,diff_name in zip(all_sheets,diff_sheets):
    data=list(w[all_name].iter_rows(values_only=True));count=len(data)-1
    for gene in genes:
        aliases={gene,"TMT1A"} if gene=="METTL7A" else {gene}
        hit=[r for r in data[1:] if isinstance(r[1],str) and set(r[1].split(";"))&aliases]
        detected.append({"assay":all_name,"symbol":gene,"status":"detected_in_all_list" if hit else "not_detected_in_all_list",
                         "protein_groups":"|".join(str(r[0]) for r in hit),"all_proteins":count,
                         "total_intensity":"|".join(str(r[3]) for r in hit),"meaning":"Detection inventory only, not DEX effect"})
    data=list(w[diff_name].iter_rows(values_only=True))
    # Preserve full selected table as a CSV so indices/header structure are auditable.
    columns=[f"column_{i+1}" for i in range(w[diff_name].max_column)]
    csv_out(dest/(diff_name+".csv"),[{k:"" if v is None else v for k,v in zip(columns,r)} for r in data],columns)
    for i,r in enumerate(data[:6],1):
        headers.append({"sheet":diff_name,"row":i,"cells":" | ".join(f"{j+1}:{v}" for j,v in enumerate(r) if v is not None)})
        print(diff_name,i,headers[-1]["cells"][:3500],flush=True)
    for i,r in enumerate(data,1):
        for gene in genes:
            aliases={gene,"TMT1A"} if gene=="METTL7A" else {gene}
            if any(isinstance(v,str) and set(v.split(";"))&aliases for v in r):
                selected.append({"assay":diff_name,"symbol":gene,"original_row":i,
                                 "cells":" | ".join(f"{j+1}:{v}" for j,v in enumerate(r) if v is not None),
                                 "scope":"Published preselected differential list, not new unselected test"})
                print("SELECTED",diff_name,gene,i,selected[-1]["cells"][:1500],flush=True)
csv_out(dest/"human_protein_detection.csv",detected)
csv_out(dest/"human_protein_selected_rows.csv",selected,["assay","symbol","original_row","cells","scope"])
csv_out(dest/"workbook_headers.csv",headers)
json_out(dest/"human_protein_inventory_run.json",{"status":"completed","utc":utc(),"complete_detection_tests":len(detected),
         "selected_rows":len(selected),"scope":"Human cross-tissue protein detection and published selected results; no candidate-wide raw quantitative re-test"})
