"""Preserve UniProt secretion annotations and compare complete published inventories."""
import argparse
import csv
import json
import re
from pathlib import Path

import openpyxl

from secretome_analyze import HUMAN, sources
from secretome_acquire import PRIOR
from set_validation import rows
from validation_sources import record, utc, fetch, csv_out, json_out


def acquire(out):
    requests = [("ANGPTL2","Q9UKU9"),("FSTL3","O95633")]
    json_out(out/"annotation_source_freeze.json",dict(utc=utc(),script=record(Path(__file__)),
        requests=requests,scope="Canonical protein secretion/function annotations, not candidate drug regulation or functional validation"))
    logs = []
    for gene,acc in requests:
        log = fetch(f"https://rest.uniprot.org/uniprotkb/{acc}.json",out/f"sources/UniProt_{gene}.json")
        logs.append(log)
    json_out(out/"annotation_acquisition.json",dict(utc=utc(),records=logs))
    assert all(r["status"] == "retrieved" for r in logs)


def compare(out):
    book = openpyxl.load_workbook(HUMAN/"Table S2.xlsx",data_only=True,read_only=True)
    qc = rows(out/"group_QC.csv")
    published, selected, cellchecks = [],[],[]
    summaries = []
    for time,all_sheet,diff_sheet in [(24,"S5_24hSecretome_All","S6_24hSecretome_DiffAbundant"),
                                      (72,"S7_72hSecretome_All","S8_72hSecretome_DiffAbundant")]:
        local = [r for r in qc if int(r["time_h"]) == time]
        raw = {r["group_id"]:r for r in local}
        original = sources(out,time)
        with original["proteinGroups.txt"].open(newline="") as stream:
            table = list(csv.DictReader(stream,delimiter="\t"))
        data = list(book[all_sheet].iter_rows(values_only=True))
        for row_num,r in enumerate(data[1:],2):
            exact = [g for g in local if g["protein_ids"] == r[0]]
            tokens = [g for g in local if set(g["protein_ids"].split(";"))&set(str(r[0]).split(";"))]
            published.append(dict(time_h=time,source_sheet=all_sheet,source_row=row_num,published_protein_ids=r[0],
                published_gene_names=r[1] or "",published_total_intensity=r[3],exact_group_matches=len(exact),
                token_group_matches=len(tokens),raw_group_ids="|".join(g["group_id"] for g in (exact or tokens)),
                raw_QC="|".join(g["QC_status"] for g in (exact or tokens))))
        data = list(book[diff_sheet].iter_rows(values_only=True))
        header_index = next(i for i,r in enumerate(data[:5]) if r[0] == "Protein Name")
        cols = []
        for i,v in enumerate(data[header_index]):
            if isinstance(v,str):
                m = re.fullmatch(r"(\d{2}\.\d{2}\.\d{2})_?(co|dex)([123])",v)
                if m:
                    cols.append((i,m[1],m[2],int(m[3])))
        assert len(cols) == 36
        samplemap = [s for s in rows(out/"sample_map.csv") if int(s["time_h"]) == time]
        for rownum,r in enumerate(data[header_index+1:],header_index+2):
            if not r[2]:
                continue
            exact = [g for g in local if g["protein_ids"] == r[7]]
            selected.append(dict(time_h=time,source_sheet=diff_sheet,source_row=rownum,published_gene_names=r[1],
                published_protein_ids=r[7],published_effect=r[5],published_pvalue=r[3],published_qvalue=r[4],
                exact_group_matches=len(exact),raw_group_ids="|".join(g["group_id"] for g in exact)))
            if len(exact) != 1:
                continue
            group = exact[0]
            source = table[int(group["source_row"])-2]
            for col,donor,condition,rep in cols:
                sample = [s for s in samplemap if s["donor"] == donor and s["condition"] == condition and int(s["technical_rep"]) == rep]
                assert len(sample) == 1
                lfq = float(source[sample[0]["source_column"]])
                import math
                log2 = math.log2(lfq) if lfq>0 else None
                value = float(r[col]) if r[col] is not None else None
                match = log2 is not None and value is not None and abs(log2-value)<=.000051
                cellchecks.append(dict(time_h=time,symbol=r[1],group_id=group["group_id"],donor=donor,condition=condition,
                    technical_rep=rep,raw_LFQ=lfq,raw_log2=log2 if log2 is not None else "",published_log2=value if value is not None else "",
                    status="matched_rounding" if match else "raw_zero_published_finite" if lfq == 0 and value is not None else "different_or_missing"))
        summaries.append(dict(time_h=time,published_all_rows=sum(r["time_h"] == time for r in published),
            all_exact_matches=sum(r["time_h"] == time and r["exact_group_matches"] == 1 for r in published),
            selected_rows=sum(r["time_h"] == time for r in selected),compared_cells=sum(r["time_h"] == time for r in cellchecks),
            matched_rounding=sum(r["time_h"] == time and r["status"] == "matched_rounding" for r in cellchecks),
            raw_zero_published_finite=sum(r["time_h"] == time and r["status"] == "raw_zero_published_finite" for r in cellchecks)))
    annotations = []
    for gene in ["ANGPTL2","FSTL3"]:
        path = out/f"sources/UniProt_{gene}.json"
        data = json.loads(path.read_text())
        for i,comment in enumerate(data.get("comments",[])):
            if comment["commentType"] in ["SUBCELLULAR LOCATION","FUNCTION"]:
                annotations.append(dict(symbol=gene,accession=data["primaryAccession"],comment_type=comment["commentType"],
                    comment_json=json.dumps(comment,ensure_ascii=False),source_path=record(path)["path"]))
        for feature in data.get("features",[]):
            if feature["type"] == "Signal":
                annotations.append(dict(symbol=gene,accession=data["primaryAccession"],comment_type="SIGNAL PEPTIDE",
                    comment_json=json.dumps(feature,ensure_ascii=False),source_path=record(path)["path"]))
    csv_out(out/"published_inventory_comparison.csv",published)
    csv_out(out/"published_selected_results.csv",selected)
    csv_out(out/"published_selected_cell_comparison.csv",cellchecks)
    csv_out(out/"secretion_annotations.csv",annotations)
    asm = rows(PRIOR/"tpm_results.csv")
    candidates = rows(out/"candidate_results.csv")
    cross = []
    for r in candidates:
        a = [a for a in asm if a["feature_id"] == r["feature_id"] and a["contrast"] == "GC_control"]
        measured = r["status"] == "estimable"
        direction = r["discovery_direction"]
        agrees = bool(measured and direction and float(r["effect"])*int(direction)>0)
        cross.append(dict(feature_id=r["feature_id"],symbol=r["symbol"],time_h=r["time_h"],
            ASM_discovery_direction=direction,adult_ASM_TPM_effect=a[0]["effect"] if a else "",
            adult_ASM_Holm54=a[0]["holm"] if a else "",secretome_status=r["status"],secretome_n_pairs=r["n_pairs"],
            secretome_effect=r["effect"],secretome_Holm112=r["holm112"],direction_concordance="same" if agrees else "opposite" if measured and direction else "unavailable_or_control",
            same_direction_corrected_support=measured and agrees and float(r["holm112"])<.05,
            scope="ASM BUD RNA vs HTPC DEX extracellular LFQ; cross-tissue/drug/dose/scale, not matched multiomics or secretion rate"))
    csv_out(out/"candidate_cross_evidence.csv",cross)
    json_out(out/"published_comparison_run.json",dict(status="completed",utc=utc(),script=record(Path(__file__)),summary=summaries,
        scope="Published same-source lists, full exact/token inventory and selected-cell comparison; no-hit not null; UniProt not drug-response validation"))
    print(json.dumps(summaries,indent=2))


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("action",choices=["acquire","compare"])
    p.add_argument("output",type=Path)
    a = p.parse_args()
    globals()[a.action](a.output.resolve())
