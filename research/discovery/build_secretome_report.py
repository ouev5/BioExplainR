"""Build a traceable Chinese secretome report with complete tables and original figures."""
import argparse
import base64
import html
import json
import shutil
import subprocess
from pathlib import Path

import fitz
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from set_validation import rows
from validation_sources import ROOT, record, utc, csv_out, json_out

NAME = "PXD033504-完整分泌蛋白组-结果报告"


def number(v, digits=3):
    return f"{float(v):.{digits}f}" if v != "" else "不可估计"


def pv(v):
    return f"{float(v):.3g}"


def build(source, dest):
    dest.mkdir(exist_ok=False)
    fig, tab = dest/"figures", dest/"tables"
    fig.mkdir(); tab.mkdir()
    candidates = rows(source/"candidate_results.csv")
    full = rows(source/"all_results.csv")
    effects = rows(source/"candidate_donor_effects.csv")
    sensitivity = rows(source/"candidate_sensitivity.csv")
    summary = rows(source/"array_summary.csv")
    detection = rows(source/"candidate_detection.csv")
    direct = rows(source/"candidate_direct_time.csv")
    versions = rows(source/"processing_versions.csv")
    sample_scope = rows(source/"sample_scope_comparison.csv")
    comparable = json.loads((source/"published_comparison_run.json").read_text())
    diagnostics = json.loads((source/"quantification_diagnostics_run.json").read_text())
    audit = json.loads((source/"r_numeric_audit.json").read_text())
    workflow = json.loads((source/"bioexplainr_workflow.json").read_text())
    assert audit["status"] == workflow["status"] == "pass"
    table1 = [{"时点":r["time_h"]+"h","原组/QC组":r["raw_groups"]+"/"+r["QC_groups"],
               "标签/技术列":r["donor_labels"]+"/"+r["technical_columns"],"全组可估计":r["estimable_groups"],
               "统一BH通过":r["pooled_BH_significant"],"候选可估计/56":r["candidate_estimable"]+"/56"} for r in summary]
    focus = [r for r in candidates if r["status"] == "estimable" or r["symbol"] in ["ANGPTL2","FSTL3"]]
    table2 = [{"基因":r["symbol"],"时点/n":r["time_h"]+"h/"+r["n_pairs"],"log₂ LFQ差":number(r["effect"]),
               "点态95%CI":number(r["ci_low"])+"至"+number(r["ci_high"]),"Holm112 P":pv(r["holm112"]),
               "↑/↓":r["positive_donors"]+"/"+r["negative_donors"] if r["status"] == "estimable" else r["status"]} for r in focus]
    table3 = []
    for gene in ["ANGPTL2","FSTL3","FBN2"]:
        for time in ["24","72"]:
            for kind in ["published_six","strict_three"]:
                r = next(r for r in sensitivity if r["symbol"] == gene and r["time_h"] == time and r["analysis"] == kind)
                table3.append({"基因/时点":gene+"/"+time+"h","方案":"发表六标签" if kind == "published_six" else "严格3/3技术测量",
                               "配对n":r["n_pairs"],"效应":number(r["effect"]),"Holm112 P":pv(r["holm112"]),"状态":r["status"]})
    for name,data in [("Table_1_array_scope",table1),("Table_2_candidate_results",table2),("Table_3_sensitivity",table3)]:
        csv_out(tab/(name+".csv"),data)
    copied = ["all_results.csv","group_QC.csv","sample_map.csv","candidate_results.csv","candidate_detection.csv",
              "candidate_measurements.csv","candidate_donor_effects.csv","candidate_sensitivity.csv","candidate_leave_one.csv",
              "candidate_direct_time.csv","candidate_direct_donor_effects.csv","published_inventory_comparison.csv",
              "published_selected_results.csv","published_selected_cell_comparison.csv","secretion_annotations.csv"]
    copied.append("candidate_cross_evidence.csv")
    copied += ["sample_scope_comparison.csv","raw_file_scope.csv","candidate_missingness_summary.csv","all_BH_significant.csv"]
    copied += ["candidate_identification.csv","candidate_raw_intensity_diagnostics.csv","published_scale_diagnostics.csv","published_total_intensity_comparison.csv"]
    for name in copied:
        shutil.copyfile(source/name,tab/name)
    for path in sorted(tab.glob("*.csv")):
        data = rows(path)
        book = openpyxl.Workbook(); sheet = book.active; sheet.title = "数据"
        sheet.append(list(data[0]))
        for r in data:
            sheet.append(list(r.values()))
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        book.create_sheet("来源").append([str(source.relative_to(ROOT)),"公开数据探索重分析；配对日期标签；LFQ0缺测；完整分母；点态CI"])
        book.save(path.with_suffix(".xlsx"))
    plt.rcParams.update({"font.family":"DejaVu Sans","font.size":9,"pdf.fonttype":42})
    f,axs = plt.subplots(2,2,figsize=(9,6.3))
    for ax,gene in zip(axs.flat,["ANGPTL2","FSTL3","FBN2","PTX3"]):
        for x,time in enumerate(["24","72"]):
            r = next(r for r in candidates if r["symbol"] == gene and r["time_h"] == time)
            vals = [float(v["effect"]) for v in effects if v["symbol"] == gene and v["time_h"] == time and v["status"] == "paired"]
            if vals:
                ax.scatter(x+np.linspace(-.10,.10,len(vals)),vals,color="#277e85",s=30)
            if r["status"] == "estimable":
                ax.errorbar(x,float(r["effect"]),yerr=[[float(r["effect"])-float(r["ci_low"])],[float(r["ci_high"])-float(r["effect"])]],
                            fmt="s",color="#bd6846",capsize=4)
            else:
                ax.text(x,.94,"Not estimable\nn="+r["n_pairs"],ha="center",va="top",transform=ax.get_xaxis_transform())
        ax.axhline(0,color="grey",ls="--"); ax.set_title(gene); ax.set_xticks([0,1],["24h","72h"])
        ax.set_ylabel("DEX-control mean log2 LFQ difference")
    f.suptitle("HTPC extracellular abundance: paired biopsy-date labels")
    f.tight_layout(); f.savefig(fig/"Figure_1_candidates.pdf",bbox_inches="tight"); f.savefig(fig/"Figure_1_candidates.png",dpi=180,bbox_inches="tight"); plt.close(f)
    f,axs = plt.subplots(1,2,figsize=(9,3.5),sharey=True)
    for ax,time in zip(axs,["24","72"]):
        z = [r for r in full if r["time_h"] == time and r["status"] == "estimable"]
        ax.scatter([float(r["effect"]) for r in z],[-np.log10(max(float(r["pooled_BH"]),1e-300)) for r in z],
                   c=["#bd6846" if float(r["pooled_BH"]) < .05 else "#a8b3bc" for r in z],s=11,alpha=.75)
        for gene in ["ANGPTL2","FSTL3","MMP19","PTX3","FBN2","CRISPLD2"]:
            hit = [r for r in z if r["gene_names"] == gene]
            for r in hit:
                ax.annotate(gene,(float(r["effect"]),-np.log10(max(float(r["pooled_BH"]),1e-300))),fontsize=7,xytext=(3,4),textcoords="offset points")
        ax.axhline(-np.log10(.05),color="grey",ls="--"); ax.axvline(0,color="grey",ls=":")
        ax.set_title(time+"h, "+str(len(z))+" estimable groups"); ax.set_xlabel("Mean paired log2 LFQ difference")
    axs[0].set_ylabel("-log10 pooled BH adjusted P")
    f.suptitle(str(len(full))+" QC protein groups, one pooled BH family; no imputation")
    f.tight_layout(); f.savefig(fig/"Figure_2_complete_array.pdf",bbox_inches="tight"); f.savefig(fig/"Figure_2_complete_array.png",dpi=180,bbox_inches="tight"); plt.close(f)
    doc = Document()
    for style in ["Normal","Title","Heading 1","Heading 2"]:
        doc.styles[style].font.name = "Noto Sans CJK SC"
        doc.styles[style]._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Noto Sans CJK SC")
    doc.styles["Normal"].font.size = Pt(10)
    sec = doc.sections[0]; sec.top_margin=sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.6)
    title = "PXD033504完整分泌蛋白组：ANGPTL2／FSTL3及固定候选的分泌层面证据"
    md,h = ["# "+title], ["<h1>"+title+"</h1>"]
    doc.add_heading(title,0)
    def heading(text,newpage=False):
        if newpage: doc.add_page_break()
        doc.add_heading(text,1); md.append("\n## "+text); h.append("<h2>"+html.escape(text)+"</h2>")
    def para(text):
        doc.add_paragraph(text); md.append(text); h.append("<p>"+html.escape(text)+"</p>")
    def addtable(data,caption):
        para(caption); t=doc.add_table(rows=1,cols=len(data[0])); t.style="Table Grid"
        for c,k in zip(t.rows[0].cells,data[0]): c.text=k
        for r in data:
            for c,v in zip(t.add_row().cells,r.values()): c.text=str(v)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs: run.font.size=Pt(7.5)
        md.extend(["|"+"|".join(data[0])+"|","|"+"|".join(["---"]*len(data[0]))+"|"])
        md.extend("|"+"|".join(str(v) for v in r.values())+"|" for r in data)
        h.append("<div class='table-wrap'><table><tr>"+"".join("<th>"+html.escape(k)+"</th>" for k in data[0])+"</tr>"+
                 "".join("<tr>"+"".join("<td>"+html.escape(str(v))+"</td>" for v in r.values())+"</tr>" for r in data)+"</table></div>")
    def figure(name,caption):
        doc.add_picture(str(fig/(name+".png")),width=Inches(6.8)); para(caption)
        md.append(f"![{caption}](figures/{name}.png)")
        h.append("<figure><img alt='"+html.escape(caption)+"' src='data:image/png;base64,"+base64.b64encode((fig/(name+".png")).read_bytes()).decode()+"'></figure>")
    para("2026-10-10｜公开数据结果后探索｜固定56基因×2时点｜BioExplainR 0.4.0.9000")
    heading("1 完成范围与独立单位")
    extractions=[json.loads((source/f"archives/{time}h/extraction.json").read_text()) for time in [24,72]]
    para(f"本轮已取得PXD033504两个完整未筛proteinGroups、原样本解码、参数及summary。目标成员跨全部固实块，{sum(e['range_parts'] for e in extractions):,}个成功HTTP206段与header段覆盖全部容器；两个容器分别核对PRIDE上游SHA1，{sum(len(e['outputs']) for e in extractions)}个成员CRC通过。完整指公开处理蛋白数组，本轮未重新搜索RAW谱图。")
    para("HTPC为人睾丸管周平滑肌样细胞，1µM DEX对溶剂对照，24/72h。供体按原日期解码，三技术重复先平均log2LFQ，每条件至少2/3可测，至少三配对日期才检验。LFQ0缺测、不填补。secretome每时点样本范围由原解码、summary与发表六标签分别核对，不能沿用旧proteome七标签结论。与旧proteome同一生物学来源，不增加独立ASM供体。")
    para("原secretome参数："+"；".join(r["time_h"]+"h MaxQuant "+r["maxquant_version"] for r in versions)+"。原发表/PRIDE记录1.6.11.0；实际是否一致依原参数，旧proteome1.6.15.0版本冲突仍仅属于其原范围。")
    for time in ["24","72"]:
        local=[r for r in sample_scope if r["time_h"] == time]
        differences=[r["donor"] for r in local if r["in_original_LFQ"] != r["in_published_selected_headers"]]
        para(time+"h原LFQ与发表标签比较："+("范围一致。" if not differences else "不一致标签为"+"、".join(differences)+"，新增标签的作者QC/年龄未知，发表子集另分析。")+
             "原数组标签："+"、".join(r["donor"] for r in local if r["in_original_LFQ"] == "True")+"。")
    addtable(table1,"表1 完整数组与固定候选分母。全数组两时点共"+str(len(full))+"个QC组统一BH；候选共112项Holm，缺测P=1占位但效应/CI为空。")
    para("全部原QC、可估计/不足配对组、检测候选与未检出位置分别公开，不按新显著性替换原名单。技术列和供体日期均从原解码及summary双源核对，并与PRIDE RAW文件名逐一匹配。")
    heading("2 候选实测结果与供体分布",True)
    addtable(table2,"表2 所有可估计候选及ANGPTL2缺测位置。效应为每供体DEX−control平均log2LFQ差，CI点态95%；Holm112仅对应本轮56×2家族。")
    for gene in ["ANGPTL2","FSTL3"]:
        for time in ["24","72"]:
            r=next(r for r in candidates if r["symbol"]==gene and r["time_h"]==time)
            para(f"{gene} {time}h：{r['status']}，{r['n_pairs']}个有效配对标签；效应{number(r['effect'])}，CI {number(r['ci_low'])}至{number(r['ci_high'])}，Holm112={pv(r['holm112'])}，精确Holm112={pv(r['exact_holm112'])}。")
    para("ANGPTL2 24h唯一组有14条肽/13条unique肽，序列覆盖37.3%；两个合格日期标签的配对差分别−.453和−1.088，与ASM转录下调方向一致，但n=2低于固定至少三供体规则，不能报告确认统计。FSTL3两时点各2条unique肽；24h仅10.09.19有co1/dex2个非零LFQ技术测量、72h该标签co0/dex2，不满足配对阈值。识别和可分泌注释不等于药物响应已确认。")
    cross = rows(source/"candidate_cross_evidence.csv")
    for gene in ["ANGPTL2","FSTL3"]:
        z = [r for r in cross if r["symbol"] == gene]
        para(gene+"跨层方向："+"；".join(r["time_h"]+"h "+r["direction_concordance"] for r in z)+
             "。ASM为BUD转录、HTPC为DEX上清蛋白，药物/组织/剂量/尺度不同；同向是线索，反向是范围限制，均不能直接替代ASM分泌验证。")
    figure("Figure_1_candidates","图1 重点候选每个可测日期标签的药物响应及均值/点态CI；两时点不按点位置连线。缺测不是无变化。")
    heading("3 敏感性、直接时间差与分泌注释",True)
    addtable(table3,"表3 发表六标签子集与严格3/3技术测量。每种方案各112项Holm，均不是新增队列。")
    para("FBN2 72h主分析7/7可测标签↑，+1.535（CI+1.112至+1.958），Holm112=.0127；严格3/3技术测量6/6↑，+1.419（.0259）。发表六标签中只有5对可测：+1.466、Holm112=.270未通过，精确校正全部未确认。此线索对供体范围敏感，不能写为发表六成人已确认；24hFBN2不通过。PTX3 24h6/6↑但Holm112=.592，72h亦未通过且一轮留一方向反转。")
    for gene in ["ANGPTL2","FSTL3"]:
        r=next(r for r in direct if r["symbol"]==gene)
        para(f"{gene}直接72h−24h药物响应差：n={r['n_pairs']}，{r['status']}，效应{number(r['effect'])}，Holm56={pv(r['holm56'])}。两个时点各自MaxQuant处理，虽secretome原文使用相同LC梯度，也不能将此差解释为绝对分泌累积、分泌速率或无混杂药物×时间机制。")
    para("UniProt reviewed Q9UKU9(ANGPTL2)、O95633(FSTL3)原JSON保存：两者有Secreted/信号肽注释。ANGPTL2的内皮出芽/旁分泌功能注释，FSTL3对activin等TGFβ家族结合/拮抗注释，只提供可分泌蛋白身份和已有功能背景；本轮没有测量候选受体通路、GR依赖或ASM收缩功能。FSTL3还存在核内/异构体注释，bottom-up共同肽不能确认分泌异构体。")
    para("培养上清LFQ是细胞外相对丰度，可能受分泌、释放、死亡、细胞数、摄取和降解影响；原MS每样本固定750ng上样，不是每细胞单位时间绝对通量。跨组织同向可作为线索，反向结果保留，不能将ASM RNA方向直接等同于HTPC分泌。")
    heading("4 完整分泌蛋白组与发表处理范围",True)
    figure("Figure_2_complete_array","图2 完整可估计蛋白组的效应与统一BH。橙色BH<.05，灰色未通过；不足配对组仍在全family分母内但不绘为零效应。")
    para("全数组不要求single-gene，结果单位是protein group；固定基因只取唯一single-gene且QC合格组。合并基因组不分配为单基因效应。全部精确统一BH与固定Holm结果另列，小样本精确校正未确认不改写为非参数成功。")
    for r in comparable["summary"]:
        para(f"原发表{r['time_h']}h ALL共{r['published_all_rows']}组，原数组exact匹配{r['all_exact_matches']}；发表DIFF选中{r['selected_rows']}组，逐格比较{r['compared_cells']}技术测量，{r['matched_rounding']}舍入一致、{r['raw_zero_published_finite']}原LFQ0却发表有限值。全部差异逐格保留；不猜作者缺失值处理，不从选中表缺席推断阴性。")
    for r in diagnostics["summary"]:
        para(f"进一步诊断{r['time_h']}h：按Protein IDs/原Majority IDs可对应{r['exact_or_Majority_resolved']}/{r['published_all_rows']}发表ALL组，总Intensity原值相同{r['total_intensity_equal']}组；选中表log2(未归一Intensity)也无逐格匹配。非零原LFQ log2−发表值范围{r['nonzeroLFQ_offset_min']:.3g}至{r['nonzeroLFQ_offset_max']:.3g}。")
    para("24h原LFQ的803个非零选中测量均不匹配发表log2值，ALL总Intensity也只有9/467一致；差不是四位舍入或统一常数偏移，作者处理/输入差异尚未核实。因此本轮24h明确是存档MaxQuant数组新重分析，不能称复现原Perseus结果。72h非零2128测量匹配舍入，原LFQ0的320个发表有限值仍无法当作原非零实测；本轮不填补。")
    heading("5 数值复核、结论与复现",True)
    para(f"独立R从原proteinGroups、解码与summary重建{audit['raw_measurements']:,}个LFQ测量、{audit['donor_rows']:,}供体组行、{audit['full_family']}全组统计/统一BH及112候选Holm、224敏感性、56直接时间差、{audit['leave_one_rows']}留一方向；全部通过。Python原输出与R是同来源复核，不增加独立证据。")
    para(f"实际源码载入BioExplainR {workflow['package_version']}，完整分泌蛋白组与固定候选两个context、项目和中英文HTML已保存；全family={workflow['full_family']}及112完整事实分母通过，在线调用0。统计由原数组及Python/R计算，BioExplainR用于事实和范围组织。")
    supported=[r for r in cross if r["same_direction_corrected_support"] == "True"]
    opposite=[r for r in cross if r["direction_concordance"] == "opposite"]
    para("候选与原ASM转录同向且Holm112通过："+("、".join(r["symbol"]+r["time_h"]+"h" for r in supported) or "无")+
         "。可估计但反向："+("、".join(r["symbol"]+r["time_h"]+"h" for r in opposite) or "无")+"。完整缺测和未通过条目不删除。")
    para("ANGPTL2/FSTL3有可分泌身份；具体药物响应须按表2/3的方向、有效配对、家族校正和敏感性解释。可测不等于支持ASM方向。培养上清或跨组织RNA-蛋白方向不能建立ASM分泌/GR/功能因果。")
    para("完整56×2、全部QC蛋白组、原测量、日期效应、敏感性/时间差、发表逐格冲突和UniProt注释见tables及research/results/discovery_secretome_v1。保存结果核验/新目录数值复现见README。实际下载重试和修订见EXECUTION_NOTES。")
    para("来源：[1] Stepanov YK et al. Cells 2022;11:3164，doi:10.3390/cells11193164，PMC9562650，TableS2。[2] PRIDE PXD033504，两SEARCH完整处理容器；来源文件清单/项目API原响应已保存。[3] UniProt Q9UKU9、O95633原公开JSON，本轮查询UTC及hash保存。[4] ASM原54集合及六成人来源的独立封存v0.12结果（本轮不重算RNA）。")
    docx=dest/(NAME+".docx"); doc.save(docx)
    (dest/"results.md").write_text("\n\n".join(md)+"\n")
    css="body{font-family:system-ui,'Noto Sans CJK SC',sans-serif;max-width:1080px;margin:2em auto;padding:0 1em;line-height:1.7;color:#26343b}h1,h2{color:#236477}table{border-collapse:collapse;width:100%;font-size:.85em}td,th{border:1px solid #ccd;padding:.4em;text-align:left}th{background:#edf4f6}.table-wrap{overflow-x:auto}img{max-width:100%;height:auto}p{overflow-wrap:anywhere}figure{margin:1em 0}"
    (dest/(NAME+".html")).write_text("<!doctype html><html lang='zh-CN'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>"+title+"</title><style>"+css+"</style><body>"+"\n".join(h)+"</body></html>")
    command=["libreoffice","-env:UserInstallation=file:///tmp/opencode/validation-report-lo","--headless","--convert-to","pdf","--outdir",str(dest),str(docx)]
    p=subprocess.run(command,capture_output=True,text=True,timeout=180)
    (dest/"export_log.txt").write_text(p.stdout+p.stderr)
    assert p.returncode==0 and docx.with_suffix(".pdf").exists()
    with fitz.open(docx.with_suffix(".pdf")) as pdf:
        pages=len(pdf)
    json_out(dest/"export.json",dict(status="completed",utc=utc(),pages=pages,tables=3,figures=2,table_files=len(list(tab.glob("*.csv"))),
        command=command,script=record(Path(__file__)),pdf_figure_mode="PNG previews in Word/PDF; separate original vector figure PDFs"))
    print("Secretome report exported",pages,"pages")


if __name__ == "__main__":
    p=argparse.ArgumentParser(); p.add_argument("source",type=Path); p.add_argument("destination",type=Path)
    a=p.parse_args(); build(a.source.resolve(),a.destination.resolve())
