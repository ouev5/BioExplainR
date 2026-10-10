"""Frozen cross-array provenance diagnostics and scoped public-source search."""
import argparse
import csv
import json
import math
import re
import sys
import urllib.parse
import xml.etree.ElementTree as ET
from pathlib import Path

import numpy as np
import openpyxl

from validation_sources import ROOT, HERE, EPMC, utc, record, json_out, csv_out, fetch

OLD = ROOT / "research/results/discovery_secretome_v1"
HUMAN = ROOT / "research/results/discovery_validation_v2/human_protein"
SET = ROOT / "research/results/discovery_set_validation_v1"
SPECS = {
    "secretome24": OLD / "archives/24h/extracted/txt/proteinGroups.txt",
    "secretome72": OLD / "archives/72h/extracted/txt/proteinGroups.txt",
    "proteome24": SET / "protein_full_tables_small/PXD033514/extracted_verified/txt/proteinGroups.txt",
    "proteome72": SET / "protein_full_tables_small/PXD033534/extracted/txt/proteinGroups.txt",
}
SHEETS = [(1, "proteome24"), (3, "proteome72"), (5, "secretome24"), (7, "secretome72")]
QUERIES = {
    "asm_secretome": 'TITLE_ABS:("airway smooth muscle" OR "bronchial smooth muscle") AND TITLE_ABS:(secretome OR "conditioned medium" OR proteomic)',
    "asm_gc_secretion": 'TITLE_ABS:("airway smooth muscle" OR "bronchial smooth muscle") AND TITLE_ABS:(dexamethasone OR glucocorticoid OR budesonide) AND TITLE_ABS:(secreted OR secretion OR extracellular)',
    "asm_candidates": '(ANGPTL2 OR FSTL3 OR FBN2) AND ("airway smooth muscle" OR "bronchial smooth muscle")',
    "asm_proteomics": 'TITLE_ABS:("human airway smooth muscle") AND TITLE_ABS:(proteome OR proteomics OR "mass spectrometry")',
    "pxd033504": 'PXD033504',
    "htpc_gc": 'TITLE_ABS:(peritubular) AND TITLE_ABS:(dexamethasone OR glucocorticoid)',
}
PRIDE_QUERIES = ["airway smooth muscle", "bronchial smooth muscle", "ASM secretome", "smooth muscle dexamethasone"]


def init(out):
    if out.exists():
        raise ValueError(f"Refusing overwrite {out}")
    from seal_secretome import verify
    verify()
    out.mkdir()
    (out / "sources").mkdir()
    (out / "rgo_before_provenance.md").write_bytes((ROOT / "rgo.md").read_bytes())
    inputs = [record(p) for p in SPECS.values()] + [record(HUMAN / "Table S2.xlsx"),
        record(HUMAN / "PMC9562650.xml"), record(OLD / "sample_map.csv")]
    json_out(out / "freeze.json", dict(utc=utc(), protocol=record(HERE / "SECRETOME_PROVENANCE_PROTOCOL.md"),
        script=record(Path(__file__)), inputs=inputs, array_paths={k:record(v) for k,v in SPECS.items()},
        queries=QUERIES, pride_queries=PRIDE_QUERIES, tolerance=.000051,
        timing="Prior secretome effects and mismatch known; cross-array source comparisons not run; new search not performed",
        scientific_scope="Post-result provenance diagnostics; no new independent sample or corrected candidate significance"))


def read_table(path):
    csv.field_size_limit(10*1024*1024)
    with path.open(newline="") as stream:
        return list(csv.DictReader(stream, delimiter="\t"))


def identity_map(table):
    exact, majority = {}, {}
    for i,r in enumerate(table,2):
        exact.setdefault(r["Protein IDs"], []).append((i,r))
        majority.setdefault(r["Majority protein IDs"], []).append((i,r))
    def lookup(ids):
        matches = exact.get(ids, [])
        rule = "exact_Protein_IDs"
        if not matches:
            matches = majority.get(ids, [])
            rule = "exact_Majority_IDs"
        if len(matches) != 1:
            return "unresolved", "", None
        i,r = matches[0]
        return rule, i, r
    return lookup


def analyze(out):
    frozen = json.loads((out / "freeze.json").read_text())
    for r in frozen["inputs"]:
        assert record(ROOT / r["path"]) == r, r["path"]
    assert record(Path(__file__)) == frozen["script"]
    tables = {k:read_table(p) for k,p in SPECS.items()}
    lookups = {k:identity_map(t) for k,t in tables.items()}
    book = openpyxl.load_workbook(HUMAN / "Table S2.xlsx", read_only=True, data_only=True)
    csv_out(out / "workbook_headers.csv", [dict(sheet=s, row=i, values=json.dumps(list(r),ensure_ascii=False))
        for s in book.sheetnames for i,r in enumerate(book[s].iter_rows(max_row=3,values_only=True),1)])
    allcells, diffcells, published = [], [], []
    # Use frozen verified secretome decoding to map numeric experiment labels.
    with (OLD / "sample_map.csv").open(newline="") as stream:
        samplemap = list(csv.DictReader(stream))
    for n,expected in SHEETS:
        allsheet = next(s for s in book.sheetnames if s.startswith(f"S{n}_"))
        diffsheet = next(s for s in book.sheetnames if s.startswith(f"S{n+1}_"))
        for i,r in enumerate(book[allsheet].iter_rows(min_row=2,values_only=True),2):
            if not r[0]:
                continue
            for key,lookup in lookups.items():
                rule, rownum, original = lookup(r[0])
                value = float(original["Intensity"]) if original else None
                pub = float(r[3])
                allcells.append(dict(sheet=allsheet, expected_array=expected, array=key, published_row=i,
                    protein_ids=r[0], symbol=r[1] or "", published_intensity=pub, source_row=rownum,
                    mapping=rule, original_intensity=value if value is not None else "", equal=value == pub))
        data = list(book[diffsheet].iter_rows(values_only=True))
        hi = next(i for i,r in enumerate(data[:5]) if r[0] == "Protein Name")
        columns = [(j,re.fullmatch(r"(\d{2}\.\d{2}\.\d{2})_?(co|dex)([123])",v))
            for j,v in enumerate(data[hi]) if isinstance(v,str)]
        columns = [(j,m) for j,m in columns if m]
        assert len(columns) == 36
        for i,r in enumerate(data[hi+1:],hi+2):
            if not r[2]:
                continue
            for j,m in columns:
                pub = float(r[j])
                published.append(dict(sheet=diffsheet,expected_array=expected,published_row=i,protein_ids=r[7],
                    symbol=r[1],donor=m[1],condition=m[2],technical_rep=int(m[3]),published_log2=pub))
                for key,lookup in lookups.items():
                    rule, rownum, original = lookup(r[7])
                    time = 24 if key.endswith("24") else 72
                    exp = next((s["experiment"] for s in samplemap if key.startswith("secretome") and
                        int(s["time_h"]) == time and s["donor"] == m[1] and s["condition"] == m[2] and
                        s["technical_rep"] == m[3]), f"{m[1]}_{m[2]}{m[3]}")
                    if original:
                        experiments = [col.removeprefix("LFQ intensity ") for col in original if col.startswith("LFQ intensity ")]
                        # Proteome headers are directly dated; never match columns by numerical values.
                        direct = [e for e in experiments if re.fullmatch(re.escape(m[1])+r"_?"+m[2]+m[3],e)]
                        if direct:
                            assert len(direct) == 1
                            exp = direct[0]
                    for measure in ["LFQ intensity", "Intensity", "iBAQ"]:
                        col = measure+" "+exp
                        value = float(original[col]) if original and col in original else None
                        log = math.log2(value) if value is not None and value>0 else None
                        diffcells.append(dict(sheet=diffsheet,expected_array=expected,array=key,measure=measure,
                            published_row=i,protein_ids=r[7],symbol=r[1],donor=m[1],condition=m[2],technical_rep=int(m[3]),
                            published_log2=pub,source_row=rownum,mapping=rule,source_column=col,
                            original=value if value is not None else "",original_log2=log if log is not None else "",
                            delta=log-pub if log is not None else "",matches_rounding=log is not None and abs(log-pub)<=.000051))
    csv_out(out / "all_intensity_cross_array.csv", allcells)
    csv_out(out / "published_measurements.csv", published)
    csv_out(out / "diff_measurement_cross_array.csv", diffcells)
    summary = []
    for n,expected in SHEETS:
        for key in SPECS:
            a = [r for r in allcells if r["expected_array"] == expected and r["array"] == key]
            for measure in ["LFQ intensity", "Intensity", "iBAQ"]:
                z = [r for r in diffcells if r["expected_array"] == expected and r["array"] == key and r["measure"] == measure]
                finite = [r for r in z if r["original_log2"] != ""]
                delta = [r["delta"] for r in finite]
                summary.append(dict(expected_array=expected,array=key,measure=measure,all_rows=len(a),
                    all_resolved=sum(r["mapping"] != "unresolved" for r in a),all_intensity_equal=sum(r["equal"] for r in a),
                    selected_cells=len(z),original_nonzero=len(finite),matches_rounding=sum(r["matches_rounding"] for r in z),
                    original_zero=sum(r["original"] == 0 for r in z),unmapped=sum(r["original"] == "" for r in z),
                    delta_min=min(delta) if delta else "",delta_max=max(delta) if delta else "",
                    correlation=float(np.corrcoef([r["original_log2"] for r in finite],[r["published_log2"] for r in finite])[0,1]) if len(finite)>2 else ""))
    csv_out(out / "cross_array_summary.csv",summary)
    bysample = []
    z = [r for r in diffcells if r["expected_array"] == "secretome24" and r["array"] == "secretome24" and r["measure"] == "LFQ intensity"]
    for donor,cond,rep in sorted({(r["donor"],r["condition"],r["technical_rep"]) for r in z}):
        finite = [r for r in z if (r["donor"],r["condition"],r["technical_rep"]) == (donor,cond,rep) and r["delta"] != ""]
        bysample.append(dict(donor=donor,condition=cond,technical_rep=rep,n_nonzero=len(finite),
            delta_mean=float(np.mean([r["delta"] for r in finite])),delta_sd=float(np.std([r["delta"] for r in finite],ddof=1)),
            delta_min=min(r["delta"] for r in finite),delta_max=max(r["delta"] for r in finite)))
    csv_out(out / "secretome24_sample_offsets.csv",bysample)
    article = ET.parse(HUMAN / "PMC9562650.xml").getroot()
    passages = []
    for i,p in enumerate(article.iter("p"),1):
        text = " ".join("".join(p.itertext()).split())
        if re.search(r"Perseus|MaxQuant|PXD|secretom|supplement|mass spectrom",text,re.I):
            passages.append(dict(paragraph=i,text=text))
    csv_out(out / "original_methods_passages.csv",passages)
    json_out(out / "analysis.json",dict(status="completed",utc=utc(),all_comparisons=len(allcells),
        selected_comparisons=len(diffcells),published_measurements=len(published),summary=summary,
        limitations="Cross-array values are diagnostics, not source correction or new candidate tests; exact LFQ does not verify imputation"))
    print(json.dumps(summary,ensure_ascii=False,indent=2))


def search(out):
    logs, papers = [], {}
    for label,query in QUERIES.items():
        query += " AND FIRST_PDATE:[1900-01-01 TO 2026-10-10]"
        cursor,returned = "*",0
        for page in range(1,11):
            path = out / f"sources/epmc_{label}_{page}.json"
            url = EPMC+"search?"+urllib.parse.urlencode(dict(query=query,format="json",resultType="core",pageSize=100,cursorMark=cursor))
            log = fetch(url,path,timeout=30)
            log.update(label=label,query=query,page=page)
            logs.append(log)
            if log["status"] != "retrieved":
                break
            data = json.loads(path.read_text()); hits = data.get("resultList",{}).get("result",[])
            returned += len(hits)
            log.update(hit_count=data["hitCount"],returned=len(hits),query_complete=returned>=data["hitCount"])
            for r in hits:
                key = r["source"]+":"+r["id"]
                if key not in papers:
                    papers[key] = dict(id=r["id"],source=r["source"],pmcid=r.get("pmcid",""),doi=r.get("doi",""),
                        title=r.get("title",""),abstract=r.get("abstractText",""),year=r.get("pubYear",""),queries=[])
                papers[key]["queries"].append(label)
            if log["query_complete"]:
                break
            cursor = data["nextCursorMark"]
        print(label,returned,flush=True)
    csv_out(out / "literature_records.csv",[dict(r,queries="|".join(sorted(set(r["queries"])))) for r in papers.values()],
        fields=["id","source","pmcid","doi","title","abstract","year","queries"])
    projects = []
    for i,query in enumerate(PRIDE_QUERIES,1):
        path=out / f"sources/pride_search_{i}.json"
        url="https://www.ebi.ac.uk/pride/ws/archive/v2/projects?"+urllib.parse.urlencode(dict(keyword=query,pageSize=100,page=0))
        log=fetch(url,path,timeout=30); log.update(label="pride",query=query)
        logs.append(log)
        if log["status"] == "retrieved":
            try:
                data=json.loads(path.read_text())
                hits=data.get("_embedded",{}).get("projects",[]) if isinstance(data,dict) else data
                log.update(returned=len(hits),pagination=data.get("page",{}) if isinstance(data,dict) else {})
                for r in hits:
                    projects.append(dict(query=query,accession=r.get("accession",""),title=r.get("title",""),metadata=json.dumps(r,ensure_ascii=False)))
            except Exception as exc:
                log["parse_error"]=str(exc)
    csv_out(out / "pride_projects.csv",projects,fields=["query","accession","title","metadata"])
    json_out(out / "search_acquisition.json",dict(utc=utc(),records=logs,papers=len(papers),
        scope="Scoped indexed search; no-hit or endpoint failure is not evidence of nonexistence"))


if __name__ == "__main__":
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("action",choices=["init","analyze","search"]);p.add_argument("output",type=Path)
    a=p.parse_args(); globals()[a.action](a.output.resolve())
