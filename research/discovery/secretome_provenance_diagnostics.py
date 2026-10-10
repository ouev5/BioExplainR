"""Fixed mechanistic-scale diagnostics, source feasibility and real sample linkage."""
import argparse
import json
import math
import re
import xml.etree.ElementTree as ET
from pathlib import Path

import numpy as np
import openpyxl

from secretome_provenance import SPECS, read_table, identity_map
from secretome_provenance_complete import rows
from validation_sources import utc, record, json_out, csv_out, fetch


def main(out):
    dest=out / "complete"; cross=dest / "diff_measurement_cross_array.csv"
    json_out(dest / "diagnostic_freeze.json",dict(utc=utc(),script=record(Path(__file__)),
        inputs=[record(cross),record(SPECS["secretome24"])],
        timing="Protein-specific approximatelyconstant offsets now known; no normalization-factor formula evaluated",
        formulas=["LFQ / Peptides","LFQ / Razor + unique peptides","LFQ / Unique peptides",
            "Intensity / Peptides","Intensity / Razor + unique peptides","Intensity / Unique peptides"],
        criterion="Only positive denominator/raw; log2 error <=.000051. All six formulas disclosed, no candidate re-estimation",
        note="Protein-specific fitted offsets not correction; hypothesized normalization diagnostics only"))
    raw=read_table(SPECS["secretome24"]); lookup=identity_map(raw)
    cells=[r for r in rows(cross) if r["expected_array"]==r["array"]=="secretome24" and r["measure"]=="LFQ intensity"]
    transformed=[]
    for r in cells:
        _,i,protein=lookup(r["protein_ids"])
        assert protein and i==int(r["source_row"])
        exp=r["source_column"].removeprefix("LFQ intensity ")
        for measure in ["LFQ intensity","Intensity"]:
            value=float(protein[measure+" "+exp])
            for denom in ["Peptides","Razor + unique peptides","Unique peptides"]:
                d=int(protein[denom]);log=math.log2(value/d) if d>0 and value>0 else None
                error=log-float(r["published_log2"]) if log is not None else None
                transformed.append(dict(protein_ids=r["protein_ids"],symbol=r["symbol"],donor=r["donor"],condition=r["condition"],
                    technical_rep=r["technical_rep"],measure=measure,denominator=denom,raw_value=value,denominator_value=d,
                    published_log2=r["published_log2"],transformed_log2=log if log is not None else "",error=error if error is not None else "",
                    matched=error is not None and abs(error)<=.000051))
    csv_out(dest / "scale_formula_comparisons.csv",transformed)
    summary=[]
    for measure in ["LFQ intensity","Intensity"]:
        for denom in ["Peptides","Razor + unique peptides","Unique peptides"]:
            z=[r for r in transformed if r["measure"]==measure and r["denominator"]==denom]
            finite=[r for r in z if r["error"]!=""]
            summary.append(dict(measure=measure,denominator=denom,cells=len(z),eligible=len(finite),matched=sum(r["matched"] for r in z),
                error_min=min(r["error"] for r in finite),error_max=max(r["error"] for r in finite),
                mean_absolute_error=float(np.mean([abs(r["error"]) for r in finite]))))
    csv_out(dest / "scale_formula_summary.csv",summary)
    wb=openpyxl.load_workbook(out / "sources/PXD055968_Link_between_files.xlsx",read_only=True,data_only=True)
    linkage=[dict(sheet=s,row=i,values=json.dumps(list(r),ensure_ascii=False)) for s in wb.sheetnames
        for i,r in enumerate(wb[s].iter_rows(values_only=True),1)]
    csv_out(dest / "EV_file_linkage.csv",linkage)
    passages=[];attachments=[]
    for label,file in [("EV","PMC12141081.xml"),("PTX3","PTX3_article.xml"),("cough","cough_article.xml")]:
        root=ET.parse(out / "sources" / file).getroot()
        for i,p in enumerate(root.iter("p"),1):
            text=" ".join("".join(p.itertext()).split())
            if label != "EV" or re.search(r"proteom|patient|subjects|vesicles.*medium|data.*request",text,re.I):
                passages.append(dict(source=label,paragraph=i,text=text))
        for el in root.iter():
            if el.tag in ["media","supplementary-material"]:
                attachments.append(dict(source=label,tag=el.tag,attributes=json.dumps(el.attrib),text="".join(el.itertext())))
    csv_out(dest / "lead_methods_passages.csv",passages)
    csv_out(dest / "lead_attachment_metadata.csv",attachments)
    # Explicit public supplement probes; fail bodies are not interpreted as PDFs.
    urls=["https://pmc.ncbi.nlm.nih.gov/articles/instance/12141081/bin/mmc1.pdf"]
    logs=[]
    for i,url in enumerate(urls,1):
        path=out / f"sources/EV_supplement_probe_{i}.bin"
        log=fetch(url,path,timeout=20)
        if path.exists():
            log["pdf_signature"]=path.read_bytes().startswith(b"%PDF")
        logs.append(log)
    json_out(dest / "diagnostics.json",dict(status="completed",utc=utc(),summary=summary,
        linkage_rows=len(linkage),supplement_probes=logs,
        scope="Scale/feasibility diagnostics; no new pairedGC candidate protein estimates"))
    print(json.dumps(dict(summary=summary,linkage=linkage,attachments=attachments),ensure_ascii=False,indent=2))


if __name__=="__main__":
    p=argparse.ArgumentParser();p.add_argument("output",type=Path);main(p.parse_args().output.resolve())
