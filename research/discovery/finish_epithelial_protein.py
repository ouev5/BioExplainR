"""Readable delivery and read-only verification for frozen epithelial protein results."""
import argparse
import ast
import csv
import html
import json
from pathlib import Path
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from epithelial_protein import OUT, HERE, ROOT, CODE, rows, check_freeze, prior_check
from validation_sources import record, utc, json_out

TABLES = ["summary", "candidate_results", "detection", "measurements", "individual_effects",
          "sensitivity", "T2_descriptive", "all_protein_results", "protein_row_qc", "A549_normalization"]


def number(value):
    if value in ["", None]:
        return "不可估计"
    v = float(value)
    return f"{v:.4g}"


def build():
    check_freeze()
    audit = json.loads((OUT / "r_audit.json").read_text())
    workflow = json.loads((OUT / "bioexplainr_workflow.json").read_text())
    assert audit["status"] == workflow["status"] == "pass"
    results = rows(OUT / "candidate_results.csv")
    summary = rows(OUT / "summary.csv")
    sensitivity = rows(OUT / "sensitivity.csv")
    lines = ["# 已下载上皮蛋白数组：固定候选集合的实际分析", "", "2026-10-10。结果后跨细胞来源探索；56固定位置（原54基因＋PER1／NR3C1）。", "",
        "## 结果摘要", "", "|来源／比较|唯一QC蛋白组|可估计／56|参数Holm通过|精确Holm通过|全数组BH通过／可估计|", "|---|---:|---:|---:|---:|---:|"]
    labels = {"DEX_48h":"A549：48h DEX−未处理", "ICS_unstimulated":"HBEC：未刺激临床ICS后−前",
              "ICS_polyIC":"HBEC：polyIC下临床ICS后−前", "ICS_x_polyIC":"HBEC：同人直接ICS×polyIC差中差"}
    for s in summary:
        lines.append(f"|{labels[s['contrast']]}|{s['unique_groups']}|{s['estimable']}/56|{s['holm_pass']}|{s['exact_holm_pass']}|{s['global_BH_pass']}/{s['global_estimable']}|")
    lines += ["", "主家族分别56项；HBEC两个次比较合为112项。每一家族均保留不可估计P=1位置。",
        "全数组BH：A549全部QC行一个家族，HBEC全部QC行×三个比较一个家族；分表呈现不表示分别校正。",
        "参数与精确检验并列；小样本的精确P离散性限制不被参数显著性替代。", "",
        "## 原12重点候选和控制：主比较逐项公开", "",
        "|候选|A549状态／每组有效孔n|A549效应[点态95%CI]／Holm56|HBEC未刺激状态／配对n|HBEC效应[点态95%CI]／Holm56|",
        "|---|---|---|---|---|"]
    genes = rows(OUT / "frozen_genes.csv")
    index = {(r["dataset"],r["symbol"],r["contrast"]):r for r in results}
    selected = [g for g in genes if not g["set_rank"] or int(g["set_rank"])<=12]
    def cell(r):
        if r["effect"] == "":
            return "不可估计；校正位置保留"
        return f"{number(r['effect'])} [{number(r['ci_low'])}, {number(r['ci_high'])}]／{number(r['holm'])}"
    for g in selected:
        a = index["A549",g["symbol"],"DEX_48h"]
        h = index["HBEC",g["symbol"],"ICS_unstimulated"]
        label = "TMT1A（METTL7A）" if g["symbol"]=="TMT1A" else g["symbol"]
        lines.append(f"|{label}|{a['status']}／{a['n_control']},{a['n_treated']}|{cell(a)}|{h['status']}／{h['n_pairs']}|{cell(h)}|")
    significant = [r for r in results if float(r["holm"]) < .05]
    lines += ["", "## 完整集合中参数Holm通过的项目", ""]
    if significant:
        lines += ["|来源／比较|候选|效应／校正P|与原ASM转录发现方向|精确校正P|", "|---|---|---|---|---|"]
        for r in significant:
            direction = "同向" if r["direction_matches"]=="True" else "反向" if r["direction_matches"]=="False" else "控制，无冻结方向"
            lines.append(f"|{labels[r['contrast']]}|{r['symbol']}|{number(r['effect'])}／{number(r['holm'])}|{direction}|{number(r['exact_holm'])}|")
    else:
        lines.append("三个预定家族均无参数Holm<0.05项目。未通过不证明等效，也不把不可测位置记为蛋白阴性。")
    normalization = rows(OUT / "A549_normalization.csv")
    raw_pass = [r for r in sensitivity if r["dataset"]=="A549" and float(r["holm_sensitivity_56"])<.05]
    lines += ["", "## 预处理和敏感性", "",
        f"- A549逐列中心以{normalization[0]['reference_rows']}个全部8孔有正强度的QC蛋白行为公共参考；不使用候选选择参考。",
        f"- 未额外归一的A549 log2强度敏感性：56位置中{len(raw_pass)}项Holm通过（" + ("、".join(r["symbol"] for r in raw_pass) or "无") + "）；主分析与敏感性效应在`sensitivity.csv`完整公开。",
        "- HBEC不二次log、不重归一、不填补。公开表量级兼容log2，但精确转换日志未取得，效应仅标导出尺度单位，不换算倍数。",
        "- HBEC每个有效配对逐人公开；留一人重新估计和T2高／低各自有效n、描述性均值见完整表。T2未追加显著性检验。",
        "- 无测量、QC排除、多蛋白组歧义、有效孔／配对不足各自保留原因；这些不是候选生物学阴性。", "",
        "## 能支持的科学范围", "",
        "A549估计的是癌源上皮细胞系48h培养孔之间的DEX关联蛋白丰度差，每组4孔不是4供体，后缀未证实配对，也未证实独立实验批次。",
        "HBEC估计的是10个体临床ICS前后的上皮蛋白关联；用同人四格直接检验polyIC背景下前后变化是否不同，不能归因培养GC因素。",
        "两者均不提供成人原代ASM受控GC上清设计，不能合并效应、P或供体数。新上皮结果不确认原ASM PER2／METTL7A候选的蛋白、节律、分泌、GR遗传因果或功能；ANGPTL2／FSTL3／FBN2原结论保持对应证据范围。",
        "没有感染或病毒复制结局，本轮不对抗病毒效果作推断。", "",
        "## 实际复核与R包应用", "",
        f"base R从原表／原作者SRDF独立核对{audit['global_tests']}全数组检验、{audit['candidate_tests']}候选位置、{audit['measurement_cells']}候选测量、{audit['individual_positions']}逐人位置、{audit['sensitivity_tests']}敏感性及{audit['T2_descriptive_rows']}描述性行，通过。",
        f"实际BioExplainR {workflow['package_version']}源码载入，3上下文（56／56／112），完整分母与显著计数核对，通过；双语HTML已保存，在线模型调用0。",
        "工作簿逐CSV单元核对，清单与本地ZIP逐字节核对；数字与文件核对不等于人工科学审阅。", "",
        "## 来源与复现", "",
        "- A549：PRIDE [PXD063431](https://www.ebi.ac.uk/pride/archive/projects/PXD063431)，公开全蛋白表5,428行，PD2.5，原项目处理说明和源表按开始freeze hash读取。",
        "- HBEC：PRIDE [PXD029135](https://www.ebi.ac.uk/pride/archive/projects/PXD029135)，公开median-normalized全表3,072行；MaxQuant1.6.15.0／THISP2020-06-01／NormalyzerDE median。",
        "- 论文对应范围沿用前轮原文核查：HBEC关联文章的蛋白比较为3例imiquimod；不能把整个43样本前后设计称为该文统计复现。",
        "- 输入、方案、完整表、缺测和独立复核见`analysis_freeze.json`、`candidate_results.csv`、`all_protein_results.csv`、`r_audit.json`。",
        "- 保存结果核验（无网络、无统计重跑）：`python3 research/discovery/finish_epithelial_protein.py verify`。",
        "- 新数值重跑须另建输出目录并另冻结；默认脚本拒绝覆盖本轮封存。", "",
        "## 下一实际材料", "",
        "成人同供体逐时GC／vehicle、完整NR3C1×GC和ASM未筛上清数组仍需要合格公开材料；24h HTPC来源差异仍需真实Perseus／版本说明。现有PXD055968最终导出缺件状态保持，按负责人指示跳过其下载。稿件真实人员信息与最终人工科学／图稿审阅仍需实际材料。",
        "", "本次结果独立保存。GitHub／外部上传仍按既有决定暂停。"]
    (OUT / "FINDINGS.md").write_text("\n".join(lines)+"\n", encoding="utf-8")
    readme = """# 上皮蛋白候选检查 v1

2026-10-10，公开数据实际下载后的统计续接。

- [中文科学结果](FINDINGS.md)
- [完整CSV-XLSX工作簿](epithelial_protein_tables.xlsx)／[候选结果CSV](candidate_results.csv)／[全部QC蛋白CSV](all_protein_results.csv)
- [实际BioExplainR中文报告](BioExplainR-epithelial-zh.html)／[英文报告](BioExplainR-epithelial-en.html)
- [冻结方案与输入](analysis_freeze.json)／[独立R复核](r_audit.json)／[执行记录](EXECUTION_NOTES.md)
- 本地续接包`epithelial_protein_bundle.zip`和独立清单`epithelial_manifest.json`在最终核验后生成。

保存结果只读核验，在仓库根目录：

```sh
python3 research/discovery/finish_epithelial_protein.py verify
```

依赖Python3.12、numpy1.26.4、scipy1.11.4、openpyxl3.1.2；独立数值实现base R4.3.3／jsonlite。
本地ZIP包含本轮全部结果、直接输入和冻结代码，可按内部完整文件路径恢复；更早v0.13/
下载封存的层叠核验仍需原文件树。清单不覆盖旧封存或把上皮结果作为ASM验证。
"""
    (OUT / "README.md").write_text(readme, encoding="utf-8")
    notes = """# 执行记录

- 开始读取磁盘维护README与全部rgo。实测系统盘53%、/data 37%，无需继续迁移。
- 临时`/tmp/opencode`原venv已不在，系统Python具备实际依赖；系统R可载入pkgload/jsonlite，BioExplainR由项目源码实际载入。
- 按既有精确快照政策核验旧下载157文件／25完整文件、v0.13绑定302记录及ZIP，通过。旧验证函数仅在进程内将唯一rgo文档依赖指向其开始快照，原文件和旧清单没有改写。
- 新方案在候选丰度提取前写入，A549后缀未证实配对，采用Welch；HBEC沿用已核对原作者SRDF，不借样本数推断供体数。
- HBEC表缺完整转换日志，保留导出尺度；缺测不填补，混合基因/污染蛋白行排除，歧义与缺测保留固定检验位置。
- 使用Python实际数值分析；base R从原始表独立重建并复核；实际BioExplainR离线上下文／项目／双语HTML保存。
- 本轮没有网络数据取得、作者通信、在线模型或外部上传。保存结果审计和工作簿核对不代替最终人工科学审阅。
"""
    (OUT / "EXECUTION_NOTES.md").write_text(notes, encoding="utf-8")
    wb = Workbook()
    info = wb.active
    info.title = "Scope"
    info.append(["key", "value"])
    for key,val in [("scope","Exploratory epithelial results, not adult ASM controlled protein validation"),
        ("A549 units","Culture wells (4 per group), independent Welch; suffix pairing not established"),
        ("HBEC units","10 individuals, clinical pre/post ICS; no controlled GC versus vehicle"),
        ("families","A54956; HBEC primary56 and secondary112, missing P=1 slots"),
        ("effect scales","A549 median-centered log2 intensity; HBEC supplied normalized export units"),
        ("CI","Pointwise95%, not simultaneous intervals"), ("online_model_calls",0)]:
        info.append([key,val])
    for name in TABLES:
        sheet = wb.create_sheet(name[:31])
        with (OUT / (name+".csv")).open(newline="",encoding="utf-8-sig") as stream:
            for row in csv.reader(stream):
                sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    wb.save(OUT / "epithelial_protein_tables.xlsx")
    print(json.dumps(delivery_checks(), ensure_ascii=False, indent=2))


def delivery_checks():
    check_freeze()
    audit = json.loads((OUT / "r_audit.json").read_text())
    run = json.loads((OUT / "run.json").read_text())
    assert audit["status"] == "pass" and run["status"] == "completed"
    assert audit["global_tests"] == run["global_tests"]
    assert audit["candidate_tests"] == run["candidate_tests"] == 224
    assert audit["measurement_cells"] == run["measurements"]
    workflow = json.loads((OUT / "bioexplainr_workflow.json").read_text())
    assert workflow["status"] == "pass" and workflow["online_model_calls"] == 0
    results = rows(OUT / "candidate_results.csv")
    for family, facts in workflow["contexts"].items():
        subset = [r for r in results if r["family"]==family]
        assert facts["n_features"]==len(subset)
        assert facts["n_significant"]==sum(float(r["holm"])<.05 for r in subset)
        assert facts["n_estimable"]==sum(r["effect"]!="" for r in subset)
    wb = load_workbook(OUT / "epithelial_protein_tables.xlsx",read_only=True,data_only=True)
    assert wb.sheetnames == ["Scope"] + [n[:31] for n in TABLES]
    cells = 0
    for name in TABLES:
        with (OUT / (name+".csv")).open(newline="",encoding="utf-8-sig") as stream:
            expected = list(csv.reader(stream))
        actual = [["" if x is None else str(x) for x in row] for row in wb[name[:31]].iter_rows(values_only=True)]
        assert actual==expected, name
        cells += sum(len(r) for r in actual)
    wb.close()
    for language in ["zh","en"]:
        content = (OUT / f"BioExplainR-epithelial-{language}.html").read_text()
        assert "PXD063431" in content and "PXD029135" in content and "not ASM" in html.unescape(content)
    for name in CODE:
        if name.endswith(".py"):
            ast.parse((HERE/name).read_text())
    return dict(status="pass", full_array_tests=audit["global_tests"], candidate_positions=224,
                independent_R=audit, workbook_sheets=len(TABLES)+1, workbook_cells_checked=cells,
                bioexplainr_contexts=3, online_model_calls=0)


def seal():
    dest = OUT / "epithelial_manifest.json"
    if dest.exists():
        raise ValueError("Already sealed")
    verification = delivery_checks()
    previous = prior_check()
    before = (OUT / "rgo_before_epithelial.md").read_bytes()
    current = (ROOT / "rgo.md").read_bytes()
    marker = "## 最新续接：公开数据实际下载补齐".encode()
    assert current[current.index(marker):]==before[before.index(marker):]
    bundle = OUT / "epithelial_protein_bundle.zip"
    if bundle.exists():
        raise ValueError("Bundle already exists")
    freeze = json.loads((OUT / "analysis_freeze.json").read_text())
    paths = {p for p in OUT.rglob("*") if p.is_file()}
    paths |= {ROOT/r["path"] for r in freeze["inputs"]+freeze["source_code"]}
    with zipfile.ZipFile(bundle,"x",compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in sorted(paths):
            z.write(p,p.relative_to(ROOT).as_posix())
    json_out(dest,dict(utc=utc(), status="sealed", files=[record(p) for p in sorted(paths) if p.is_relative_to(OUT)],
        dependencies=freeze["inputs"], source_code=freeze["source_code"], documentation=[record(ROOT/"rgo.md")],
        archive=record(bundle), prior_integrity=previous, delivery_checks=verification,
        scope="Two epithelial tables / full fixed families; no ASM controlled validation",online_model_calls=0))
    verify()


def verify():
    data = json.loads((OUT / "epithelial_manifest.json").read_text())
    for group in ["files","dependencies","source_code","documentation"]:
        for r in data[group]:
            assert record(ROOT/r["path"])==r, r["path"]
    assert record(ROOT/data["archive"]["path"])==data["archive"]
    all_records = {r["path"]:r for group in ["files","dependencies","source_code"] for r in data[group]}
    import hashlib
    with zipfile.ZipFile(ROOT/data["archive"]["path"]) as z:
        assert z.testzip() is None
        assert len(z.namelist())==len(set(z.namelist())) and set(z.namelist())==set(all_records)
        for name,r in all_records.items():
            blob = z.read(name)
            assert len(blob)==r["bytes"] and hashlib.sha256(blob).hexdigest()==r["sha256"],name
    assert delivery_checks()==data["delivery_checks"]
    assert prior_check()==data["prior_integrity"]
    print(json.dumps(dict(status="pass",new_files=len(data["files"]),dependencies=len(data["dependencies"]),
        archive_members=len(all_records),full_array_tests=data["delivery_checks"]["full_array_tests"],
        candidate_positions=224,workbook_cells=data["delivery_checks"]["workbook_cells_checked"],
        independent_R="pass",previous_download_seal="pass",online_model_calls=0),ensure_ascii=False,indent=2))


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("action",choices=["build","seal","verify"])
    globals()[p.parse_args().action]()
