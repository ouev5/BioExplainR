"""Check report tables against CSV and PDF/vector-figure geometry."""
import csv
import json
import sys
from pathlib import Path

import fitz
from docx import Document


def rows(p):
    with Path(p).open(newline="") as f:
        return list(csv.DictReader(f))


src, report = map(Path, sys.argv[1:3])
doc = Document(report / "BioExplainR-生物学发现-结果报告-v1.docx")
assert len(doc.tables) == 3 and len(doc.inline_shapes) == 5
summary = rows(src / "evidence_summary.csv")
table = doc.tables[0]
assert len(table.rows) == 13
for r, row in zip(summary, table.rows[1:]):
    values = [c.text for c in row.cells]
    assert values[0] == r["symbol"]
    if r["external_effect"]:
        assert abs(float(values[2])-float(r["external_effect"])) <= .00051
    else:
        assert values[2] == "不可估计"
for table_index, filename, condition, effectfield, pfield in [
    (1, "gr_candidate_results.csv", lambda r: r["symbol"] == "PER2", "log2_effect", "holm_all_tests"),
    (2, "direct_temporal_results.csv", lambda r: r["symbol"] in ["PER2", "PER3", "METTL7A", "MAOA"], "effect", "holm_24_tests")]:
    rr = [r for r in rows(src / filename) if condition(r)]
    for r, row in zip(rr, doc.tables[table_index].rows[1:]):
        values = [c.text for c in row.cells]
        position = 1 if table_index == 1 else 3
        assert abs(float(values[position])-float(r[effectfield])) <= .00051
        assert abs(float(values[-1])-float(r[pfield])) <= max(1e-14, float(r[pfield])*.006)
pdf = fitz.open(report / "BioExplainR-生物学发现-结果报告-v1.pdf")
text = "".join(p.get_text() for p in pdf)
for word in ["PER2", "PER3", "METTL7A", "54", "82,227", "不可估计", "LINC00607"]:
    assert word in text, word
outside = []; blocks = 0
for n, page in enumerate(pdf):
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                blocks += 1; box = fitz.Rect(span["bbox"])
                if not (page.rect + (-1, -1, 1, 1)).contains(box):
                    outside.append({"page": n+1, "text": span["text"], "bbox": list(box)})
assert not outside, outside
figures = []
for file in sorted((src / "figures").glob("*.pdf")):
    document = fitz.open(file); assert len(document) == 1
    page = document[0]; words = page.get_text("words"); assert len(words) > 15
    for w in words:
        assert (page.rect + (-1, -1, 1, 1)).contains(fitz.Rect(w[:4])), (file.name, w)
    assert len(page.get_drawings()) > 0
    figures.append({"file": file.name, "pages": 1, "text_words": len(words), "vector_paths": len(page.get_drawings()), "within_page": True})
assert len(figures) == 5
for lang in ["zh", "en"]:
    content = (src / ("BioExplainR-report-"+lang+"-v2.html")).read_text()
    assert "scope_limitations" in content and "Two experiments" in content
    for name in ["PER2", "PER3", "LINC00607"]:
        # The absent gene is carried as its stable feature ID in package contexts.
        assert name in content or name == "LINC00607" and "ENSG00000235770" in content
result = {"status": "pass", "pdf_pages": len(pdf), "word_tables": 3, "word_figures": 5,
          "table_data_source": "saved original CSV results", "pdf_text_spans_checked": blocks,
          "pdf_text_outside_page": outside, "figures": figures,
          "BioExplainR_v2_scope_limits_present": True, "visual_review": "not performed; model cannot read images"}
(report / "audit.json").write_text(json.dumps(result, indent=2)+"\n")
print(json.dumps(result, indent=2))
