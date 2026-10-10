"""Build vector panel PDFs, compose them as PDF pages, and export separate tables."""
import csv
import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path

import fitz
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import Circle, Ellipse, FancyArrowPatch, FancyBboxPatch, Rectangle
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/paper_cn_v1"
RESULTS = ROOT / "research/results"
CASES = RESULTS / "paper_cn_cases_v1"
MAE = RESULTS / "miniacc_installed_review"
COLORS = ["#247A6B", "#C45B52", "#387FA7", "#C09442", "#7E6D99"]
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 12, "pdf.fonttype": 42,
                     "axes.titlesize": 14, "axes.titleweight": "bold", "axes.spines.top": False,
                     "axes.spines.right": False, "savefig.facecolor": "white"})


def csv_rows(path):
    with path.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def js(path):
    return json.loads(path.read_text(encoding="utf-8"))


def numeric(rows, key):
    return np.array([float(r[key]) if r[key] not in ("", "NA") else np.nan for r in rows])


def panel(number, label, title, draw):
    fig, ax = plt.subplots(figsize=(6.4, 4.7))
    fig.subplots_adjust(left=.17, right=.96, bottom=.21, top=.82)
    if (number,label) in ((5,"A"),(6,"C")):
        fig.subplots_adjust(left=.28)
    fig.text(.035, .945, label, fontsize=20, weight="bold", color="#243E4B")
    fig.text(.105, .945, title, fontsize=14, weight="bold", color="#243E4B")
    draw(fig, ax)
    path = OUT / "figures/panels" / f"Figure_{number}{label}.pdf"
    fig.savefig(path)
    plt.close(fig)
    return path


def bars(ax, names, values, ylabel, colors=None):
    ax.bar(range(len(values)), values, color=colors or COLORS[:len(values)], width=.65)
    ax.set_xticks(range(len(values)), names)
    ax.set_ylabel(ylabel)
    ax.grid(axis="y", alpha=.16)
    ax.set_axisbelow(True)
    ax.set_ylim(0, max(values) * 1.24 if max(values) else 1)
    for i, value in enumerate(values):
        ax.text(i, value + max(values)*.035, f"{value:g}", ha="center", weight="bold", fontsize=12)


def card(ax, x, y, w, h, title, detail, color=COLORS[0]):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=.01,rounding_size=.025",
                               facecolor="#F4F8FA", edgecolor=color, linewidth=1.5))
    ax.text(x+.03, y+h-.01, title, va="top", weight="bold", fontsize=10 if h < .18 else 12, color=color)
    ax.text(x+.03, y+.012, detail, va="bottom", fontsize=8 if h < .18 else 10, linespacing=1.3)


def icon(ax, kind, x, y, size=.1, color=COLORS[0]):
    """Original vector icons: no external icon font or raster asset."""
    if kind == "database":
        ax.add_patch(Rectangle((x, y), size, size*.7, facecolor="white", edgecolor=color, lw=2))
        for dy in (0, .35, .7):
            ax.add_patch(Ellipse((x+size/2, y+dy*size), size, size*.28, facecolor="white", edgecolor=color, lw=2))
    elif kind == "chart":
        for i, h in enumerate((.4, .8, .6)):
            ax.add_patch(Rectangle((x+i*size*.33, y), size*.2, h*size, color=color))
    elif kind == "network":
        points = [(x, y), (x+size, y), (x+size*.5, y+size*.8)]
        for a, b in ((0,1),(1,2),(2,0)):
            ax.plot(*zip(points[a], points[b]), color=color, lw=2)
        for p in points:
            ax.add_patch(Circle(p, size*.14, facecolor="white", edgecolor=color, lw=2))
    elif kind == "check":
        ax.add_patch(Circle((x+size*.5,y+size*.4), size*.48, facecolor="white", edgecolor=color, lw=2))
        ax.plot([x+.2*size,x+.43*size,x+.8*size], [y+.4*size,y+.17*size,y+.65*size], color=color,lw=3)
    elif kind == "report":
        ax.add_patch(Rectangle((x,y),size*.8,size,facecolor="white",edgecolor=color,lw=2))
        for dy in (.25,.45,.65,.82):
            ax.plot([x+.15*size,x+.65*size],[y+dy*size]*2,color=color,lw=1.5)


def framework():
    def workflow(fig, ax):
        ax.axis("off"); ax.set(xlim=(0,1),ylim=(0,1))
        stages = [("database","R objects", "Full data / results", COLORS[0]),
                  ("chart","Deterministic facts", "Counts / scope / units",COLORS[0]),
                  ("network","Optional language", "Templates or model",COLORS[2]),
                  ("check","Separate checks", "Claims / plots / relations",COLORS[4]),
                  ("report","Traceable outputs", "Shiny / HTML / PDF",COLORS[0])]
        for i,(kind,title,detail,color) in enumerate(stages):
            y=.86-i*.18
            icon(ax,kind,.035,y-.03,.09,color)
            card(ax,.2,y-.065,.74,.145,title,detail,color)
            if i<4: ax.add_patch(FancyArrowPatch((.6,y-.075),(.6,y-.105),arrowstyle="-|>",mutation_scale=15,color="#526C7A"))
    def boundary(fig, ax):
        ax.axis("off"); ax.set(xlim=(0,1),ylim=(0,1))
        card(ax,.03,.5,.41,.35,"Local R computation","Full matrices\nSample maps / coordinates\nStatistics + source object")
        card(ax,.57,.5,.4,.35,"Model interface","Complete scalar facts\nBounded preview\nExplicit budget",COLORS[2])
        icon(ax,"database",.18,.39,.09); icon(ax,"network",.7,.39,.09,COLORS[2])
        ax.add_patch(FancyArrowPatch((.44,.67),(.55,.67),arrowstyle="-|>",mutation_scale=17,color=COLORS[2]))
        card(ax,.04,.04,.9,.22,"Budget exceeded: explicit failure","No silent fact truncation; no automatic chunking.",COLORS[1])
    def checks(fig, ax):
        ax.axis("off"); ax.set(xlim=(0,1),ylim=(0,1))
        for i,(name,detail,color) in enumerate([( "pass","Specified value agrees",COLORS[0]),("fail","Specified value disagrees",COLORS[1]),("partial","Check scope is incomplete",COLORS[3]),("unverified","No sufficient explicit claim", "#7D8990")]):
            y=.82-i*.21
            icon(ax,"check",.04,y-.035,.08,color)
            card(ax,.2,y-.075,.74,.14,name,detail,color)
    def project(fig, ax):
        ax.axis("off"); ax.set(xlim=(0,1),ylim=(0,1))
        card(ax,.25,.68,.5,.22,"Cohort / source","Explicit denominator")
        card(ax,.02,.28,.43,.24,"RNA result","Feature + effect + units")
        card(ax,.55,.28,.43,.24,"Protein result","Feature + effect + units",COLORS[2])
        for end in ((.23,.53),(.77,.53)):
            ax.add_patch(FancyArrowPatch((.5,.67),end,arrowstyle="-|>",mutation_scale=17,color="#526C7A"))
        ax.add_patch(FancyArrowPatch((.46,.4),(.54,.4),arrowstyle="<->",mutation_scale=15,color=COLORS[4]))
        ax.text(.5,.1,"Same patients and derived outputs\nare not independent confirmation.",ha="center",fontsize=12,color="#526C7A")
    return [panel(1,label,title,draw) for label,title,draw in zip("ABCD",
        ["Object-to-report workflow","Computation and payload boundary","Explicit validation states","Dependency-aware projects"],
        [workflow,boundary,checks,project])]


def figures():
    all_panels = {1: framework()}
    de = csv_rows(CASES / "airway_de.csv")
    coords = csv_rows(CASES / "airway_pca.csv")
    var = numeric(csv_rows(CASES / "airway_variance.csv"), "variance")*100
    thresholds = csv_rows(CASES / "airway_threshold_checks.csv")
    def volcano(fig, ax):
        fc,p = numeric(de,"log2FoldChange"),numeric(de,"padj")
        good = np.isfinite(p)&np.isfinite(fc)
        significant = good&(p<.05)&(abs(fc)>1)
        ax.scatter(fc[good],-np.log10(np.maximum(p[good],1e-300)),s=2,c="#BDC8CE",rasterized=False)
        for mask,color in [(significant&(fc>0),COLORS[1]),(significant&(fc<0),COLORS[2])]:
            ax.scatter(fc[mask],-np.log10(p[mask]),s=4,c=color)
        ax.set(xlabel="log2 fold change (trt / untrt)",ylabel="-log10 adjusted P")
        ax.text(.04,.93,"999 selected: 519 up / 480 down",transform=ax.transAxes,fontsize=11)
    def pca(fig, ax):
        for group,color in [("untrt",COLORS[2]),("trt",COLORS[1])]:
            rows=[r for r in coords if r["treatment"]==group]
            ax.scatter(numeric(rows,"PC1"),numeric(rows,"PC2"),c=color,s=65,label=group)
        for donor in {r["donor_code"] for r in coords}:
            pair=[r for r in coords if r["donor_code"]==donor]
            ax.plot(numeric(pair,"PC1"),numeric(pair,"PC2"),color="#AEBEC7",zorder=0)
        ax.set(xlabel=f"PC1 ({var[0]:.2f}%)",ylabel=f"PC2 ({var[1]:.2f}%)")
        ax.legend(frameon=False); ax.text(.03,.03,"8 libraries / 4 paired donors",transform=ax.transAxes,fontsize=11)
    def sensitivity(fig, ax):
        for lfc,color in [(0,COLORS[0]),(1,COLORS[2]),(2,COLORS[1])]:
            rows=[r for r in thresholds if float(r["lfc"])==lfc and r["fact"]=="n_significant"]
            ax.plot(numeric(rows,"padj"),numeric(rows,"actual"),"o-",label=f"|log2FC| > {lfc}",color=color,lw=2)
        ax.set(xlabel="Adjusted P threshold",ylabel="Selected features"); ax.legend(frameon=False,fontsize=10)
        ax.set_xticks([.01,.05,.1]); ax.grid(alpha=.15)
    def methods(fig, ax):
        bars(ax,["DESeq2","limma","edgeR"],[999,1301,1292],"Selected features")
        ax.text(.5,.91,"Method-specific upstream results",transform=ax.transAxes,ha="center",fontsize=10)
    all_panels[2]=[panel(2,l,t,d) for l,t,d in zip("ABCD",["Full-table differential expression","Paired-design PCA","Threshold sensitivity: 45/45 facts","Multiple upstream object classes"],[volcano,pca,sensitivity,methods])]
    pbmc=csv_rows(CASES/"pbmc_embedding_qc.csv")
    def umap(fig, ax):
        for group,color in [("mt <= 5%",COLORS[0]),("mt > 5%",COLORS[1])]:
            rows=[r for r in pbmc if r["mito_qc"]==group]
            ax.scatter(numeric(rows,"UMAP1"),numeric(rows,"UMAP2"),c=color,s=3,label=group)
        ax.set(xlabel="UMAP 1",ylabel="UMAP 2"); ax.legend(frameon=False,markerscale=3,fontsize=10)
    def qc(fig, ax):
        ax.hist(numeric(pbmc,"percent.mt"),bins=45,color=COLORS[0]);ax.axvline(5,color=COLORS[1],ls="--",lw=2)
        ax.set(xlabel="Mitochondrial UMI (%)",ylabel="Cells")
    def units(fig, ax):
        ax.axis("off");ax.set(xlim=(0,1),ylim=(0,1))
        card(ax,.06,.57,.88,.27,"2,700 cells / 32,738 features","SCE and Seurat object summaries agree")
        card(ax,.06,.12,.88,.27,"1 donor","Cell count does not become donor count",COLORS[1])
    def exp_set(fig, ax):
        bars(ax,["ALL","bladderbatch"],[128,57],"Samples",[COLORS[2],COLORS[4]])
        ax.text(.03,.9,"ALL: 12,625 features / 10 BT groups\nbladderbatch: 22,283 features / 5 batches",transform=ax.transAxes,fontsize=10)
    all_panels[3]=[panel(3,l,t,d) for l,t,d in zip("ABCD",["PBMC3k embedding","QC distribution","Independent-unit safeguard","ExpressionSet design summaries"],[umap,qc,units,exp_set])]
    pc=csv_rows(RESULTS/"related_tools/pasilla_pca.csv")
    pv=csv_rows(RESULTS/"related_tools/pasilla_volcano.csv")
    pf=csv_rows(RESULTS/"related_tools/pasilla_fact_checks.csv")
    def pas_pca(fig, ax):
        for condition,color in zip(sorted({r["condition"] for r in pc}),COLORS):
            rows=[r for r in pc if r["condition"]==condition]
            ax.scatter(numeric(rows,"PC1"),numeric(rows,"PC2"),s=60,c=color,label=condition)
        ax.set(xlabel="PC1",ylabel="PC2");ax.legend(frameon=False,fontsize=10)
    def pas_volcano(fig, ax):
        fc,p=numeric(pv,"log2FC"),numeric(pv,"padj")
        ok=np.isfinite(fc)&np.isfinite(p);sig=ok&(p<.05)&(abs(fc)>1)
        ax.scatter(fc[ok],-np.log10(np.maximum(p[ok],1e-300)),s=3,c="#C8D2D6")
        ax.scatter(fc[sig],-np.log10(p[sig]),s=5,c=COLORS[1]);ax.set(xlabel="log2FC (treated / untreated)",ylabel="-log10 adjusted P")
    def pas_count(fig, ax):bars(ax,["Up","Down"],[108,110],"Selected features",[COLORS[1],COLORS[2]])
    def pas_check(fig, ax):
        for lfc,color in [(0,COLORS[0]),(1,COLORS[2])]:
            rows=[r for r in pf if float(r["lfc"])==lfc and r["fact"]=="n_significant"]
            ax.plot(numeric(rows,"alpha"),numeric(rows,"actual"),"o-",label=f"|log2FC| > {lfc}",color=color,lw=2)
        ax.set(xlabel="Adjusted P threshold",ylabel="Selected features");ax.legend(frameon=False,loc="lower right",fontsize=10);ax.set_xticks([.01,.05,.1])
        ax.text(.04,.94,"30/30 saved source-table checks",transform=ax.transAxes,fontsize=11)
    all_panels[4]=[panel(4,l,t,d) for l,t,d in zip("ABCD",["pasilla: separate-source PCA","Adjusted-P volcano","218 features at preset thresholds","Six threshold combinations"],[pas_pca,pas_volcano,pas_count,pas_check])]
    coverage=csv_rows(MAE/"modality_coverage.csv");features=csv_rows(MAE/"feature_results.csv")
    def cov(fig, ax):
        names=[r["modality"] for r in coverage];vals=numeric(coverage,"n_units")
        ax.barh(range(5),vals,color=COLORS[0]);ax.barh(range(5),92-vals,left=vals,color="#E6ECEF")
        ax.set_yticks(range(5),["RNA","Copy number","Protein","Mutation","miRNA"]);ax.set(xlabel="Patients",xlim=(0,104))
        for i,v in enumerate(vals):ax.text(v+1,i,f"{int(v)}/92",va="center",fontsize=11)
        ax.text(.03,1.025,"43/92 complete across all modalities",transform=ax.transAxes,fontsize=11)
    def effects(fig, ax):
        for direction,color in [("same",COLORS[0]),("opposite",COLORS[1])]:
            rows=[r for r in features if r["direction"]==direction]
            ax.scatter(numeric(rows,"rna_effect"),numeric(rows,"protein_effect"),c=color,s=45,label=f"{direction} ({len(rows)})")
        ax.axhline(0,color="#8E9DA5",ls="--",lw=1);ax.axvline(0,color="#8E9DA5",ls="--",lw=1)
        ax.set(xlabel="RNA mean difference (late - early)",ylabel="RPPA mean difference (late - early)")
        ax.legend(frameon=False,fontsize=10,loc="upper left")
    def bh(fig, ax):
        values=np.array([numeric(features,"rna_padj"),numeric(features,"protein_padj")])
        image=ax.imshow(values,aspect="auto",cmap="viridis_r",vmin=0,vmax=1)
        ax.set_yticks([0,1],["RNA","Protein"]);ax.set_xticks([0,8,16,24,32],["1","9","17","25","33"])
        ax.set_xlabel("Shared feature index (alphabetical)")
        fig.colorbar(image,ax=ax,fraction=.05,pad=.04,label="BH-adjusted P")
        ax.text(.02,1.10,"BH < 0.05: RNA 3 / protein 0",transform=ax.transAxes,fontsize=11)
    def design(fig, ax):
        bars(ax,["Early","Late"],[28,18],"Paired patients",[COLORS[2],COLORS[1]])
        ax.text(.04,.93,"46 paired patients / 33 shared features",transform=ax.transAxes,fontsize=11)
        ax.text(.04,.55,"Sign agreement is not\nsignificance or independence.",transform=ax.transAxes,fontsize=9,
                bbox={"facecolor":"white","edgecolor":"none","alpha":.9})
    all_panels[5]=[panel(5,l,t,d) for l,t,d in zip("ABCD",["Multimodal coverage denominator","Explicit direction agreement","Significance kept separate","Stage-comparison design"],[cov,effects,bh,design])]
    ab=csv_rows(RESULTS/"main_scale_480/ablation_summary.csv")
    fam=csv_rows(RESULTS/"main_scale_480/family_summary.csv")
    def matching(fig, ax):
        bars(ax,["Bounded\nraw","R facts","R facts +\nevidence"],numeric(ab,"accuracy")*100,"Target matching (%)",[COLORS[1],COLORS[0],COLORS[2]])
        ax.text(.03,.93,"437/1440     1440/1440     1440/1440",transform=ax.transAxes,fontsize=10)
    def coverage_plot(fig, ax):
        for i,key in enumerate(["coverage","answered_accuracy"]):
            ax.bar(np.arange(3)+(i-.5)*.32,numeric(ab,key)*100,width=.32,label=["Coverage","Answered accuracy"][i],color=[COLORS[3],COLORS[2]][i])
        ax.set_xticks(range(3),["Bounded\nraw","R facts","R facts +\nevidence"]);ax.set(ylabel="Rate (%)",ylim=(0,122));ax.legend(frameon=False,fontsize=10,loc="upper left")
    def family(fig, ax):
        families=sorted({r["family"] for r in fam});methods=["direct_bounded","facts_only","facts_evidence"]
        vals=[[float(next(r for r in fam if r["family"]==f and r["method"]==m)["accuracy"])*100 for m in methods] for f in families]
        ax.imshow(vals,cmap="YlGn",vmin=0,vmax=100,aspect="auto")
        ax.set_yticks(range(6),["Composition","DE","Enrichment","Generic","PCA","Single-cell"]);ax.set_xticks(range(3),["Raw","Facts","Facts + E"])
        for i,row in enumerate(vals):
            for j,v in enumerate(row):ax.text(j,i,f"{v:.1f}",ha="center",va="center",fontsize=11,color="#162B24")
    def models(fig, ax):
        bars(ax,["StepFun\nstep-5-preview","DeepSeek\ndeepseek-flash"],[480,480],"Matched target facts",[COLORS[2],COLORS[0]])
        ax.text(.5,.9,"120 calls / 480 targets per model",transform=ax.transAxes,ha="center",fontsize=10)
        ax.text(.5,.45,"Historical structured-transfer endpoint;\nno new provider calls in this revision.",transform=ax.transAxes,ha="center",fontsize=10,
                bbox={"facecolor":"white","edgecolor":"none","alpha":.9})
    all_panels[6]=[panel(6,l,t,d) for l,t,d in zip("ABCD",["Fixed-target input ablation","Coverage and answered accuracy","Six functional families (%)","Two-model protocol replay audit"],[matching,coverage_plot,family,models])]
    scenarios=csv_rows(CASES/"scenario_checks.csv")
    def claims(fig, ax):bars(ax,["Injected\nexplicit errors","Path only"],[72,18],"Controlled claims",[COLORS[1],COLORS[3]])
    def plots(fig, ax):bars(ax,["Correct\nplots","Injected\nplot errors"],[3,6],"Controlled plots",[COLORS[0],COLORS[1]])
    def new_cases(fig, ax):
        counts=Counter(r["expected"] for r in scenarios)
        bars(ax,["pass","fail","partial","unverified"],[counts[k] for k in ["pass","fail","partial","unverified"]],"Expected-state cases",[COLORS[0],COLORS[1],COLORS[3],"#7D8990"])
        ax.text(.5,.91,"15/15 expected states observed",transform=ax.transAxes,ha="center",fontsize=11)
    def repairs(fig, ax):
        ax.axis("off");ax.set(xlim=(0,1),ylim=(0,1))
        card(ax,.04,.58,.9,.28,"Reference preserved","Runtime OK + invariants_passed")
        card(ax,.04,.13,.9,.28,"Reference reversed","Runtime OK + invariants_failed",COLORS[1])
    all_panels[7]=[panel(7,l,t,d) for l,t,d in zip("ABCD",["Explicit-claim mechanism challenge","Source-to-plot challenge","Added development scenarios","Runnable versus semantic-preserving"],[claims,plots,new_cases,repairs])]
    browser=RESULTS/"w00_baseline/paper-cn-shiny-browser-final"
    def website(fig, ax):
        from PIL import Image
        img=Image.open(browser/"miniacc-project-desktop.png")
        # Crop the analysis content at its actual screenshot scale, never fabricate a screen.
        img=img.crop((0,0,img.width,min(img.height,1250)))
        ax.imshow(img);ax.axis("off")
    def mobile(fig, ax):
        from PIL import Image
        img=Image.open(browser/"miniacc-project-mobile.png")
        img=img.crop((0,0,img.width,min(img.height,1100)))
        ax.imshow(img);ax.axis("off")
    def pdf_download(fig, ax):
        doc=fitz.open(browser/"generated-pca.pdf")
        pix=doc[0].get_pixmap(matrix=fitz.Matrix(1.2,1.2))
        ax.imshow(np.frombuffer(pix.samples,dtype=np.uint8).reshape(pix.height,pix.width,pix.n));ax.axis("off")
    def engineering(fig, ax):
        ax.axis("off");ax.set(xlim=(0,1),ylim=(0,1))
        card(ax,.02,.68,.95,.19,"Source: 676 pass / 1 skip","677 assertions; no failures or warnings")
        card(ax,.02,.37,.95,.19,"Installed package: check OK","Imports-only workflow; legacy API 74/74")
        card(ax,.02,.06,.95,.19,"Browser acceptance: 7 groups","Project / PDF / EN-ZH / two sessions")
    all_panels[8]=[panel(8,l,t,d) for l,t,d in zip("ABCD",["Real miniACC project in Shiny","Mobile project display","Actual generated PDF download","Layered engineering checks"],[website,mobile,pdf_download,engineering])]
    return all_panels


def compose(number, paths):
    doc=fitz.open();page=doc.new_page(width=921.6,height=676.8)
    for i,path in enumerate(paths):
        with fitz.open(path) as source:
            rect=fitz.Rect((i%2)*460.8,(i//2)*338.4,(i%2+1)*460.8,(i//2+1)*338.4)
            page.show_pdf_page(rect,source,0)
    output=OUT/"figures/composites"/f"Figure_{number}.pdf"
    doc.save(output,deflate=True)
    page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).save(OUT/"figures/previews"/f"Figure_{number}.png")
    doc.close()
    return {"number":number,"composite":output.relative_to(ROOT).as_posix(),
            "panels":[p.relative_to(ROOT).as_posix() for p in paths],"assembly":"PyMuPDF show_pdf_page; 2x2 vector-preserving composition"}


def export_table(number,title,headers,rows,note,sources):
    path=OUT/"tables"/f"Table_{number}.xlsx"
    book=Workbook();sheet=book.active;sheet.title=f"Table {number}"
    sheet.append(headers)
    for row in rows:sheet.append(row)
    sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="247A6B")
    from openpyxl.utils import get_column_letter
    for col in range(1,len(headers)+1):
        width=max(len(str(sheet.cell(r,col).value or "")) for r in range(1,sheet.max_row+1))
        sheet.column_dimensions[get_column_letter(col)].width=min(max(width+2,16),65)
    for row in sheet:
        for cell in row:cell.alignment=Alignment(vertical="top",wrap_text=True)
    meta=book.create_sheet("图表说明与来源");meta.append(["标题",title]);meta.append(["表注",note])
    for source in sources:meta.append(["来源",source])
    book.save(path)
    with (OUT/"tables"/f"Table_{number}.csv").open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.writer(f);writer.writerow(headers);writer.writerows(rows)
    return {"number":number,"title":title,"headers":headers,"rows":rows,"note":note,"sources":sources,
            "xlsx":path.relative_to(ROOT).as_posix(),"csv":f"docs/paper_cn_v1/tables/Table_{number}.csv"}


def tables():
    specs=[]
    def add(*args):specs.append(export_table(*args))
    add(1,"案例数据、版本与分析边界",["数据","对象/来源版本","规模","分析边界"],
        [["airway","airway 1.22.0; DESeq2 1.42.0","8文库、4供体; 22,369保留特征","配对DE、PCA、ORA及阈值核对"],
         ["PBMC3k","SCE 1.24.0; SeuratObject 5.0.1","2,700细胞、32,738特征; 1供体","QC/UMAP/独立单位摘要"],
         ["pasilla","pasilla 1.30.0; DESeq2 1.42.0","7文库、14,599特征","6组阈值、30项保存数值核对"],
         ["ALL","ALL 1.44.0; Biobase 2.62.0","128样本、12,625特征","表型与10个BT分组摘要"],
         ["bladderbatch","bladderbatch 1.40.0","57样本、22,283特征","3癌症分组、5批次摘要"],
         ["miniACC","MultiAssayExperiment 1.28.0","92病例、5模态; 46配对","开发案例:覆盖、方向、BH及安装包复现"]],
        "各数据承担不同验证终点；对象/细胞/文库/病例不合并为研究样本量。",["research/results/data_manifest.json","research/results/public_datasets/dataset_tests.json","research/results/miniacc_installed_review/summary.json"])
    thresholds=csv_rows(CASES/"airway_threshold_checks.csv")
    add(2,"airway 阈值敏感性及适配器核对",["P_adj阈值","|log2FC|阈值","事实","源表值","适配器值","一致"],
        [[float(r["padj"]),float(r["lfc"]),r["fact"],int(r["expected"]),int(r["actual"]),r["pass"]] for r in thresholds],
        "9组阈值×5事实=45项。严格使用P_adj<α、|log2FC|>τ；NA不作零。",["research/results/paper_cn_cases_v1/airway_threshold_checks.csv"])
    pf=csv_rows(RESULTS/"related_tools/pasilla_fact_checks.csv")
    add(3,"pasilla 独立来源数值事实核对",["P_adj阈值","|log2FC|阈值","事实","直接源表值","适配器值","一致"],
        [[float(r["alpha"]),float(r["lfc"]),r["fact"],int(r["expected"]),int(r["actual"]),r["pass"]] for r in pf],
        "沿用0.3.1保存结果，30/30；本轮没有重拟合pasilla上游模型。",["research/results/related_tools/pasilla_fact_checks.csv"])
    c=csv_rows(MAE/"modality_coverage.csv")
    add(4,"miniACC 模态覆盖及统一分母",["模态","观测病例","缺失病例","主病例分母"],[[r["modality"],int(r["n_units"]),int(r["n_missing"]),92] for r in c],
        "映射pass表示观测ID对应合法，不表示五模态齐全；完全观测43病例、未观测0病例。",["research/results/miniacc_installed_review/modality_coverage.csv","research/results/miniacc_installed_review/alignment_summary.json"])
    f=csv_rows(MAE/"feature_results.csv");h=list(f[0])
    add(5,"miniACC 33特征效应、方向及显著性",h,[[float(r[k]) if k not in ("feature_id","direction") else r[k] for k in h] for r in f],
        "RNA为log2(归一化表达+1)，蛋白保持RPPA值；效应为late均值−early均值。两模态量纲不同，不比较绝对效应大小。BH分别在33特征内计算。",["research/results/miniacc_installed_review/feature_results.csv"])
    ab=csv_rows(RESULTS/"main_scale_480/ablation_summary.csv")
    add(6,"保存的主消融实验",["输入","匹配","目标","匹配率","覆盖率","已回答准确率","API请求","中位秒"],
        [[r["method"],int(r["correct"]),int(r["expected"]),float(r["accuracy"]),float(r["coverage"]),float(r["answered_accuracy"]),int(r["calls"]),float(r["median_elapsed_sec"])] for r in ab],
        "120程序案例、480唯一事实、3方法×3重复；1,080实际请求、4,320评分。100%是显式事实传递终点。",["research/results/main_scale_480/ablation_summary.csv"])
    family=csv_rows(RESULTS/"main_scale_480/family_summary.csv")
    add(7,"六类功能家族与双模型事实传递",["家族/模型","方法","匹配","已回答","目标"],
        [[r["family"],r["method"],int(r["correct"]),int(r["answered"]),int(r["expected"])] for r in family]+[["step-5-preview","同输入单轮",480,480,480],["deepseek-flash","同输入单轮",480,480,480]],
        "家族汇总属主消融三轮；双模型各120请求，只评价结构化事实传递。未配置代理和缺失replay不进入API分母。",["research/results/main_scale_480/family_summary.csv","research/results/common_task_scale_480/common_task_summary.csv"])
    scenarios=csv_rows(CASES/"scenario_checks.csv")
    add(8,"新增开发场景的期望状态与实际状态",list(scenarios[0]),[list(r.values()) for r in scenarios],
        "15/15期望一致；受控修改标明controlled mutation，手算fixture不冒充公开生物数据。",["research/results/paper_cn_cases_v1/scenario_checks.csv"])
    add(9,"工程验收与复现层级",["层级","结果","记录","范围"],
        [["源码","676 pass / 1 skip / 677","paper-cn-source.*","0失败/错误/警告"],
         ["包检查","Status: OK","paper-cn-package-check.*","本轮实际构建包"],
         ["历史API","74/74","review-legacy-api*-result.json","10-08源码及安装包名称/位置参数"],
         ["浏览器","7组pass","paper-cn-shiny-browser-final/summary.json","项目、PDF、英中、两会话、双视口"],
         ["miniACC安装复现","8核心表逐字节相同","review-miniacc-boundaries-final.*","同一数据工程复现"],
         ["Python复核","66效应/方向/覆盖/BH","review-miniacc-installed-audit.*","原始Wilcoxon P未独立重算"]],
        "工程检查和数值复核不代替独立专家、用户或科学泛化评价。",["research/results/w00_baseline/paper-cn-source.txt","research/results/w00_baseline/paper-cn-shiny-browser-final/summary.json","docs/implementation_records/REVIEW_CONTINUATION.md"])
    return specs


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--refresh-draft",action="store_true",help="Regenerate this session's unfrozen draft assets")
    args=parser.parse_args()
    for name in ["figures/panels","figures/composites","figures/previews","tables"]:
        (OUT/name).mkdir(parents=True,exist_ok=True)
    if (OUT/"assets_manifest.json").exists() and not args.refresh_draft:
        raise ValueError("Assets already frozen; choose a new version directory")
    figures_spec=[compose(n,paths) for n,paths in figures().items()]
    specs=tables()
    files=[p for directory in (OUT/"figures",OUT/"tables") for p in directory.rglob("*") if p.is_file()]
    manifest={"status":"generated","figures":figures_spec,"tables":specs,"files":[
        {"path":p.relative_to(ROOT).as_posix(),"bytes":p.stat().st_size,"sha256":hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(files)]}
    (OUT/"assets_manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(f"Generated {len(figures_spec)} PDF composites from 32 panel PDFs, {len(specs)} separate CSV/XLSX tables")


if __name__=="__main__":main()
