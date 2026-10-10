"""Audit Chinese manuscript contents, numeric sources, vector layout and external tables."""
import argparse
import csv
import hashlib
import json
import math
import re
from pathlib import Path
import fitz
from docx import Document
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/"docs/paper_cn_v1"
NAME="BioExplainR-中文论文-v0.8"


def require(condition,message):
    if not condition:raise ValueError(message)


def rows(path):
    with path.open(encoding="utf-8-sig",newline="") as f:return list(csv.DictReader(f))


def geometry(path):
    d=fitz.open(path);p=d[0]
    require(len(d)==1,f"Figure not one page: {path}")
    spans=[(s["text"],fitz.Rect(s["bbox"])) for b in p.get_text("dict")["blocks"] if "lines" in b
           for line in b["lines"] for s in line["spans"] if s["text"].strip()]
    for text,rect in spans:
        require(p.rect.contains(rect),f"Figure text outside page: {path.name}: {text}")
    for i,(a,x) in enumerate(spans):
        for b,y in spans[i+1:]:
            rect=x&y
            require(rect.is_empty or rect.width*rect.height <= .04*min(x.width*x.height,y.width*y.height),
                    f"Overlapping figure text: {path.name}: {a} / {b}")
    require(len(p.get_drawings())>0 or len(p.get_images())>0,f"Empty figure {path}")
    result={"file":path.relative_to(ROOT).as_posix(),"text_spans":len(spans),"drawings":len(p.get_drawings()),
            "images":len(p.get_images()),"page_bounds":"pass","text_overlap":"pass"}
    d.close();return result


def audit():
    text=(OUT/"manuscript.md").read_text(encoding="utf-8")
    body,appendix=text.split("# 附录 A：",1)
    require("## 2.11 外部评分" not in text and "1.751" not in text and "91.98" not in text,"Human results present")
    require(text.rstrip().endswith("内容审定后整理。"),"Future supplement must be last")
    references=body.split("# 参考文献\n",1)[1]
    numbers=[int(n) for n in re.findall(r"^(\d+)\. ",references,re.M)]
    require(numbers==list(range(1,30)),"Reference numbering mismatch")
    mainbody=body.split("# 参考文献\n",1)[0]
    cited=set()
    for value in re.findall(r"\[([\d,，–—\- ]+)\]",mainbody):
        for part in re.split(r"[,，]",value):
            limits=re.split(r"[–—\-]",part.strip())
            if len(limits)==2:cited.update(range(int(limits[0]),int(limits[1])+1))
            else:cited.add(int(limits[0]))
    require(cited==set(numbers),f"Uncited/missing references: {cited ^ set(numbers)}")
    for n in range(1,9):require(f"图 {n}" in mainbody and f"**图 {n}." in appendix,f"Missing figure citation/caption {n}")
    for n in range(1,10):require(f"表 {n}" in mainbody and f"## 表 {n}" in appendix,f"Missing table citation/note {n}")
    doc=Document(OUT/f"{NAME}.docx")
    require(len(doc.tables)==0 and len(doc.inline_shapes)==8,"DOCX embedded table/figure count")
    require(len(doc.sections)==1,"Unexpected landscape sections")
    html=(OUT/f"{NAME}.html").read_text(encoding="utf-8")
    require("<table" not in html and html.count("<img ")==8,"HTML table/figure count")
    manifest=json.loads((OUT/"assets_manifest.json").read_text())
    for r in manifest["files"]:
        p=ROOT/r["path"]
        require(p.stat().st_size==r["bytes"] and hashlib.sha256(p.read_bytes()).hexdigest()==r["sha256"],f"Asset hash {p}")
    require(len(manifest["figures"])==8 and len(manifest["tables"])==9,"Asset count")
    for table in manifest["tables"]:
        book=load_workbook(ROOT/table["xlsx"],read_only=True,data_only=True)
        data=list(book.worksheets[0].values)
        require(data[0]==tuple(table["headers"]) and len(data)==len(table["rows"])+1,"XLSX row/header count")
        for actual,expected in zip(data[1:],table["rows"]):
            for a,e in zip(actual,expected):
                require(math.isclose(a,e,rel_tol=1e-12,abs_tol=1e-15) if isinstance(e,(int,float)) else a==e,
                        f"XLSX data {table['number']}: {a} != {e}")
        with (ROOT/table["csv"]).open(encoding="utf-8-sig",newline="") as f:
            values=list(csv.reader(f))
        require(values==[table["headers"]]+[[str(v) for v in row] for row in table["rows"]],"CSV/table mismatch")
        book.close()
    de=rows(ROOT/"research/results/paper_cn_cases_v1/airway_de.csv")
    thresholds=rows(ROOT/"research/results/paper_cn_cases_v1/airway_threshold_checks.csv")
    for r in thresholds:
        a,b=float(r["padj"]),float(r["lfc"])
        chosen=[d for d in de if d["padj"] and float(d["padj"])<a and abs(float(d["log2FoldChange"]))>b]
        expected={"n_features":len(de),"n_significant":len(chosen),"n_up":sum(float(d["log2FoldChange"])>0 for d in chosen),
                  "n_down":sum(float(d["log2FoldChange"])<0 for d in chosen),"n_na_padj":sum(not d["padj"] for d in de)}[r["fact"]]
        require(expected==int(r["actual"])==int(r["expected"]),"Airway independent recomputation")
    from sys import path as module_paths
    module_paths.insert(0,str(ROOT/"research"))
    from multiomics.audit_miniacc import audit as miniacc
    multi=miniacc(ROOT/"research/results/miniacc_installed_review")
    require(multi["status"]=="pass","miniACC recheck")
    scenarios=rows(ROOT/"research/results/paper_cn_cases_v1/scenario_checks.csv")
    require(len(scenarios)==15 and all(r["expected"]==r["actual"] for r in scenarios),"Scenarios mismatch")
    geometries=[geometry(p) for p in sorted((OUT/"figures/panels").glob("*.pdf"))]
    composites=[geometry(p) for p in sorted((OUT/"figures/composites").glob("*.pdf"))]
    require(len(geometries)==32 and len(composites)==8,"PDF counts")
    d=fitz.open(OUT/f"{NAME}.pdf")
    require(all(abs(p.rect.width-595.3)<1 for p in d),"Non-portrait paper page")
    pdftext="\n".join(p.get_text() for p in d)
    require("附录 C" in pdftext and "用户使用体验" in pdftext and "1.751" not in pdftext,"PDF sections mismatch")
    vector=json.loads((OUT/"export_report.json").read_text())["vector_composites_embedded"]
    for r in vector:
        p=d[r["page"]-1]
        require(len(p.get_drawings())>0,f"Composite missing vector content {r['figure']}")
        require(f"图 {r['figure']}." in p.get_text(),"Caption separated from figure")
    for p in d:
        for b in p.get_text("dict")["blocks"]:
            if "lines" not in b:continue
            for line in b["lines"]:
                for span in line["spans"]:
                    require(p.rect.contains(fitz.Rect(span["bbox"])),f"Paper text outside page {p.number+1}")
    downloaded=fitz.open(ROOT/"research/results/w00_baseline/paper-cn-shiny-browser-final/generated-pca.pdf")
    require(len(downloaded)==1 and len(downloaded[0].get_drawings())>0,"Downloaded PDF not vector")
    result={"status":"pass","references":len(numbers),"figures":8,"panel_pdfs":32,"composite_pdfs":8,
            "external_tables":9,"docx_tables":0,"pdf_pages":len(d),"airway_independent_numeric_rechecks":45,
            "development_scenarios":15,"miniacc":multi,"geometry":geometries+composites,
            "visual_review_scope":"Automated text bounding/overlap and content audit; model cannot view images/PDF directly"}
    d.close();downloaded.close();return result


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument("--output",type=Path,required=True)
    args=parser.parse_args()
    result=audit()
    args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({k:v for k,v in result.items() if k!="geometry"},ensure_ascii=False,indent=2))


if __name__=="__main__":main()
