"""Create simple, answer-key-free distribution packages for human studies."""
import csv
import io
import shutil
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "research/study_materials"


def workbook_with_rater(source, rater_id):
    workbook = load_workbook(source)
    sheet = workbook["Ratings"]
    for row in sheet.iter_rows(min_row=2):
        row[1].value = rater_id
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def csv_with_rater(source, rater_id):
    rows = []
    with source.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = reader.fieldnames
        for row in reader:
            row["rater_id"] = rater_id
            rows.append(row)
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=fields)
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8")


def add_file(archive, path, name):
    archive.writestr(name, path.read_bytes())


def build_expert_package(output):
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("README.md", """# 专家评审包

请将 `expert_E01`、`expert_E02`、`expert_E03` 三个文件夹分别发给三位专家。每位专家只填写自己文件夹中的一个 Excel 或 CSV 表，不要互相交换评分。

填写方法：

1. 打开 `expert_review_workbook.xlsx` 的 `Ratings` 工作表。
2. 每个 `R001-R030` 是一个匿名解释案例；每个案例有 8 行，需要给每个维度填一个分数：`0` = 错误/缺失，`1` = 部分正确，`2` = 正确且可追溯。
3. `comment` 可简短写一句依据，例如“数字正确，但把富集写成通路激活”。没有意见可留空。
4. 不需要填写 `Case_Index`，不要猜测数据集、模型或作者身份。
5. 填完后保存 Excel；CSV 是同一份空表的备用格式，二选一即可。

匿名案例在根目录的 `expert_cases/` 中。不要打开或传播研究负责人保留的答案密钥；本包不包含答案密钥。
""")
        for rater_id in ("E01", "E02", "E03"):
            folder = f"expert_{rater_id}"
            archive.writestr(f"{folder}/expert_review_workbook.xlsx",
                             workbook_with_rater(SOURCE / "expert_review_workbook.xlsx", rater_id))
            archive.writestr(f"{folder}/expert_ratings_template.csv",
                             csv_with_rater(SOURCE / "expert_ratings_template.csv", rater_id))
            add_file(archive, SOURCE / "expert_review_index.csv", f"{folder}/expert_review_index.csv")
        for path in sorted((SOURCE / "expert_cases").glob("R*.json")):
            add_file(archive, path, f"expert_cases/{path.name}")


def build_user_package(output):
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("README.md", """# 用户功能评分表

这个包只有一个 Excel 评分表，预置 30 位学生 × 5 项 BioExplainR 功能 = 150 行。

填写方法：

1. 每位学生只填写自己对应的 5 行。
2. `score_0_to_100` 填 0-100 分。
3. `useful_1_yes_2_no` 填 1（有用）或 2（没用）。
4. `comment` 可以写一句简短意见，没有意见可留空。
5. 不要填写姓名、邮箱或临床身份信息，也不要编造评分。

五项功能已经写在表格中：运行并记录 R 代码、解释 R 生物信息学对象、生成常用分析图形、证据与来源追踪、生成分析报告。
""")
        add_file(archive, SOURCE / "user_study_workbook.xlsx", "user_study_workbook.xlsx")


def main():
    expert = SOURCE / "expert_review_distribution.zip"
    user = SOURCE / "user_study_distribution.zip"
    build_expert_package(expert)
    build_user_package(user)
    for path in (expert, user):
        with zipfile.ZipFile(path) as archive:
            assert archive.testzip() is None
            assert "investigator_answer_key.json" not in archive.namelist()
    print(f"expert={expert} bytes={expert.stat().st_size}")
    print(f"user={user} bytes={user.stat().st_size}")


if __name__ == "__main__":
    main()
