"""Build v0.9 with a concise main text, documented tool comparison and separate supplement."""
import csv
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys

import fitz
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/paper_cn_v2"
OLD = ROOT / "docs/paper_cn_v1"
CASE = ROOT / "research/results/discovery_sparcl1_v1"
sys.path.insert(0, str(ROOT / "research"))
from export_paper import native_docx


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def json_write(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def record(path):
    return {"path": path.relative_to(ROOT).as_posix(), "bytes": path.stat().st_size,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def external_table(name, headers, rows, note):
    csv_path = OUT / "tables" / f"{name}.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f); writer.writerow(headers); writer.writerows(rows)
    book = Workbook(); sheet = book.active; sheet.title = "Data"
    sheet.append(headers)
    for row in rows: sheet.append(row)
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]: cell.font = Font(bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 25
        for cell in column: cell.alignment = Alignment(wrap_text=True, vertical="top")
    book.create_sheet("说明与来源").append([note])
    book.save(OUT / "tables" / f"{name}.xlsx")


def discovery_figure():
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "pdf.fonttype": 42,
                         "axes.spines.top": False, "axes.spines.right": False})
    colors = ["#387FA7", "#C09442", "#247A6B"]
    figure, axes = plt.subplots(2, 2, figsize=(11, 7.5))
    figure.subplots_adjust(left=.16, right=.95, bottom=.11, top=.9, wspace=.45, hspace=.5)
    candidate = read_csv(CASE / "candidate.csv")[0]
    axes[0, 0].axis("off")
    axes[0, 0].text(0, .92, "A  Discovery and frozen candidate", fontsize=12, weight="bold")
    axes[0, 0].text(0, .7, "airway / GSE52778: 4 paired donors\nFull DESeq2 results; padj < 0.05, log2FC > 1", va="top")
    axes[0, 0].text(0, .40, "Rank by padj, then feature ID\nFirst candidate: SPARCL1 / ENSG00000152583", va="top", color=colors[2])
    axes[0, 0].text(0, .12, f"Discovery log2FC = +{float(candidate['log2FoldChange']):.3f}\nFreeze before candidate-expression testing", va="top")
    measurements = read_csv(CASE / "gene_measurements.csv")
    ax = axes[0, 1]
    ax.set_title("B  All external culture replicates", loc="left", fontsize=12, weight="bold", pad=12)
    for i, group in enumerate(["untreated", "dex4h", "dex24h"]):
        values = [float(r["value"]) for r in measurements if r["group"] == group]
        jitter = [(j-(len(values)-1)/2)*.08 for j in range(len(values))]
        ax.scatter([i+j for j in jitter], values, color=colors[i], s=38)
    ax.set_xticks([0, 1, 2], ["Untreated\nn=4", "Dex 4h\nn=3", "Dex 24h\nn=3"])
    ax.set_ylabel("Published processed-expression units")
    ax.grid(axis="y", alpha=.15)
    ax = axes[1, 0]
    ax.set_title("C  Predefined contrasts and Welch CI", loc="left", fontsize=12, weight="bold", pad=12)
    results = read_csv(CASE / "validation_results.csv")
    for i, r in enumerate(results):
        effect = float(r["effect"]); lo = float(r["ci_low"]); hi = float(r["ci_high"])
        ax.errorbar(effect, i, xerr=[[effect-lo], [hi-effect]], fmt="o", capsize=5, color=colors[2-i], markersize=7)
        ax.text(effect, i+.16, f"P = {float(r['pvalue']):.2e}", ha="center", fontsize=9)
    ax.axvline(0, color="#999999", linewidth=1, linestyle="--")
    ax.set_yticks([0, 1], ["24h (primary)", "4h (secondary)"])
    ax.set(xlabel="Mean difference: dex - untreated", xlim=(-.25, 5), ylim=(-.4, 1.5))
    ax = axes[1, 1]; ax.axis("off")
    ax.text(0, .94, "D  Evidence and next experiments", fontsize=12, weight="bold")
    ax.text(0, .70, "External experiment: GSE34313 / GPL6480\nOne HASM1 cell line; culture replicates\nNot an independent-donor cohort", va="top")
    ax.text(0, .37, "Primary permutation P = 1/35 = 0.02857\nExpression response: preliminarily replicated\nMechanism / protein / clinical effects: untested", va="top", color=colors[2])
    ax.text(0, .06, "Next: new donors, GR perturbation,\nSPARCL1 intervention and matrix endpoints", va="top")
    pdf_path = OUT / "figures/Figure_2.pdf"
    figure.savefig(pdf_path); figure.savefig(OUT / "figures/Figure_2.png", dpi=190)
    plt.close(figure)


def renumber_figure_text(text):
    return re.sub(r"图\s*([1-8])", r"补充图 S\1", text)


def make_sources(comparison):
    old = (OLD / "manuscript.md").read_text(encoding="utf-8")
    refs = old.split("# 参考文献\n", 1)[1].split("# 附录 A：", 1)[0].strip()
    refs += "\n30. Masuno K, et al. Expression profiling identifies Klf15 as a glucocorticoid target that regulates airway hyperresponsiveness. Am J Respir Cell Mol Biol. 2011;45:642–649. doi: [10.1165/rcmb.2010-0369OC](https://doi.org/10.1165/rcmb.2010-0369OC)."
    headers, rows = comparison["headers"], comparison["rows"]
    markdown = "| " + " | ".join(headers) + " |\n|" + "---|" * len(headers)
    markdown += "\n" + "\n".join("| " + " | ".join(row) + " |" for row in rows)
    main = (OUT / "manuscript.md").read_text(encoding="utf-8")
    main = main.replace('::: {.revision-references}\n:::', refs).replace('::: {.revision-table}\n:::', markdown)
    (OUT / "manuscript_expanded.md").write_text(main, encoding="utf-8")
    methods = old.split("# 2 方法\n", 1)[1].split("# 3 结果", 1)[0]
    methods = re.sub(r"## 2\.(\d+)", r"## S\1", methods)
    methods = renumber_figure_text(methods)
    methods = re.sub(r"表\s*([1-9])", r"补充表 S\1", methods)
    figures = old.split("# 附录 A：", 1)[1].split("# 附录 B：", 1)[0]
    figures = renumber_figure_text(figures).replace("paper_cn_v1/", "paper_cn_v2/")
    figures = figures.replace("论文 PDF 的本附录", "补充 PDF 的本部分")
    tables = old.split("# 附录 B：", 1)[1].split("# 附录 C：", 1)[0]
    tables = re.sub(r"表\s*([1-9])", r"补充表 S\1", tables)
    tables = re.sub(r"Table_(\d)", r"Table_S\1", tables)
    scope = "本补充保留 v0.8 的详细方法及历史结果图表，软件版本与执行日期按原记录标注。S1–S10 是既有方法，S11 为本轮外部发现案例。主稿聚焦四项贡献；本文件并不新增专家或用户评分。参考文献编号与主稿一致。"
    supplement = '---\ntitle: "BioExplainR 中文论文 v0.9：补充材料"\ndate: "2026-10-08"\nlang: zh-CN\n---\n\n' + scope + "\n\n# 补充方法\n\n" + methods
    supplement += "\n## S11 候选固定、外部实验及独立实现复核\n\n" + (ROOT / "research/discovery/PROTOCOL.md").read_text(encoding="utf-8").split("\n", 1)[1]
    supplement += "\n\n外部数值结果见补充表 S10；Python/SciPy 与 R 的十项效应/P/CI/置换值及两项 Holm 值相符。BioExplainR 生成两个 scope/unit 明确区分的上下文与项目，报告并不把软件 pass 解释为机制成立。\n\n# 补充结果图\n\n" + figures
    supplement += "\n# 补充结果表\n\n" + tables
    supplement += "\n## 补充表 S10 SPARCL1 外部实验与重算\n\n**表注：** 24 小时为主比较，4 小时为次要比较；两侧 Welch、95% CI、两个时间点 Holm 值和 35 种分组枚举置换分别提供。培养重复来自同一细胞系，不作为独立供体。**来源：** `research/results/discovery_sparcl1_v1/validation_results.csv`、公开 GEO 输入及 `r_numeric_checks.csv`。**文件：** `tables/Table_S10.csv`、`tables/Table_S10.xlsx`。\n\n# 后续研究与投稿信息\n\n真实作者、基金、伦理、独立专家评价与人工图稿审阅待负责人确认；用户研究 deferred。DOI 只有归档平台实际签发后填写。\n"
    (OUT / "supplementary.md").write_text(supplement, encoding="utf-8")
    old_body = old.split("# 1 引言", 1)[1].split("# 数据、代码与稿件信息", 1)[0]
    new_body = main.split("# 1 引言", 1)[1].split("# 数据与代码可用性", 1)[0]
    json_write(OUT / "compression.json", {"metric": "non-whitespace Unicode characters from introduction through conclusion",
        "v0_8_body": len(re.sub(r"\s", "", old_body)), "v0_9_body": len(re.sub(r"\s", "", new_body)),
        "core_contributions": 4, "technical_methods_moved": 10, "supplementary_methods": 11})


def export(source, name, specs, figures):
    # Use actual Word superscript runs so the minus sign cannot disappear with font fallback.
    text = source.read_text(encoding="utf-8")
    digits = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹", "0123456789")
    text = re.sub(r"10⁻([⁰¹²³⁴⁵⁶⁷⁸⁹]+)", lambda m: "10^-"+m[1].translate(digits)+"^", text)
    rendered = OUT / f"{source.stem}_rendered.md"
    rendered.write_text(text, encoding="utf-8")
    base = ["pandoc", str(rendered), "--standalone", f"--resource-path={ROOT/'docs'}:{ROOT}"]
    ast = json.loads(subprocess.check_output(base + ["-t", "json"]))
    doc = native_docx(ast, OUT / f"{name}.docx", specs, legacy_layout=False)
    # Headings contain their literal number once. Disable inherited Word numbering.
    for style in doc.styles:
        if style.name.startswith("Heading"):
            pr = style.element.get_or_add_pPr()
            existing = pr.find(qn("w:numPr"))
            if existing is not None: pr.remove(existing)
            num = OxmlElement("w:numPr"); ident = OxmlElement("w:numId"); ident.set(qn("w:val"), "0")
            num.append(ident); pr.append(num)
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading") and (p.text in figures or p.text == "图表"):
            p.paragraph_format.page_break_before = True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        from docx.shared import Pt
                        run.font.size = Pt(8.5)
    doc.save(OUT / f"{name}.docx")
    subprocess.run(base + ["--embed-resources", f"--css={OUT/'paper.css'}", "-o", str(OUT / f"{name}.html")], check=True)
    intermediate = OUT / "rendered_docx"; intermediate.mkdir(exist_ok=True)
    subprocess.run(["libreoffice", "-env:UserInstallation=file:///tmp/opencode/bioexplainr-paper-v09-lo", "--headless",
        "--convert-to", "pdf", "--outdir", str(intermediate), str(OUT / f"{name}.docx")], check=True, timeout=180)
    pdf = fitz.open(intermediate / f"{name}.pdf"); replaced = []
    for page in pdf:
        images = page.get_image_info(xrefs=True)
        if not images: continue
        text = page.get_text()
        matches = [caption for caption in figures if caption in text]
        if len(matches) != 1 or len(images) != 1:
            raise ValueError(f"Ambiguous figure page {page.number+1}: {matches}, images {len(images)}")
        rect = fitz.Rect(images[0]["bbox"])
        page.add_redact_annot(rect, fill=(1, 1, 1)); page.apply_redactions(images=2, graphics=0, text=0)
        with fitz.open(figures[matches[0]]) as vector:
            page.show_pdf_page(rect, vector, 0, keep_proportion=True)
        replaced.append({"caption": matches[0], "page": page.number+1, "pdf": record(figures[matches[0]])})
    assert len(replaced) == len(figures)
    pdf.set_metadata({"title": name, "author": "", "subject": "BioExplainR revision v0.9"})
    pdf.save(OUT / f"{name}.pdf", garbage=4, deflate=True)
    result = {"name": name, "pages": len(pdf), "docx_tables": len(doc.tables),
              "figures": len(doc.inline_shapes), "vector_figures": replaced}
    pdf.close(); return result


def main():
    for sub in ["tables", "figures", "figures/composites", "figures/panels", "figures/previews"]:
        (OUT / sub).mkdir(parents=True, exist_ok=True)
    comparison = json.loads((Path(__file__).parent / "tool_comparison.json").read_text(encoding="utf-8"))
    external_table("Table_1", comparison["headers"], comparison["rows"], comparison["scope"])
    for n in range(1, 10):
        for ext in ["csv", "xlsx"]:
            (OUT / f"tables/Table_S{n}.{ext}").write_bytes((OLD / f"tables/Table_{n}.{ext}").read_bytes())
    for old in (OLD / "figures").rglob("*"):
        if old.is_file():
            dest = OUT / "figures" / old.relative_to(OLD / "figures")
            dest.parent.mkdir(parents=True, exist_ok=True); dest.write_bytes(old.read_bytes())
    for ext, directory in [("pdf", "composites"), ("png", "previews")]:
        (OUT / f"figures/Figure_1.{ext}").write_bytes((OLD / f"figures/{directory}/Figure_1.{ext}").read_bytes())
    rows = read_csv(CASE / "validation_results.csv")
    external_table("Table_S10", list(rows[0]), [[r[k] for k in rows[0]] for r in rows],
        "GSE34313; fixed SPARCL1 candidate; single-cell-line culture replication; Welch/CI/Holm/permutation")
    discovery_figure()
    make_sources(comparison)
    spec = {"id": "comparison", "headers": comparison["headers"], "rows": comparison["rows"],
            "widths_cm": [2.5, 2.6, 2.7, 3.1, 2.0, 4.1]}
    main_export = export(OUT / "manuscript_expanded.md", "BioExplainR-中文论文-v0.9", [spec],
        {"图 1 BioExplainR 的事实约束工作流": OUT / "figures/Figure_1.pdf",
         "图 2 SPARCL1 候选固定与外部实验初步重复": OUT / "figures/Figure_2.pdf"})
    captions = ["平台框架", "配对转录组案例", "单细胞与微阵列案例", "pasilla 数值基准", "miniACC 多组学案例", "模型事实传递", "受控错误与新增场景", "工作台与工程验收"]
    supplement = export(OUT / "supplementary.md", "BioExplainR-补充材料-v0.9", [],
        {f"补充图 S{n} {caption}": OUT / f"figures/composites/Figure_{n}.pdf" for n, caption in enumerate(captions, 1)})
    json_write(OUT / "export_report.json", {"status": "pass", "main": main_export, "supplement": supplement})
    assets = [p for sub in ["tables", "figures"] for p in (OUT / sub).rglob("*") if p.is_file()]
    json_write(OUT / "assets_manifest.json", {"files": [record(p) for p in sorted(assets)],
        "comparison_sources": comparison["sources"], "reuse": "Historical v0.8 figures/tables copied byte-for-byte; discovery figure newly generated"})
    print(json.dumps({"main": main_export, "supplement": supplement}, ensure_ascii=False, indent=2))


if __name__ == "__main__": main()
