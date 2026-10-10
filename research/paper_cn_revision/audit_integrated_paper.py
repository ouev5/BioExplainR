"""Verify v0.13 against preserved source evidence, donor arithmetic and exports."""
import argparse
import ast
from collections import Counter, defaultdict
import contextlib
import hashlib
import io
import json
import math
import re
import sys
import zipfile
from pathlib import Path

import fitz
import numpy as np
from scipy import stats
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "research/discovery"))
from set_validation import rows
from validation_sources import record, utc, json_out
from build_integrated_paper import PAPER, OLD, SECRETOME, PROVENANCE, PUBLIC, EV, TABLE_SOURCES, draft_json


def prior_evidence():
    """Only substitute the latest historical root-document dependency."""
    import audit_public_proteome as previous
    manifest = json.loads(previous.MANIFEST.read_text())
    snapshot = PAPER / "evidence/rgo_before_integration.md"
    saved = record(snapshot)
    expected = next(r for r in manifest["documentation"] if r["path"] == "rgo.md")
    assert (saved["bytes"], saved["sha256"]) == (expected["bytes"], expected["sha256"])
    original = previous.record

    def snapshot_record(path):
        if path.resolve() == (ROOT / "rgo.md").resolve():
            return dict(original(snapshot), path="rgo.md")
        return original(path)

    previous.record = snapshot_record
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            result = previous.verify()
    finally:
        previous.record = original
    before = snapshot.read_bytes()
    current = (ROOT / "rgo.md").read_bytes()
    marker = "## 最新续接：公开蛋白组分页补齐与样本结构复核（2026-10-10）".encode()
    assert current[current.index(marker):] == before[before.index(marker):]
    assert current[:current.index("## 最新续接：".encode())] == before[:before.index("## 最新续接：".encode())]
    return result


def holm(values):
    p = np.asarray(values, dtype=float)
    order = np.argsort(p, kind="stable")
    adjusted = np.minimum(1, np.maximum.accumulate(p[order] * np.arange(len(p), 0, -1)))
    result = np.empty(len(p))
    result[order] = adjusted
    return result


def bh(values):
    p = np.asarray(values, dtype=float)
    order = np.argsort(p, kind="stable")
    adjusted = np.minimum(1, np.minimum.accumulate((p[order] * len(p) / np.arange(1, len(p) + 1))[::-1])[::-1])
    result = np.empty(len(p))
    result[order] = adjusted
    return result


def numbers():
    full = rows(SECRETOME / "all_results.csv")
    candidate = rows(SECRETOME / "candidate_results.csv")
    assert len(full) == 846 and len(candidate) == 112
    assert np.allclose(bh([float(r["pvalue"]) for r in full]), [float(r["pooled_BH"]) for r in full], atol=1e-12, rtol=1e-9)
    assert np.allclose(holm([float(r["pvalue"]) for r in candidate]), [float(r["holm112"]) for r in candidate], atol=1e-12, rtol=1e-9)
    assert np.allclose(holm([float(r["exact_pvalue"]) for r in candidate]), [float(r["exact_holm112"]) for r in candidate], atol=1e-12, rtol=1e-9)
    by_gene = defaultdict(list)
    for r in rows(SECRETOME / "candidate_donor_effects.csv"):
        if r["status"] == "paired":
            by_gene[(r["symbol"], r["time_h"])].append(float(r["effect"]))
    checked = 0
    for r in candidate:
        values = by_gene[(r["symbol"], r["time_h"])]
        assert len(values) == int(r["n_pairs"])
        if r["status"] != "estimable":
            assert r["effect"] == r["ci_low"] == r["ci_high"] == ""
            assert float(r["pvalue"]) == float(r["holm112"]) == 1
            continue
        n = len(values)
        mean = float(np.mean(values))
        se = float(np.std(values, ddof=1) / np.sqrt(n))
        interval = stats.t.interval(.95, n - 1, loc=mean, scale=se)
        p = float(stats.ttest_1samp(values, 0).pvalue)
        for key, value in [("effect", mean), ("se", se), ("ci_low", interval[0]), ("ci_high", interval[1]), ("pvalue", p)]:
            assert math.isclose(float(r[key]), value, rel_tol=1e-8, abs_tol=1e-10), (r["symbol"], key)
        assert int(r["positive_donors"]) == sum(v > 0 for v in values)
        assert int(r["negative_donors"]) == sum(v < 0 for v in values)
        checked += 1
    summary = rows(SECRETOME / "array_summary.csv")
    assert [(r["time_h"], r["donor_labels"], r["technical_columns"], r["estimable_groups"], r["pooled_BH_significant"]) for r in summary] == [("24", "6", "36", "194", "7"), ("72", "8", "48", "153", "16")]
    assert checked == 4 and sum(float(r["holm112"]) < .05 for r in candidate) == 1
    for time in ["24", "72"]:
        assert sum(r["status"] == "estimable" for r in full if r["time_h"] == time) == int(next(r for r in summary if r["time_h"] == time)["estimable_groups"])
    assert sum(float(r["exact_holm112"]) < .05 for r in candidate) == 0
    sensitivity = rows(SECRETOME / "candidate_sensitivity.csv")
    for kind in ["published_six", "strict_three"]:
        data = [r for r in sensitivity if r["analysis"] == kind]
        assert len(data) == 112
        assert np.allclose(holm([float(r["pvalue"]) for r in data]), [float(r["holm112"]) for r in data], atol=1e-12, rtol=1e-9)
    display = rows(PAPER / "tables/Table_4.csv")
    for r in display:
        gene, time = r["基因/时点"].split("/")
        source = candidate if r["范围"] == "主分析" else sensitivity
        original = next(s for s in source if s["symbol"] == gene and s["time_h"] == time[:-1] and (r["范围"] == "主分析" or s["analysis"] == ("strict_three" if r["范围"] == "严格3/3" else "published_six")))
        assert r["配对n"] == original["n_pairs"]
        assert math.isclose(float(r["Holm112 P"]), float(original["holm112"]), rel_tol=.006, abs_tol=1e-12)
        if original["status"] == "estimable":
            assert abs(float(r["效应"]) - float(original["effect"])) <= .00051
            lo, hi = map(float, r["点态95%CI"].split("至"))
            assert abs(lo - float(original["ci_low"])) <= .00051 and abs(hi - float(original["ci_high"])) <= .00051
        else:
            assert r["效应"] == "不可估计" and r["点态95%CI"] == "—"
    matches = rows(PROVENANCE / "cross_array_summary.csv")
    same = [r for r in matches if r["expected_array"] == r["array"] and r["measure"] == "LFQ intensity"]
    assert {r["array"]: (int(r["original_nonzero"]), int(r["matches_rounding"])) for r in same} == {
        "proteome24": (4543, 4543), "proteome72": (14829, 14829), "secretome24": (803, 0), "secretome72": (2128, 2128)}
    a = next(r for r in same if r["array"] == "secretome24")
    assert (a["all_intensity_equal"], a["all_positive_equal"], a["original_zero"]) == ("9", "0", "25")
    versions = rows(EV / "version_evidence_complete.csv")
    assert len(versions) == 18 and all(r["software"] == "3.1.0.638" and r["protein_abundance_columns"] == "0" and r["full_final_export"] == "False" for r in versions)
    decoded = rows(PUBLIC / "HBEC_sample_decode.csv")
    cells = Counter((r["post_ICS"], r["polyIC"], r["T2high"]) for r in decoded if r["imiquimod"] == "0")
    assert len(decoded) == 43 and len({r["individual"] for r in decoded}) == 10 and len(cells) == 8 and set(cells.values()) == {5}
    return dict(secretome_family=846, candidate_family=112, estimable_candidates=checked, candidate_Holm_pass=1,
                candidate_exact_pass=0, candidate_donor_arithmetic=True, sensitivity_families=2,
                source_comparison_arrays=4, EV_single_run_versions=18, HBEC_balanced_cells=8)


def audit():
    prior = prior_evidence()
    frozen = json.loads((PAPER / "integration_freeze.json").read_text())
    for r in frozen["inputs"]:
        assert record(ROOT / r["path"]) == r, r["path"]
    original_script = ROOT / frozen["script"]["path"]
    snapshot = PAPER / "evidence/build_integrated_initial.py"
    original = snapshot if snapshot.exists() else original_script
    assert (record(original)["bytes"], record(original)["sha256"]) == (frozen["script"]["bytes"], frozen["script"]["sha256"])
    assert frozen["new_candidate_statistics"] is False
    assets = json.loads((PAPER / "assets_manifest.json").read_text())
    for r in assets["files"]:
        assert record(ROOT / r["path"]) == r
    for r in assets["copied_tables"] + assets["figure_sources"]:
        assert r["new"]["sha256"] == r["source"]["sha256"] and record(ROOT / r["source"]["path"]) == r["source"]
    for item in assets["table_sources"]:
        assert record(ROOT / item["source"]["path"]) == item["source"]
        assert rows(PAPER / f'tables/Table_{item["label"]}.csv') == rows(ROOT / item["source"]["path"])
    for r in assets["editorial_sources"]:
        assert record(ROOT / r["path"]) == r
    pairs = 0
    for path in sorted((PAPER / "tables").glob("*.xlsx")):
        data = rows(path.with_suffix(".csv"))
        wb = load_workbook(path, data_only=True, read_only=True)
        sheet = wb["Data"] if "Data" in wb.sheetnames else wb["数据"]
        assert [["" if v is None else str(v) for v in row] for row in sheet.values] == [list(data[0])] + [list(r.values()) for r in data], path.name
        wb.close()
        pairs += 1
    assert pairs == 51
    numeric = numbers()
    md = (PAPER / "manuscript_expanded.md").read_text()
    body, remaining = md.split("# 参考文献", 1)
    citations = set()
    for match in re.findall(r"\[([\d,–\-]+)\]", body):
        for part in match.split(","):
            if "–" in part or "-" in part:
                a, b = map(int, re.split("[–-]", part))
                citations.update(range(a, b + 1))
            else:
                citations.add(int(part))
    assert citations == set(range(1, 20))
    assert len(re.findall(r"^\d+\. ", remaining.split("# 图表", 1)[0], re.M)) == 19
    ref17 = remaining.split("17. ", 1)[1].split("\n", 1)[0]
    metadata = json.loads((ROOT / assets["editorial_sources"][0]["path"]).read_text())["resultList"]["result"][0]
    assert metadata["title"] in ref17 and ref17.startswith(metadata["authorString"].split(",")[0])
    for term in ["45/54", "53/53", "LINC03140", "Holm54", "Holm112", "846", "1.535", ".270", ".187",
                 "0/803", "4,543", "14,829", "2,128", "3.1.0.638", "未冻结／启动", "草稿、未发送", "5／5", "纠正历史", "临床ICS前后", "同一来源", "本稿没有新增候选", "结果后探索"]:
        assert term in md, term
    compact = re.sub(r"\s", "", body)
    # Expand explicitly cited inclusive table ranges before checking each table.
    for start, end in re.findall(r"S(\d+)[–−-]S(\d+)", compact):
        compact += " ".join(f"S{i}" for i in range(int(start), int(end) + 1))
    for term in ["表1", "表2", "表3", "表4", "图1", "图2", "图3", "图4", "S27", "S28", "S29", "S30", "S31", "S32", "S33", "S34", "S35", "S36", "S37", "S38", "S39", "S40", "S41", "S42", "补充图S5", "补充图S6"]:
        assert term in compact, term
    for source in [PAPER / "manuscript.md", PAPER / "supplementary.md"]:
        for relative in re.findall(r"!?\[[^\]]*\]\(([^)]+)\)", source.read_text()):
            if "://" not in relative:
                assert (ROOT / "docs" / relative).is_file(), (source.name, relative)
    exports = json.loads((PAPER / "export_report.json").read_text())
    documents = []
    for name, part, ntables, nfigures in [("BioExplainR-中文论文-v0.13", "main", 4, 4), ("BioExplainR-补充材料-v0.13", "supplement", 5, 6)]:
        document = Document(PAPER / (name + ".docx"))
        assert len(document.tables) == ntables and len(document.inline_shapes) == nfigures
        names = [f"Table_{i}" for i in range(1, 5)] if part == "main" else ["F1_common_task", "F2_paired_bootstrap", "F3_receptor_control", "F4_chip_primary", "Table_S26"]
        for table, filename in zip(document.tables, names):
            data = rows(PAPER / "tables" / (filename + ".csv"))
            assert [[c.text for c in r.cells] for r in table.rows] == [list(data[0])] + [list(r.values()) for r in data]
        with fitz.open(PAPER / (name + ".pdf")) as pdf:
            assert len(pdf) == exports[part]["pages"]
            spans = 0
            texts = "".join(page.get_text() for page in pdf)
            for term in ["PER2", "FBN2", "PXD055968", "846", "803", "391"]:
                assert term in texts, (name, term)
            for page in pdf:
                for block in page.get_text("dict")["blocks"]:
                    for line in block.get("lines", []):
                        for span in line["spans"]:
                            assert (page.rect + (-1, -1, 1, 1)).contains(fitz.Rect(span["bbox"])), (name, page.number + 1, span["text"])
                            spans += 1
            for item in exports[part]["vector_figures"]:
                page = pdf[item["page"] - 1]
                assert record(ROOT / item["pdf"]["path"]) == item["pdf"]
                with fitz.open(ROOT / item["pdf"]["path"]) as original:
                    assert len(page.get_drawings()) >= len(original[0].get_drawings()) > 0
                    assert sorted(r["digest"] for r in page.get_image_info(hashes=True)) == sorted(r["digest"] for r in original[0].get_image_info(hashes=True))
            documents.append(dict(name=name, pages=len(pdf), tables=ntables, figures=nfigures, spans=spans, original_PDF_embedding=True))
    figures = []
    for path in sorted((PAPER / "figures").glob("*.pdf")):
        with fitz.open(path) as pdf:
            assert len(pdf) == 1
            for block in pdf[0].get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        assert (pdf[0].rect + (-1, -1, 1, 1)).contains(fitz.Rect(span["bbox"])), (path.name, span["text"])
            figures.append(dict(file=path.name, images=len(pdf[0].get_images()), drawings=len(pdf[0].get_drawings())))
    for path in Path(__file__).parent.glob("*integrated*.py"):
        ast.parse(path.read_text(), str(path))
        assert all(s == s.rstrip() for s in path.read_text().splitlines())
    result = dict(status="pass", utc=utc(), prior=prior, numeric=numeric, csv_xlsx_pairs=pairs, references=19,
                  documents=documents, figures=figures, new_candidate_statistics=False,
                  scope="Saved statistics/donor arithmetic/source integrity, table/Word parity, links, PDF geometry; no human signoff")
    if not (PAPER / "manuscript_manifest.json").exists():
        draft_json(PAPER / "audit.json", result)
    print(json.dumps({k: v for k, v in result.items() if k != "prior"}, ensure_ascii=False, indent=2))
    return result


def archive_checks(path, entries):
    expected = {r["path"]: r for r in entries}
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        assert len(archive.namelist()) == len(set(archive.namelist())) == len(expected)
        assert set(archive.namelist()) == set(expected)
        for name, r in expected.items():
            data = archive.read(name)
            assert len(data) == r["bytes"] and hashlib.sha256(data).hexdigest() == r["sha256"]


def seal():
    manifest = PAPER / "manuscript_manifest.json"
    archive = PAPER / "BioExplainR-中文论文-v0.13-附件.zip"
    for p in [manifest, archive]:
        if p.exists():
            raise ValueError(f"Refusing overwrite {p}")
    result = audit()
    browser = json.loads((PAPER / "browser_final_v2/audit.json").read_text())
    assert browser["status"] == "pass"
    assert len(browser["inputs"]) == 2 and len(browser["results"]) == 4
    for r in browser["inputs"]:
        assert record(ROOT / r["path"]) == r
    assert "## 最新续接：论文v0.13" in (ROOT / "rgo.md").read_text()
    paths = [p for p in sorted(PAPER.rglob("*")) if p.is_file()]
    files = [record(p) for p in paths]
    codes = [Path(__file__).resolve(), SCRIPT_PATH, Path(__file__).with_name("capture_integrated_draft.py"), Path(__file__).with_name("browser_integrated_paper.cjs"), Path(__file__).with_name("INTEGRATION_PROTOCOL.md"),
             Path(__file__).with_name("build_revision.py"), Path(__file__).with_name("build_discovery_paper.py"),
             ROOT / "research/export_paper.py", ROOT / "research/discovery/set_validation.py", ROOT / "research/discovery/validation_sources.py"]
    source_code = [record(p) for p in codes]
    documentation = [record(ROOT / "rgo.md")]
    dependencies = {r["path"]: r for r in json.loads((PAPER / "integration_freeze.json").read_text())["inputs"]}
    for path in [EV / "AUTHOR_MATERIAL_REQUEST.md", ROOT / "research/results/discovery_asm_ev_materials_v1/MATERIALS_RECEIPT.md"]:
        r = record(path)
        dependencies[r["path"]] = r
    for item in json.loads((PAPER / "assets_manifest.json").read_text())["copied_tables"] + json.loads((PAPER / "assets_manifest.json").read_text())["figure_sources"]:
        dependencies[item["source"]["path"]] = item["source"]
    for r in json.loads((PAPER / "assets_manifest.json").read_text())["editorial_sources"]:
        dependencies[r["path"]] = r
    dependency_list = list(dependencies.values())
    portable = {r["path"]: r for r in files + source_code + documentation + dependency_list}
    with zipfile.ZipFile(archive, "x", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as zipped:
        for r in portable.values():
            zipped.write(ROOT / r["path"], arcname=r["path"])
    json_out(manifest, dict(status="sealed", utc=utc(), version="v0.13", files=files, source_code=source_code,
        documentation=documentation, dependencies=dependency_list, archive=record(archive), archive_members=list(portable.values()),
        prior_evidence=result["prior"], new_candidate_statistics=False, publication="paused", human_signoff="not performed",
        scope="Saved-evidence manuscript and direct result dependencies; original arrays indexed by prior seals",
        snapshot_policy="Only previous public_proteome root rgo dependency uses exact pre-integration snapshot; all scientific paths/hash original"))
    verify()


def verify():
    m = json.loads((PAPER / "manuscript_manifest.json").read_text())
    for group in ["files", "source_code", "documentation", "dependencies"]:
        for r in m[group]:
            assert record(ROOT / r["path"]) == r, r["path"]
    assert record(ROOT / m["archive"]["path"]) == m["archive"]
    archive_checks(ROOT / m["archive"]["path"], m["archive_members"])
    browser = json.loads((PAPER / "browser_final_v2/audit.json").read_text())
    for r in browser["inputs"]:
        assert record(ROOT / r["path"]) == r
    result = audit()
    result = dict(status="pass", files=len(m["files"]), archive_members=len(m["archive_members"]), archive_bytes=m["archive"]["bytes"],
                  previous_evidence=result["prior"], scope=m["scope"])
    (PAPER / "verification.json").write_text(json.dumps(dict(utc=utc(), result=result, note="Verification output excluded from own manifest"), ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({k: v for k, v in result.items() if k != "previous_evidence"}, ensure_ascii=False, indent=2))


SCRIPT_PATH = Path(__file__).with_name("build_integrated_paper.py")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=["audit", "seal", "verify"])
    globals()[parser.parse_args().action]()
