"""Verify the concise manuscript, supplement, comparison table and discovery provenance."""
import csv
import hashlib
import json
from pathlib import Path
import re
import sys

import fitz
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/paper_cn_v2"
sys.path.insert(0, str(ROOT / "research"))
sys.path.insert(0, str(ROOT / "research/discovery"))
from paper_cn.audit_manuscript import geometry
from audit_saved import audit as discovery_audit


def main():
    manifest = json.loads((OUT / "assets_manifest.json").read_text())
    for row in manifest["files"]:
        p = ROOT / row["path"]
        assert len(p.read_bytes()) == row["bytes"] and hashlib.sha256(p.read_bytes()).hexdigest() == row["sha256"]
    text = (OUT / "manuscript_expanded.md").read_text()
    body, refs = text.split("# 参考文献\n", 1)
    references = re.findall(r"^(\d+)\. ", refs, re.M)
    assert references == [str(n) for n in range(1, 31)]
    cited = set()
    for expression in re.findall(r"\[([\d,，–—\- ]+)\]", body):
        for part in re.split(r"[,，]", expression):
            limits = re.split(r"[–—\-]", part.strip())
            cited.update(range(int(limits[0]), int(limits[-1])+1))
    assert cited == set(range(1, 31))
    assert "1.1.1" not in text and "CRISPR2" not in text and "CRISPLD2" in text
    comparison = json.loads((Path(__file__).parent / "tool_comparison.json").read_text())
    expected = [comparison["headers"]] + comparison["rows"]
    with (OUT / "tables/Table_1.csv").open(encoding="utf-8-sig") as f: assert list(csv.reader(f)) == expected
    doc = Document(OUT / "BioExplainR-中文论文-v0.9.docx")
    assert len(doc.tables) == 1 and len(doc.inline_shapes) == 2
    assert [[cell.text for cell in row.cells] for row in doc.tables[0].rows] == expected
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert len(headings) == len(set(headings)) and not any(re.match(r"^(\d\.){3}", h) for h in headings)
    for style in doc.styles:
        if style.name.startswith("Heading"):
            assert style.element.find(".//" + qn("w:numId")).get(qn("w:val")) == "0"
    for name in ["Table_1"] + [f"Table_S{n}" for n in range(1, 11)]:
        with (OUT / f"tables/{name}.csv").open(encoding="utf-8-sig") as f: data = list(csv.reader(f))
        book = load_workbook(OUT / f"tables/{name}.xlsx", read_only=True, data_only=True)
        xlsx = list(book.worksheets[0].values)
        assert len(data) == len(xlsx)
        for a, b in zip(data, xlsx):
            for v, w in zip(a, b):
                if isinstance(w, (float, int)):
                    assert abs(float(v)-w) <= max(1e-14, abs(w)*1e-12)
                else: assert v == ("" if w is None else w)
        book.close()
    reports = json.loads((OUT / "export_report.json").read_text())
    pdfs = {}
    for key, count, shapes in [("main", 1, 2), ("supplement", 0, 8)]:
        report = reports[key]; name = report["name"]
        document = Document(OUT / f"{name}.docx")
        assert len(document.tables) == count and len(document.inline_shapes) == shapes
        with fitz.open(OUT / f"{name}.pdf") as pdf:
            content = "\n".join(p.get_text() for p in pdf)
            assert len(pdf) == report["pages"] and "CRISPR2" not in content and "1.1.1" not in content
            for row in report["vector_figures"]:
                assert len(pdf[row["page"]-1].get_drawings()) > 0
            for page in pdf:
                for block in page.get_text("dict")["blocks"]:
                    for line in block.get("lines", []):
                        for span in line["spans"]:
                            assert page.rect.contains(fitz.Rect(span["bbox"])), (name, page.number+1, span["text"])
            pdfs[key] = {"pages": len(pdf), "tables": count, "figures": shapes}
    compression = json.loads((OUT / "compression.json").read_text())
    assert compression["core_contributions"] == 4 and compression["v0_9_body"] < compression["v0_8_body"] * .5
    discovery = discovery_audit()
    # Sensitive magnitudes are read from the saved results, not inferred from plot labels.
    result = json.loads((ROOT / "research/results/discovery_sparcl1_v1/summary.json").read_text())["results"][0]
    assert f"{result['effect']:.3f}" in body and "单一细胞系" in body and "DOI 尚未建立" in body
    figures = [geometry(OUT / "figures/Figure_2.pdf")]
    audit = {"status": "pass", "references": 30, "figures": 2, "docx_tables": 1,
             "external_tables": 11, "pdf_pages": pdfs["main"]["pages"], "supplement": pdfs["supplement"],
             "comparison_tools": 6, "body_reduction_percent": round(100*(1-compression["v0_9_body"]/compression["v0_8_body"]), 2),
             "discovery": discovery, "geometry": figures,
             "heading_numbering": "disabled inherited Word numbering; literal numbers only",
             "visual_scope": "Automated PDF text bounds/figure geometry and document content, not independent human review"}
    (OUT / "audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2)+"\n")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__": main()
