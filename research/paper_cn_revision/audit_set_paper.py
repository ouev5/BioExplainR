"""Audit v0.12 source-bound tables, references, narrative links and PDF exports."""
import argparse
import ast
import json
import re
import sys
from pathlib import Path
import fitz
import numpy as np
from docx import Document
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/"research/discovery"))
from set_validation import rows
from validation_sources import record,utc
from set_integrity import paper_integrity
SET=ROOT/"research/results/discovery_set_validation_v1"


def audit(out):
    historical=paper_integrity(SET)
    for name in ["numeric_audit.json","sensitivity_audit.json","r_tpm_audit.json","bioexplainr_workflow.json"]:
        assert json.loads((SET/name).read_text())["status"]=="pass"
    for name in ["r_audit.json","r_sensitivity_audit.json","final_audit.json","bioexplainr_workflow.json"]:
        assert json.loads((SET/"unselected_protein"/name).read_text())["status"]=="pass"
    assets=json.loads((out/"assets_manifest.json").read_text())
    for r in assets["files"]:assert record(ROOT/r["path"])==r,r["path"]
    for r in assets["copied_tables"]+assets["figure_sources"]:
        assert r["new"]["sha256"]==r["source"]["sha256"] and record(ROOT/r["source"]["path"])==r["source"]
    pairs=0
    for path in sorted((out/"tables").glob("*.xlsx")):
        rr=rows(path.with_suffix(".csv"));wb=load_workbook(path,data_only=True);sheet=wb["Data"] if "Data" in wb.sheetnames else wb["数据"]
        values=list(sheet.values)
        assert [["" if v is None else str(v) for v in row] for row in values]==[list(rr[0])]+[list(r.values()) for r in rr],path
        pairs+=1
    summary=rows(SET/"set_summary.csv")
    for r in rows(out/"tables/Table_3.csv"):
        original=next(a for a in summary if a["symbol"]==r["基因"])
        for field,key in [("TPM配对差","tpm_effect"),("计数效应","count_effect")]:assert abs(float(r[field])-float(original[key]))<.00051
        for field,key in [("Holm54 P","tpm_holm54"),("计数Holm54 P","count_holm54")]:assert np.isclose(float(r[field]),float(original[key]),rtol=.006,atol=1e-18)
    for item in assets["table_sources"]:
        for r in item["sources"]:assert record(ROOT/r["path"])==r
        if len(item["sources"])==1:assert rows(out/f'tables/Table_{item["label"]}.csv')==rows(ROOT/item["sources"][0]["path"])
    protein=SET/"unselected_protein"
    protein_targets=[next(r for r in rows(protein/name) if r["symbol"]=="TMT1A" and (not time or r.get("time_h")==time))
        for name,time in [("results.csv","24"),("results.csv","72"),("published_label_sensitivity.csv","24"),
                          ("published_label_sensitivity.csv","72"),("direct_time_results.csv",None),("direct_published_labels.csv",None)]]
    for displayed,original in zip(rows(out/"tables/Table_S26.csv"),protein_targets):
        assert abs(float(displayed["效应"])-float(original["effect"]))<.00051
        lows=displayed["点态95%CI"].split("至")
        assert abs(float(lows[0])-float(original["ci_low"]))<.00051 and abs(float(lows[1])-float(original["ci_high"]))<.00051
        p=float(displayed["校正P"].split("=")[1].replace("−","-"))
        expected=original.get("holm112",original.get("holm56"))
        assert np.isclose(p,float(expected),rtol=.006,atol=1e-18)
    md=(out/"manuscript_expanded.md").read_text();body,remaining=md.split("# 参考文献",1)
    citations=set()
    for m in re.findall(r"\[([\d,–\-]+)\]",body):
        for part in m.split(","):
            if "–" in part or "-" in part:
                a,b=map(int,re.split("[–-]",part));citations.update(range(a,b+1))
            else:citations.add(int(part))
    assert citations==set(range(1,17)),citations
    assert len(re.findall(r"^\d+\. ",remaining.split("# 图表",1)[0],re.M))==16
    required=["45/54","53/53","LINC00607","LINC03140","结果后探索","点态","未确认","109,636","GNAS",".7831"]
    required += [".0574",".00990","1.6.15.0","2998","142","Holm112","7/7"]
    for term in required:assert term in md,term
    compact=re.sub(r"\s","",body)
    for term in ["图1","图2","图3","表1","表2","表3","补充表S1","S1b","S2","S12","S13a","S13b","S13c","S13d","S14","S17","S18","S19","S20","S21","S22","S23","S24","S24b","S24c","S24d","S25","S25b","S26","补充图S1","补充图S2","补充图S3","补充图S4"]:assert term in compact,term
    exports=json.loads((out/"export_report.json").read_text());docs=[]
    for name,part,ntables,nfigures in [("BioExplainR-中文论文-v0.12","main",3,3),("BioExplainR-补充材料-v0.12","supplement",5,4)]:
        document=Document(out/(name+".docx"));assert len(document.tables)==ntables and len(document.inline_shapes)==nfigures
        files=[f"Table_{i}" for i in range(1,4)] if part=="main" else ["F1_common_task","F2_paired_bootstrap","F3_receptor_control","F4_chip_primary","Table_S26"]
        for t,filename in zip(document.tables,files):
            rr=rows(out/"tables"/(filename+".csv"));assert [[c.text for c in row.cells] for row in t.rows]==[list(rr[0])]+[list(r.values()) for r in rr]
        with fitz.open(out/(name+".pdf")) as pdf:
            assert len(pdf)==exports[part]["pages"];spans=0
            for page in pdf:
                for block in page.get_text("dict")["blocks"]:
                    for line in block.get("lines",[]):
                        for span in line["spans"]:
                            assert (page.rect+(-1,-1,1,1)).contains(fitz.Rect(span["bbox"])),(name,page.number+1,span["text"])
                            spans+=1
            for item in exports[part]["vector_figures"]:
                page=pdf[item["page"]-1]
                with fitz.open(ROOT/item["pdf"]["path"]) as original:
                    assert len(page.get_drawings())>=len(original[0].get_drawings())>0
                    assert sorted(r["digest"] for r in page.get_image_info(hashes=True))==sorted(r["digest"] for r in original[0].get_image_info(hashes=True))
            docs.append(dict(name=name,pages=len(pdf),tables=ntables,figures=nfigures,spans=spans,original_PDF_embedding=True))
    figures=[]
    for path in sorted((out/"figures").glob("*.pdf")):
        with fitz.open(path) as pdf:
            assert len(pdf)==1
            for block in pdf[0].get_text("dict")["blocks"]:
                for line in block.get("lines",[]):
                    for span in line["spans"]:assert (pdf[0].rect+(-1,-1,1,1)).contains(fitz.Rect(span["bbox"])),(path.name,span["text"])
            figures.append(dict(file=path.name,images=len(pdf[0].get_images()),drawings=len(pdf[0].get_drawings())))
    for path in list((ROOT/"research/discovery").glob("*set*.py"))+list(Path(__file__).parent.glob("*set*.py")):
        ast.parse(path.read_text(),str(path));assert all(s==s.rstrip() for s in path.read_text().splitlines()),path
    result=dict(status="pass",utc=utc(),historical=historical,source_counts=json.loads((SET/"numeric_audit.json").read_text())["set_counts"],
        csv_xlsx_pairs=pairs,references=16,documents=docs,figure_geometry=figures,
        scope="New statistical audits + source-bound tables, originalfigure embedding, citations, narrative and PDF geometry; no human signoff")
    if not (out/"manuscript_manifest.json").exists():(out/"audit.json").write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument("directory",type=Path)
    audit(parser.parse_args().directory.resolve())
